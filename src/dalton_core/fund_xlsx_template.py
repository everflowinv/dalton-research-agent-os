"""AMZN-source fund workbook style contract and openpyxl application plan.

The contract contains layout and presentation only.  Company names, financial
lines, values, formulas, and period labels are supplied by the export consumer.
"""
from __future__ import annotations

import hashlib
import json
import re
from copy import copy
from importlib.resources import files
from typing import Any, Mapping, Sequence


STYLE_SCHEMA_VERSION = "fund-xlsx-template-style-0.1"
PLAN_SCHEMA_VERSION = "fund-xlsx-template-apply-plan-0.1"
STYLE_RESOURCE = "fund_xlsx_template_style.json"
# Updated only when a new source-reviewed style contract is versioned.
STYLE_RESOURCE_SHA256 = "673062d7ebe25980743e5db82727f7c25d6de81281cdd22e90ae887aa580b1b1"

_SHEET_ROLES = ("valuation", "financials", "driver")
_MODEL_ROLES = frozenset({"financials", "driver"})
_ROW_STYLES = frozenset({
    "major_section", "section", "label", "subtotal", "growth_label",
    "valuation_section",
})
_CELL_STYLES = frozenset({
    "hardcoded_input", "cross_sheet_formula", "local_formula",
    "assumption_input",
})
_NUMBER_KINDS = frozenset({
    "amount", "amount_one_decimal", "percentage", "multiple", "date_month",
    "per_share",
})
_HEX64 = re.compile(r"[0-9a-f]{64}")
_CELL_RANGE = re.compile(r"[A-Z]+[1-9][0-9]*(?::[A-Z]+[1-9][0-9]*)?")


class FundXlsxTemplateError(ValueError):
    pass


def _need(value: Any, reason: str) -> None:
    if not value:
        raise FundXlsxTemplateError(reason)


def _resource_bytes() -> bytes:
    return files("dalton_core").joinpath(STYLE_RESOURCE).read_bytes()


def _canonical_hash(value: Any) -> str:
    wire = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n"
    return hashlib.sha256(wire.encode("utf-8")).hexdigest()


def load_fund_xlsx_template_style() -> dict[str, Any]:
    """Load the packaged, hash-pinned AMZN-source style contract."""

    payload = _resource_bytes()
    _need(hashlib.sha256(payload).hexdigest() == STYLE_RESOURCE_SHA256,
          "fund XLSX template style bytes differ from the reviewed contract")
    try:
        style = json.loads(payload)
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise FundXlsxTemplateError("fund XLSX template style is invalid JSON") from exc
    _need(isinstance(style, dict) and set(style) == {
        "schema_version", "template_ref", "source_proof", "sheet_order",
        "model_grid", "valuation_grid", "styles", "number_formats",
    }, "fund XLSX template style has an invalid closed shape")
    _need(style["schema_version"] == STYLE_SCHEMA_VERSION,
          "fund XLSX template style version is unsupported")
    _need(style["sheet_order"] == list(_SHEET_ROLES),
          "fund XLSX template sheet order differs")
    proof = style["source_proof"]
    _need(isinstance(proof, dict) and set(proof) == {
        "workbook_sha256", "exact_source_style_sha256", "theme_xml_sha256",
    } and all(isinstance(value, str) and _HEX64.fullmatch(value)
              for value in proof.values()),
          "fund XLSX template source proof differs")
    _need(set(style["styles"]) == _ROW_STYLES | _CELL_STYLES | {
        "unit_header", "period_header",
    },
          "fund XLSX template style vocabulary differs")
    _need(set(style["number_formats"]) == _NUMBER_KINDS,
          "fund XLSX template number formats differ")
    return style


def _text_list(values: Sequence[str], name: str) -> list[str]:
    _need(not isinstance(values, (str, bytes)) and values,
          f"{name} must be a non-empty sequence")
    result = []
    for value in values:
        _need(isinstance(value, str) and value.strip() == value and value,
              f"{name} contains an invalid label")
        result.append(value)
    _need(len(result) == len(set(result)), f"{name} contains a duplicate")
    return result


def _row_plan(value: Mapping[str, Sequence[Mapping[str, Any]]] | None,
              sheet_names: Mapping[str, str]) -> dict[str, list[dict[str, Any]]]:
    source = value or {}
    _need(isinstance(source, Mapping) and set(source).issubset(sheet_names),
          "row styles name an unmanaged sheet role")
    result: dict[str, list[dict[str, Any]]] = {}
    for role, rows in source.items():
        seen: set[int] = set()
        normalized = []
        for row in rows:
            _need(isinstance(row, Mapping) and set(row) == {"row", "style", "level"},
                  "row style has an invalid closed shape")
            number, style, level = row["row"], row["style"], row["level"]
            _need(isinstance(number, int) and not isinstance(number, bool) and number >= 2
                  and number not in seen, "row style has an invalid or duplicate row")
            _need(style in _ROW_STYLES, "row style is unsupported")
            _need(isinstance(level, int) and not isinstance(level, bool)
                  and 0 <= level <= 3, "row hierarchy level is unsupported")
            _need(role != "valuation" or level == 0,
                  "valuation rows cannot use the model hierarchy gutter")
            seen.add(number)
            normalized.append({"row": number, "style": style, "level": level,
                               "label_column": level + 1})
        result[role] = sorted(normalized, key=lambda item: item["row"])
    return result


def _cell_plan(value: Mapping[str, Sequence[Mapping[str, Any]]] | None,
               sheet_names: Mapping[str, str]) -> dict[str, list[dict[str, Any]]]:
    source = value or {}
    _need(isinstance(source, Mapping) and set(source).issubset(sheet_names),
          "cell styles name an unmanaged sheet role")
    result: dict[str, list[dict[str, Any]]] = {}
    for role, cells in source.items():
        normalized = []
        for cell in cells:
            _need(isinstance(cell, Mapping) and set(cell) == {"range", "style", "number_kind"},
                  "cell style has an invalid closed shape")
            cell_range, style, number_kind = cell["range"], cell["style"], cell["number_kind"]
            _need(isinstance(cell_range, str) and _CELL_RANGE.fullmatch(cell_range),
                  "cell style range is invalid")
            _need(style in _CELL_STYLES, "cell style is unsupported")
            _need(number_kind is None or number_kind in _NUMBER_KINDS,
                  "cell number kind is unsupported")
            normalized.append({"range": cell_range, "style": style,
                               "number_kind": number_kind})
        result[role] = normalized
    return result


def build_fund_xlsx_template_plan(
    *,
    sheet_names: Mapping[str, str],
    annual_periods: Sequence[str],
    quarterly_periods: Sequence[str],
    unit_labels: Mapping[str, str],
    annual_support_headers: Sequence[str] = ("CAGR",),
    hidden_periods: Sequence[str] = (),
    row_styles: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
    cell_styles: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
) -> dict[str, Any]:
    """Build a closed application plan around company-specific rows and periods.

    ``row_styles`` contains row numbers, presentation roles, and outline levels;
    it never contains line-item labels.  The consumer writes its own company
    specification labels in the returned per-row ``label_column`` before
    applying the plan.
    """

    _need(isinstance(sheet_names, Mapping) and set(sheet_names) == set(_SHEET_ROLES),
          "sheet names must bind valuation, financials, and driver")
    _need(all(isinstance(name, str) and name.strip() == name and name
              for name in sheet_names.values())
          and len(set(sheet_names.values())) == len(sheet_names),
          "sheet names must be distinct non-empty strings")
    _need(isinstance(unit_labels, Mapping) and set(unit_labels) == _MODEL_ROLES
          and all(isinstance(value, str) and value.strip() == value and value
                  for value in unit_labels.values()),
          "unit labels must bind the two model sheets")
    annual = _text_list(annual_periods, "annual periods")
    quarterly = _text_list(quarterly_periods, "quarterly periods")
    support = _text_list(annual_support_headers, "annual support headers")
    _need(not set(annual).intersection(quarterly), "annual and quarterly periods overlap")
    _need(not (set(support) & (set(annual) | set(quarterly))),
          "annual support headers overlap periods")
    hidden = list(hidden_periods)
    _need(len(hidden) == len(set(hidden)) and set(hidden).issubset(set(annual) | set(quarterly)),
          "hidden periods must be unique selected periods")
    first_annual = 5
    first_support_gutter = first_annual + len(annual)
    first_support = first_support_gutter + 1
    second_support_gutter = first_support + len(support)
    first_quarter = second_support_gutter + 1
    columns = {
        "period_row": 1,
        "hierarchy_label_columns": [1, 2, 3, 4],
        "annual": [{"label": label, "column": first_annual + index}
                   for index, label in enumerate(annual)],
        "annual_support_gutter_before": first_support_gutter,
        "annual_support": [{"label": label, "column": first_support + index}
                           for index, label in enumerate(support)],
        "annual_support_gutter_after": second_support_gutter,
        "quarterly": [{"label": label, "column": first_quarter + index}
                      for index, label in enumerate(quarterly)],
        "hidden_columns": sorted(
            first_annual + index for index, label in enumerate(annual) if label in hidden
        ) + sorted(
            first_quarter + index for index, label in enumerate(quarterly) if label in hidden
        ),
    }
    plan = {
        "schema_version": PLAN_SCHEMA_VERSION,
        "template_ref": "fund-xlsx-template-style:amzn-source:0.1",
        "style_resource_sha256": STYLE_RESOURCE_SHA256,
        "sheet_names": dict(sheet_names),
        "unit_labels": dict(unit_labels),
        "columns": columns,
        "row_styles": _row_plan(row_styles, sheet_names),
        "cell_styles": _cell_plan(cell_styles, sheet_names),
    }
    plan["content_hash"] = _canonical_hash(plan)
    return plan


def _font(Font: Any, value: Mapping[str, Any]) -> Any:
    return Font(name=value["name"], size=value["size"],
                bold=value.get("bold", False), italic=value.get("italic", False),
                color=value.get("color"))


def _apply_style(target: Any, style: Mapping[str, Any], *,
                 Font: Any, PatternFill: Any, Border: Any, Side: Any,
                 Alignment: Any) -> None:
    font = copy(target.font)
    for name, value in style["font"].items():
        setattr(font, name, value)
    target.font = font
    if "fill" in style:
        target.fill = PatternFill("solid", fgColor=style["fill"])
    if "horizontal" in style:
        target.alignment = Alignment(horizontal=style["horizontal"])
    if "border" in style:
        target.border = Border(**{
            edge: Side(style=value["style"], color=value["color"])
            for edge, value in style["border"].items()
        })


def apply_fund_xlsx_template(workbook: Any, plan: Mapping[str, Any]) -> None:
    """Apply reviewed styles/headers without changing financial values or formulas."""

    _need(isinstance(plan, Mapping) and set(plan) == {
        "schema_version", "template_ref", "style_resource_sha256", "sheet_names",
        "unit_labels", "columns", "row_styles", "cell_styles", "content_hash",
    } and plan["schema_version"] == PLAN_SCHEMA_VERSION
          and plan["template_ref"] == "fund-xlsx-template-style:amzn-source:0.1"
          and plan["style_resource_sha256"] == STYLE_RESOURCE_SHA256,
          "fund XLSX template application plan differs")
    unsigned = dict(plan)
    asserted_hash = unsigned.pop("content_hash")
    _need(asserted_hash == _canonical_hash(unsigned),
          "fund XLSX template application plan hash differs")
    style = load_fund_xlsx_template_style()
    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.cell import range_boundaries
    except ImportError as exc:
        raise FundXlsxTemplateError(
            "fund XLSX template application needs the optional `prior-models` dependency"
        ) from exc

    sheet_names = plan["sheet_names"]
    _need(isinstance(sheet_names, Mapping) and set(sheet_names) == set(_SHEET_ROLES),
          "fund XLSX template plan sheet binding differs")
    try:
        sheets = {role: workbook[name] for role, name in sheet_names.items()}
    except KeyError as exc:
        raise FundXlsxTemplateError("fund XLSX template sheet is missing") from exc

    desired = [sheets[role] for role in style["sheet_order"]]
    for index, sheet in enumerate(desired):
        workbook.move_sheet(sheet, offset=index - workbook.index(sheet))
    columns = plan["columns"]
    all_periods = columns["annual"] + columns["annual_support"] + columns["quarterly"]
    _need(all_periods and columns["annual"] and columns["quarterly"],
          "fund XLSX template plan needs annual and quarterly periods")
    last_column = max(item["column"] for item in all_periods)
    model_grid = style["model_grid"]
    default_font = model_grid["font"]
    for role in _MODEL_ROLES:
        sheet = sheets[role]
        sheet.sheet_view.showGridLines = model_grid["show_grid_lines"]
        sheet.freeze_panes = model_grid["freeze_cell"]
        sheet.sheet_format.defaultRowHeight = model_grid["default_row_height"]
        widths = model_grid["column_widths"]
        for column, width in zip(("A", "B", "C"),
                                 (widths["gutter_1"], widths["gutter_2"], widths["gutter_3"])):
            sheet.column_dimensions[column].width = width
        sheet.column_dimensions["D"].width = widths[f"{role}_label"]
        for item in all_periods:
            sheet.column_dimensions[get_column_letter(item["column"])].width = model_grid["period_column_width"]
        support = model_grid["annual_support"]
        before_gutter = get_column_letter(columns["annual_support_gutter_before"])
        after_gutter = get_column_letter(columns["annual_support_gutter_after"])
        sheet.column_dimensions[before_gutter].width = support[f"{role}_gutter_before_width"]
        sheet.column_dimensions[after_gutter].width = support[f"{role}_gutter_after_width"]
        for item in columns["annual_support"]:
            sheet.column_dimensions[get_column_letter(item["column"])].width = support[f"{role}_width"]
        for column in columns["hidden_columns"]:
            sheet.column_dimensions[get_column_letter(column)].hidden = True
        max_row = max(sheet.max_row, 1)
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=last_column):
            for cell in row:
                cell.font = _font(Font, default_font)
        sheet.cell(1, 1, plan["unit_labels"][role])
        for item in all_periods:
            sheet.cell(1, item["column"], item["label"])
        for cell in next(sheet.iter_rows(min_row=1, max_row=1,
                                         min_col=1, max_col=last_column)):
            header_style = "unit_header" if cell.column == 1 else "period_header"
            _apply_style(cell, style["styles"][header_style], Font=Font,
                         PatternFill=PatternFill, Border=Border, Side=Side,
                         Alignment=Alignment)
        for row in plan["row_styles"].get(role, []):
            sheet.row_dimensions[row["row"]].outlineLevel = row["level"]
            for cell in next(sheet.iter_rows(min_row=row["row"], max_row=row["row"],
                                             min_col=1, max_col=last_column)):
                _apply_style(cell, style["styles"][row["style"]], Font=Font,
                             PatternFill=PatternFill, Border=Border, Side=Side,
                             Alignment=Alignment)
        for item in plan["cell_styles"].get(role, []):
            min_col, min_row, max_col, max_row = range_boundaries(item["range"])
            _need(min_row >= 2 and max_col <= last_column,
                  "cell style range is outside the model grid")
            for row in sheet.iter_rows(min_row=min_row, max_row=max_row,
                                       min_col=min_col, max_col=max_col):
                for cell in row:
                    _apply_style(cell, style["styles"][item["style"]], Font=Font,
                                 PatternFill=PatternFill, Border=Border, Side=Side,
                                 Alignment=Alignment)
                    if item["number_kind"] is not None:
                        cell.number_format = style["number_formats"][item["number_kind"]]

    valuation = sheets["valuation"]
    valuation_grid = style["valuation_grid"]
    valuation.sheet_view.showGridLines = valuation_grid["show_grid_lines"]
    valuation.sheet_format.defaultRowHeight = valuation_grid["default_row_height"]
    for column, width in valuation_grid["column_widths"].items():
        valuation.column_dimensions[column].width = width
    for row in valuation.iter_rows():
        for cell in row:
            cell.font = _font(Font, valuation_grid["font"])
    valuation_last = max(valuation.max_column, 5)
    for row in plan["row_styles"].get("valuation", []):
        for cell in next(valuation.iter_rows(min_row=row["row"], max_row=row["row"],
                                             min_col=1, max_col=valuation_last)):
            _apply_style(cell, style["styles"][row["style"]], Font=Font,
                         PatternFill=PatternFill, Border=Border, Side=Side,
                         Alignment=Alignment)
    for item in plan["cell_styles"].get("valuation", []):
        min_col, min_row, max_col, max_row = range_boundaries(item["range"])
        for row in valuation.iter_rows(min_row=min_row, max_row=max_row,
                                       min_col=min_col, max_col=max_col):
            for cell in row:
                _apply_style(cell, style["styles"][item["style"]], Font=Font,
                             PatternFill=PatternFill, Border=Border, Side=Side,
                             Alignment=Alignment)
                if item["number_kind"] is not None:
                    cell.number_format = style["number_formats"][item["number_kind"]]


__all__ = [
    "FundXlsxTemplateError",
    "PLAN_SCHEMA_VERSION",
    "STYLE_SCHEMA_VERSION",
    "apply_fund_xlsx_template",
    "build_fund_xlsx_template_plan",
    "load_fund_xlsx_template_style",
]
