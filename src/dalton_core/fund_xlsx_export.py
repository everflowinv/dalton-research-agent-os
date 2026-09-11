"""Read-only export of one governed forecast into an auditable fund workbook."""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import sqlite3
import tempfile
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Mapping, Sequence

from .company_model_inputs import build_model_inputs
from .coverage_mission import CoverageMissionAuthority
from .fund_xlsx_template import (
    apply_fund_xlsx_template,
    build_fund_xlsx_template_plan,
)
from .model_forecast_driver import ForecastModelAuthority, validate_forecast_model
from .store import DaltonStore, canonical_json, content_hash
from .valuation_snapshot import ValuationSnapshotAuthority

LAYOUT_VERSION = "fund-xlsx-layout:0.4"
FUND_MONETARY_DISPLAY_SCALE = 1_000_000


class FundWorkbookExportError(RuntimeError):
    pass


def _xlsx() -> tuple[Any, Any, Any, Any]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        return Workbook, load_workbook, Font, PatternFill
    except ImportError as exc:
        raise FundWorkbookExportError(
            "XLSX export needs the optional `prior-models` dependency") from exc


def _number(value: Any) -> float:
    number = Decimal(str(value))
    if not number.is_finite():
        raise FundWorkbookExportError("workbook number must be finite")
    converted = float(number)
    if not math.isfinite(converted):
        raise FundWorkbookExportError("workbook number exceeds Excel numeric range")
    if number and abs((Decimal(str(converted)) - number) / number) > Decimal("1e-15"):
        raise FundWorkbookExportError("workbook number exceeds Excel's reliable precision")
    return converted


def _col(index: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(index)


def _style_sheet(ws: Any, font: Any, fill: Any) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 37
    for cell in ws[1]:
        cell.font = font(name="Arial", size=15, bold=True, color="FFFFFF")
        cell.fill = fill("solid", fgColor="17365D")
    for cell in ws[4]:
        cell.font = font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = fill("solid", fgColor="1F4E78")


def _quarter_label(
    period: str,
    binding: Mapping[str, Any] | None,
    historical: set[str],
) -> str:
    suffix = "A" if period in historical else "E"
    if binding is None:
        return f"{period}{suffix}"
    year = int(period[:4])
    month = int(period[5:7])
    fiscal_end = binding["fiscal_year_end_month"]
    fiscal_year = year if month <= fiscal_end else year + 1
    months_to_end = (fiscal_end - month) % 12
    quarter = 4 - (months_to_end // 3)
    if months_to_end % 3:
        return f"{period}{suffix}"
    return f"Q{quarter} FY{fiscal_year}{suffix}"


def _number_format(unit: str, *, per_share: bool = False) -> str:
    normalized = str(unit).casefold()
    if normalized == "ratio":
        return "0.0%;(0.0%);-"
    if per_share or normalized.endswith("_per_share"):
        return '"$"#,##0.00;[Red]("$"#,##0.00);-'
    if normalized == "usd":
        return '"$"#,##0.0,,;[Red]("$"#,##0.0,,);-'
    if re.fullmatch(r"[a-z]{3}", normalized):
        return "#,##0.0,,;[Red](#,##0.0,,);-"
    return "#,##0;[Red](#,##0);-"


def _display_unit(unit: str) -> str:
    normalized = str(unit).casefold()
    if re.fullmatch(r"[a-z]{3}", normalized):
        return f"{normalized.upper()} millions"
    matched = re.fullmatch(r"([a-z]{3})_per_share", normalized)
    if matched:
        return f"{matched.group(1).upper()} per share"
    return unit


def _template_number_kind(unit: str, *, per_share: bool = False) -> str:
    normalized = str(unit).casefold()
    if normalized == "ratio":
        return "percentage"
    if per_share or normalized.endswith("_per_share"):
        return "per_share"
    return "amount_one_decimal" if normalized == "usd" else "amount"


def _template_value(
    value: Any, unit: str, *, role: str | None = None,
) -> float:
    """Translate raw authority units into the source workbook display units."""

    number = _number(value)
    normalized = str(unit).casefold()
    if re.fullmatch(r"[a-z]{3}", normalized):
        return number / FUND_MONETARY_DISPLAY_SCALE
    if role == "diluted_weighted_average_shares":
        return number / FUND_MONETARY_DISPLAY_SCALE
    return number


def _financial_row_presentation(result: Mapping[str, Any]) -> tuple[int, str]:
    role = str(result.get("role") or "")
    subtotal_roles = {
        "revenue", "gross_profit", "operating_income", "pretax_income",
        "income_from_continuing_operations", "net_income", "parent_net_income",
        "diluted_eps", "free_cash_flow", "company_presented_subtotal",
    }
    if role in subtotal_roles or str(result.get("ref")) in {
        "result:gross_profit", "result:operating_income", "result:net_income",
        "result:free_cash_flow",
    }:
        return 1, "subtotal"
    return 2, "label"


def _financial_display_unit(result: Mapping[str, Any]) -> str:
    if result.get("role") == "diluted_weighted_average_shares":
        return "shares millions"
    return _display_unit(str(result.get("unit") or "number"))


_FINANCIAL_LINE_LABELS = {
    "result:revenue": "Revenue",
    "result:cost_of_revenue": "Cost of revenue",
    "result:gross_profit": "Gross profit",
    "result:operating_income": "Operating income",
    "result:income_tax_expense": "Income tax expense",
    "result:net_income": "Net income",
    "result:operating_cash_flow": "Operating cash flow",
    "result:capital_expenditure": "Capital expenditures",
    "result:free_cash_flow": "Free cash flow",
}

_READER_LABELS = {
    "CostOfGoodsAndServicesSold": "Cost of goods and services sold",
    "IncomeTaxExpenseBenefit": "Income tax expense",
    "SellingGeneralAndAdministrativeExpense":
        "Selling, general & administrative",
    "ResearchAndDevelopmentExpense": "Research & development",
}


def _reader_label(value: Any) -> str:
    label = str(value or "Financial line").strip()
    if label in _READER_LABELS:
        return _READER_LABELS[label]
    humanized = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", label)
    humanized = humanized.replace("_", " ").strip()
    return humanized[:1].upper() + humanized[1:] if humanized else "Financial line"


def _financial_line_label(result: Mapping[str, Any]) -> str:
    ref = str(result["ref"])
    if ref in _FINANCIAL_LINE_LABELS:
        return _FINANCIAL_LINE_LABELS[ref]
    if ref.startswith("result:operating_expense:"):
        concept = ref.rsplit(":", 1)[-1]
        if concept in _READER_LABELS:
            return _READER_LABELS[concept]
    return _reader_label(result.get("label"))


def _is_operating_expense_share_formula(
    result: Mapping[str, Any], formula: str,
) -> bool:
    prefix = "result:operating_expense:"
    ref = str(result.get("ref") or "")
    if not ref.startswith(prefix):
        return False
    # The authority formula is frozen against the XBRL concept while the
    # result label is intentionally reader-facing and may be humanized or
    # localized. Derive only the exact local concept from the closed result
    # ref; never use display text to decide executable formula semantics.
    concept = ref[len(prefix):].rsplit(":", 1)[-1]
    return bool(concept) and formula == f"{concept}[k] = revenue[k] * share[k]"


def _formula_for(
    result: Mapping[str, Any], cell: Mapping[str, Any], period: str,
    result_cells: Mapping[tuple[str, str], str], assumption_cells: Mapping[str, str],
    history_cells: Mapping[tuple[str, str], str],
) -> str | None:
    refs = [result_cells.get((r["ref"], r["period_end"]))
            for r in cell.get("result_refs") or []]
    if any(ref is None for ref in refs):
        return None
    assumptions = [assumption_cells.get(ref) for ref in cell.get("assumption_refs") or []]
    if any(ref is None for ref in assumptions):
        return None
    formula = result["formula"]
    if ((formula == "revenue[k] = revenue[k-1] * (1 + growth[k])"
         or formula.endswith("[k-1] * (1 + growth[k])"))
            and len(refs) <= 1 and len(assumptions) == 1):
        prior_refs = list(cell.get("result_refs") or [])
        prior = refs[0] if refs else None
        if prior is None and cell.get("input_cell_refs"):
            source = cell["input_cell_refs"][0]
            prior = history_cells.get((source.get("concept"), source.get("period_end")))
        return None if prior is None else f"={prior}*(1+{assumptions[0]})"
    if (formula in {
            "cost_of_revenue[k] = revenue[k] * share[k]",
            "income_tax_expense[k] = operating_income[k] * share[k]",
            "operating_cash_flow[k] = revenue[k] * share[k]",
            "capital_expenditure[k] = revenue[k] * share[k]",
            } or _is_operating_expense_share_formula(result, formula)
              or formula.endswith("[k] * share[k]")) \
            and len(refs) == 1 and len(assumptions) == 1:
        return f"={refs[0]}*{assumptions[0]}"
    if (formula in {
            "gross_profit[k] = revenue[k] - cost_of_revenue[k]",
            "operating_income[k] = gross_profit[k] - sum(operating_expense[k])",
            "net_income[k] = operating_income[k] - income_tax[k]",
            "free_cash_flow[k] = operating_cash_flow[k] - capital_expenditure[k]",
            } and len(refs) >= 2 and not assumptions):
        return f"={refs[0]}-SUM({','.join(refs[1:])})"
    if not assumptions:
        try:
            contract = json.loads(formula)
        except (TypeError, ValueError, json.JSONDecodeError):
            contract = None
        if isinstance(contract, Mapping) and contract.get("operator") == "sum":
            terms = contract.get("terms") or []
            if len(terms) == len(refs):
                expression = "".join(
                    (("+" if index and term.get("coefficient") == "1" else "")
                     + ("-" if term.get("coefficient") == "-1" else "") + refs[index])
                    for index, term in enumerate(terms))
                return "=" + expression
        if (isinstance(contract, Mapping) and contract.get("operator") == "divide"
                and len(refs) == 2):
            return f"={refs[0]}/{refs[1]}"
    return None


def _verify_translated_value(
    result: Mapping[str, Any], cell: Mapping[str, Any],
    result_values: Mapping[tuple[str, str], Decimal],
    assumption_values: Mapping[str, Decimal],
    history_values: Mapping[tuple[str, str], Decimal],
) -> None:
    refs = [result_values.get((item["ref"], item["period_end"]))
            for item in cell.get("result_refs") or []]
    assumptions = [assumption_values.get(item)
                   for item in cell.get("assumption_refs") or []]
    formula = result["formula"]
    expected: Decimal | None = None
    if (formula == "revenue[k] = revenue[k-1] * (1 + growth[k])"
            or formula.endswith("[k-1] * (1 + growth[k])")) and assumptions:
        prior = refs[0] if refs else None
        if prior is None and cell.get("input_cell_refs"):
            source = cell["input_cell_refs"][0]
            prior = history_values.get((source.get("concept"), source.get("period_end")))
        if prior is not None:
            expected = prior * (Decimal(1) + assumptions[0])
    if expected is None and (formula in {
            "cost_of_revenue[k] = revenue[k] * share[k]",
            "income_tax_expense[k] = operating_income[k] * share[k]",
            "operating_cash_flow[k] = revenue[k] * share[k]",
            "capital_expenditure[k] = revenue[k] * share[k]",
            } or _is_operating_expense_share_formula(result, formula)
              or formula.endswith("[k] * share[k]")):
        if refs and assumptions:
            expected = refs[0] * assumptions[0]
    elif formula in {
            "gross_profit[k] = revenue[k] - cost_of_revenue[k]",
            "operating_income[k] = gross_profit[k] - sum(operating_expense[k])",
            "net_income[k] = operating_income[k] - income_tax[k]",
            "free_cash_flow[k] = operating_cash_flow[k] - capital_expenditure[k]",
            } and refs and all(item is not None for item in refs):
        expected = refs[0] - sum(refs[1:], Decimal(0))
    if expected is None and not assumptions:
        try:
            contract = json.loads(formula)
        except (TypeError, ValueError, json.JSONDecodeError):
            contract = None
        if (isinstance(contract, Mapping) and contract.get("operator") == "sum"
                and len(contract.get("terms") or []) == len(refs)
                and all(item is not None for item in refs)):
            expected = sum(
                (Decimal(str(term["coefficient"])) * refs[index]
                 for index, term in enumerate(contract["terms"])), Decimal(0))
        elif (isinstance(contract, Mapping) and contract.get("operator") == "divide"
              and len(refs) == 2 and all(item is not None for item in refs)
              and refs[1] != 0):
            expected = refs[0] / refs[1]
    if expected is None or "value" not in cell:
        return
    actual = Decimal(str(cell["value"]))
    tolerance = max(Decimal("0.000001"), abs(actual) * Decimal("1e-12"))
    if abs(expected - actual) > tolerance:
        raise FundWorkbookExportError(
            f"translated formula disagrees with {cell['ref']}: {expected} != {actual}")


def _calendar_binding(value: Mapping[str, Any] | None) -> dict[str, Any] | None:
    if value is None:
        return None
    binding = dict(value)
    allowed = {"calendar_ref", "source_hash", "as_of",
               "fiscal_year_end_month", "content_hash"}
    if set(binding) != allowed:
        raise FundWorkbookExportError("fiscal calendar binding has an invalid closed shape")
    asserted = binding.pop("content_hash")
    if asserted != content_hash(binding):
        raise FundWorkbookExportError("fiscal calendar binding content hash is invalid")
    month = binding["fiscal_year_end_month"]
    if isinstance(month, bool) or not isinstance(month, int) or not 1 <= month <= 12:
        raise FundWorkbookExportError("fiscal_year_end_month must be an integer from 1 to 12")
    return {**binding, "content_hash": asserted}


def _fiscal_groups(periods: Sequence[str], binding: Mapping[str, Any] | None,
                   historical: set[str]) -> list[tuple[str, list[str]]]:
    if binding is None:
        return []
    groups: list[tuple[str, list[str]]] = []
    pending: list[str] = []
    for period in periods:
        pending.append(period)
        if int(period[5:7]) == binding["fiscal_year_end_month"]:
            kinds = {"A" if item in historical else "E" for item in pending}
            suffix = next(iter(kinds)) if len(kinds) == 1 else "A/E"
            groups.append((f"FY{period[:4]}{suffix}", pending))
            pending = []
    if pending:
        kinds = {"A" if item in historical else "E" for item in pending}
        suffix = next(iter(kinds)) if len(kinds) == 1 else "A/E"
        groups.append((f"FY{pending[-1][:4]}{suffix} (partial)", pending))
    return groups


def _four_quarter_flow_cells(
    by_period: Mapping[str, Mapping[str, Any]], group: Sequence[str]
) -> bool:
    """Prove that a fiscal-year aggregate contains four duration quarters."""

    if len(group) != 4:
        return False
    prior_end: date | None = None
    for period_end in group:
        cell = by_period.get(period_end)
        period = cell.get("period") if isinstance(cell, Mapping) else None
        if not isinstance(period, Mapping):
            return False
        if (
            period.get("kind") != "quarter"
            or period.get("end") != period_end
            or not isinstance(period.get("start"), str)
        ):
            return False
        try:
            start = date.fromisoformat(period["start"])
            end = date.fromisoformat(period_end)
        except ValueError:
            return False
        if start > end or (prior_end is not None and start <= prior_end):
            return False
        prior_end = end
    return True


def _duration_quarter_cell(cell: Mapping[str, Any], period_end: str) -> bool:
    period = cell.get("period")
    if not isinstance(period, Mapping):
        return False
    if (
        period.get("kind") != "quarter"
        or period.get("end") != period_end
        or not isinstance(period.get("start"), str)
    ):
        return False
    try:
        return date.fromisoformat(period["start"]) <= date.fromisoformat(period_end)
    except ValueError:
        return False


def _inclusive_duration_days(cell: Mapping[str, Any]) -> int:
    try:
        return (
            date.fromisoformat(str(cell.get("period_end")))
            - date.fromisoformat(str(cell.get("period_start")))
        ).days + 1
    except ValueError:
        return 0


def _annual_line_definition_ref(
    structure: Mapping[str, Any], line: Mapping[str, Any], concept: Any,
) -> str:
    definition = {
        "structure_ref": structure.get("structure_ref"),
        "structure_hash": structure.get("content_hash"),
        "line_ref": line.get("ref"), "role": line.get("role"),
        "concept": concept, "unit": line.get("unit"),
        "annual_semantics": line.get("annual_semantics"),
    }
    return f"statement-line-definition:{content_hash(definition)}"


def _annual_structure_facts(
    inputs: Mapping[str, Any], structure: Mapping[str, Any],
    line: Mapping[str, Any], group: Sequence[str], label: str,
    calendar_binding: Mapping[str, Any],
    formula: Mapping[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Bind raw typed durations to one exact fiscal group and structure line."""

    concept = line.get("concept") if line.get("kind") == "filed" else None
    if concept is None and formula is not None:
        concept = formula.get("tie_out_concept")
    source = next(
        (item for item in (inputs.get("filed_lines") or [])
         if item.get("concept") == concept),
        None,
    )
    if not isinstance(source, Mapping) or len(group) != 4:
        return []
    definition_ref = _annual_line_definition_ref(structure, line, concept)
    fiscal_year = f"{calendar_binding['calendar_ref']}:{label}"
    calendar = calendar_binding["content_hash"]
    selected: list[dict[str, Any]] = []
    for raw in source.get("duration_facts") or []:
        if not isinstance(raw, Mapping):
            continue
        try:
            start = date.fromisoformat(str(raw.get("period_start")))
            end = date.fromisoformat(str(raw.get("period_end")))
        except ValueError:
            continue
        unit = str(raw.get("unit") or "").casefold()
        if unit != str(line.get("unit") or "").casefold():
            continue
        period_kind = raw.get("period_kind")
        if period_kind == "quarter" and str(raw.get("period_end")) in group:
            normalized_kind = "quarter"
        elif (
            str(raw.get("period_end")) == group[-1]
            and 290 < (end - start).days <= 380
        ):
            normalized_kind = "annual"
        else:
            continue
        selected.append({
            "fiscal_year": fiscal_year, "period_kind": normalized_kind,
            "period_start": start.isoformat(), "period_end": end.isoformat(),
            "value": str(raw.get("value")), "unit": unit,
            "calendar": calendar, "definition_ref": definition_ref,
            "source_accessions": list(raw.get("source_accessions") or []),
            "source_forms": list(raw.get("source_forms") or []),
        })
    return selected


def _structured_annual_eps(
    inputs: Mapping[str, Any], structure: Mapping[str, Any],
    annual_groups: Sequence[tuple[str, list[str]]],
    calendar_binding: Mapping[str, Any] | None,
) -> dict[str, dict[str, Any]]:
    """Compute only historically disclosed, fiscal-authority-bound annual EPS."""

    if calendar_binding is None:
        return {}
    from .company_financial_statement_structure import (
        aggregate_fiscal_year, annual_diluted_eps,
    )

    lines = {str(item["ref"]): item for item in (structure.get("lines") or [])}
    formulas = {str(item["output_ref"]): item
                for item in (structure.get("formulas") or [])}
    by_role = {str(item["role"]): item for item in lines.values()}
    numerator = by_role.get("diluted_eps_numerator")
    shares = by_role.get("diluted_weighted_average_shares")
    eps = by_role.get("diluted_eps")
    if numerator is None or shares is None or eps is None:
        return {}
    results: dict[str, dict[str, Any]] = {}
    for label, group in annual_groups:
        if label.endswith("(partial)") or not label.endswith("A"):
            continue
        fiscal_year = f"{calendar_binding['calendar_ref']}:{label}"
        numerator_cells = _annual_structure_facts(
            inputs, structure, numerator, group, label, calendar_binding,
            formulas.get(str(numerator["ref"])),
        )
        share_cells = _annual_structure_facts(
            inputs, structure, shares, group, label, calendar_binding,
            formulas.get(str(shares["ref"])),
        )
        eps_cells = _annual_structure_facts(
            inputs, structure, eps, group, label, calendar_binding,
            formulas.get(str(eps["ref"])),
        )
        outcome = annual_diluted_eps(
            diluted_eps_numerator_cells=numerator_cells,
            diluted_weighted_share_cells=share_cells,
            diluted_eps_cells=eps_cells,
            fiscal_year=fiscal_year,
        )
        results[label] = {
            **outcome,
            "direct_numerator": aggregate_fiscal_year(
                numerator_cells, semantic="direct_annual", fiscal_year=fiscal_year),
            "direct_shares": aggregate_fiscal_year(
                share_cells, semantic="direct_annual", fiscal_year=fiscal_year),
        }
    return results


def _structured_annual_share_forecasts(
    model: Mapping[str, Any], annual_groups: Sequence[tuple[str, list[str]]],
    calendar_binding: Mapping[str, Any] | None,
) -> dict[str, dict[str, Any]]:
    """Execute the structure's separately authorized annual share method."""

    if calendar_binding is None:
        return {}
    from .company_financial_statement_structure import day_weighted_annual_shares

    structure = model.get("financial_statement_structure") or {}
    share_line = next(
        (item for item in (structure.get("lines") or [])
         if item.get("role") == "diluted_weighted_average_shares"), None,
    )
    if (
        not isinstance(share_line, Mapping)
        or share_line.get("annual_forecast_method") != "day_weighted_quarters"
    ):
        return {}
    replay = model.get("financial_statement_structure_replay") or {}
    share_replay = next(
        (item for item in (replay.get("forecast_methods") or [])
         if item.get("line_ref") == share_line.get("ref")), None,
    )
    if not isinstance(share_replay, Mapping) or share_replay.get("annual_status") != "validated":
        return {}
    share_result = next(
        (item for item in (model.get("results") or [])
         if item.get("role") == "diluted_weighted_average_shares"), None,
    )
    share_driver = next(
        (item for item in (model.get("drivers") or [])
         if item.get("structure_line_ref") == share_line.get("ref")), None,
    )
    if not isinstance(share_result, Mapping) or not isinstance(share_driver, Mapping):
        return {}
    by_end: dict[str, dict[str, Any]] = {
        str(item.get("period_end")): {
            "period_start": item.get("period_start"),
            "period_end": item.get("period_end"), "value": item.get("value"),
            "unit": share_line.get("unit"),
        }
        for item in (share_driver.get("history") or [])
        if item.get("period_start") and item.get("period_end")
    }
    for cell in share_result.get("cells") or []:
        period = cell.get("period")
        if (
            cell.get("status") == "computed" and isinstance(period, Mapping)
            and period.get("kind") == "quarter"
        ):
            by_end[str(period.get("end"))] = {
                "period_start": period.get("start"), "period_end": period.get("end"),
                "value": cell.get("value"), "unit": share_line.get("unit"),
            }
    calendar = calendar_binding["content_hash"]
    definition_ref = _annual_line_definition_ref(
        structure, share_line, share_line.get("concept"),
    )
    results: dict[str, dict[str, Any]] = {}
    for label, group in annual_groups:
        if label.endswith("A") or label.endswith("(partial)") or len(group) != 4:
            continue
        cells = []
        for end in group:
            cell = by_end.get(end)
            if cell is None:
                continue
            cells.append({
                **cell, "calendar": calendar, "definition_ref": definition_ref,
            })
        results[label] = day_weighted_annual_shares(
            cells, required_calendar=calendar,
            required_definition_ref=definition_ref,
        )
    return results


def export_fund_workbook(
    output: Path, *, model: Mapping[str, Any], spec: Mapping[str, Any],
    inputs: Mapping[str, Any], valuation: Mapping[str, Any] | None = None,
    valuation_scenario: Mapping[str, Any] | None = None,
    calendar_binding: Mapping[str, Any] | None = None,
    mission_binding: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Export verified authority records without changing any authority."""
    model_wire = dict(model)
    model_wire.pop("status", None)
    model = validate_forecast_model(model_wire)
    if model["spec_ref"] != spec.get("spec_id") or model["spec_hash"] != spec.get("content_hash"):
        raise FundWorkbookExportError("forecast model does not bind this specification")
    if model["inputs_hash"] != content_hash(json.loads(canonical_json(inputs))):
        raise FundWorkbookExportError("forecast model does not bind these model inputs")
    if inputs.get("company_ref") != model["company_ref"]:
        raise FundWorkbookExportError("model inputs belong to another company")
    Workbook, _, Font, PatternFill = _xlsx()
    wb = Workbook()
    wb.remove(wb.active)
    valuation_ws = wb.create_sheet("Valuation")
    financials = wb.create_sheet("Financials")
    driver = wb.create_sheet("Driver")
    sources = wb.create_sheet("Sources")
    manifest = wb.create_sheet("Formula Map")
    periods = list(dict.fromkeys([*model["history_periods"], *[p["end"] for p in model["realised_periods"] + model["forecast_periods"]]]))
    if not periods:
        periods = sorted({c["period"]["end"] for r in model["results"] for c in r["cells"]})
    history = list(model["history_periods"])
    calendar_binding = _calendar_binding(calendar_binding)
    annual_groups = _fiscal_groups(periods, calendar_binding, set(history))
    annual_labels = [label for label, _ in annual_groups]
    displayed_annual_labels = annual_labels or ["Annual unavailable"]
    quarter_labels = [
        _quarter_label(period, calendar_binding, set(history)) for period in periods
    ]
    template_plan = build_fund_xlsx_template_plan(
        sheet_names={"valuation": "Valuation", "financials": "Financials", "driver": "Driver"},
        annual_periods=displayed_annual_labels, quarterly_periods=quarter_labels,
        unit_labels={"financials": "(USD MM)", "driver": "(US$mm)"},
    )
    annual_columns = {
        item["label"]: item["column"] for item in template_plan["columns"]["annual"]
        if item["label"] in annual_labels
    }
    quarter_start = template_plan["columns"]["quarterly"][0]["column"]
    period_columns = {
        period: quarter_start + index for index, period in enumerate(periods)
    }
    cagr_column = template_plan["columns"]["annual_support"][0]["column"]
    formula_map: list[dict[str, Any]] = []
    statement_structure = (
        model.get("financial_statement_structure") or {}
        if model.get("schema_version") == "0.3" else {}
    )
    driver_roles = {
        str(line["ref"]): str(line.get("role") or "")
        for line in statement_structure.get("lines") or []
    }

    for ws, title in ((valuation_ws, "Valuation"), (sources, "Sources and bindings"),
                      (manifest, "Formula map")):
        ws.append([title])
        company_label = (
            (mission_binding or {}).get("entity_name")
            or (mission_binding or {}).get("ticker")
            or model["company_ref"]
        )
        ws.append([company_label, f"As of {model['created_at']}"])
        ws.append([])

    driver.cell(2, 1, "Key drivers")
    template_row_styles: dict[str, list[dict[str, Any]]] = {
        "driver": [{"row": 2, "style": "major_section", "level": 0}],
        "financials": [{"row": 2, "style": "major_section", "level": 0}],
        "valuation": [],
    }
    template_cell_styles: dict[str, list[dict[str, Any]]] = {
        "driver": [], "financials": [], "valuation": [],
    }
    history_cells: dict[tuple[str, str], str] = {}
    history_values: dict[tuple[str, str], Decimal] = {}
    history_flow_periods: set[tuple[str, str]] = set()
    row = 3
    for item in model["drivers"]:
        values = {c["period_end"]: c for c in item.get("history") or []}
        item_role = driver_roles.get(str(item.get("structure_line_ref") or ""))
        unit_label = (
            "shares millions" if item_role == "diluted_weighted_average_shares"
            else _display_unit(item.get("unit") or "number")
        )
        driver.cell(row, 2, f"{_reader_label(item['label'])} ({unit_label}) — actual")
        template_row_styles["driver"].append({"row": row, "style": "label", "level": 1})
        for period in periods:
            ci = period_columns[period]
            value = values.get(period)
            if value:
                driver.cell(
                    row, ci,
                    _template_value(
                        value["value"], item.get("unit") or "number", role=item_role,
                    ),
                )
                template_cell_styles["driver"].append({
                    "range": driver.cell(row, ci).coordinate, "style": "hardcoded_input",
                    "number_kind": _template_number_kind(item.get("unit") or "number"),
                })
                history_cells[(item.get("concept"), period)] = f"'Driver'!{_col(ci)}{row}"
                history_cells[(item["ref"], period)] = f"'Driver'!{_col(ci)}{row}"
                history_values[(item.get("concept"), period)] = Decimal(str(value["value"]))
                history_values[(item["ref"], period)] = Decimal(str(value["value"]))
                if isinstance(value.get("period_start"), str):
                    history_flow_periods.add((item["ref"], period))
        row += 1
    driver.cell(row, 2, "Forecast assumptions")
    template_row_styles["driver"].append({"row": row, "style": "section", "level": 1})
    row += 1
    assumption_cells: dict[str, str] = {}
    assumption_values: dict[str, Decimal] = {}
    driver_labels = {
        item["ref"]: _reader_label(item["label"]) for item in model["drivers"]
    }
    for assumption in model["assumptions"]:
        label = driver_labels.get(assumption["driver_ref"], assumption["driver_ref"])
        driver.cell(
            row,
            3,
            f"{label} ({_display_unit(assumption['unit'])}) — "
            f"{assumption['measure']} — assumption",
        )
        period = assumption["period"]["end"]
        if period in periods:
            ci = period_columns[period]
            driver.cell(row, ci, _number(assumption["value"]))
            template_cell_styles["driver"].append({
                "range": driver.cell(row, ci).coordinate, "style": "assumption_input",
                "number_kind": _template_number_kind(assumption["unit"]),
            })
            assumption_cells[assumption["ref"]] = f"'Driver'!{_col(ci)}{row}"
            assumption_values[assumption["ref"]] = Decimal(str(assumption["value"]))
        template_row_styles["driver"].append({
            "row": row, "style": "growth_label", "level": 2,
        })
        row += 1

    financials.cell(2, 1, "Income statement")
    result_rows = {item["ref"]: 3 + i for i, item in enumerate(model["results"])}
    structured_lines: dict[str, Mapping[str, Any]] = {}
    structured_formulas: dict[str, Mapping[str, Any]] = {}
    if model.get("schema_version") == "0.3":
        from .model_forecast_driver import _structure_result_refs
        structure = statement_structure
        refs_by_line = _structure_result_refs(structure)
        structured_lines = {
            refs_by_line[str(line["ref"])]: line for line in structure.get("lines") or []
        }
        structured_formulas = {
            refs_by_line[str(formula["output_ref"])]: formula
            for formula in structure.get("formulas") or []
        }
    annual_eps_results = _structured_annual_eps(
        inputs, model.get("financial_statement_structure") or {},
        annual_groups, calendar_binding,
    ) if model.get("schema_version") == "0.3" else {}
    annual_share_forecasts = _structured_annual_share_forecasts(
        model, annual_groups, calendar_binding,
    ) if model.get("schema_version") == "0.3" else {}
    result_cells: dict[tuple[str, str], str] = {}
    result_flow_periods: set[tuple[str, str]] = set()
    result_values = {
        (result["ref"], cell["period"]["end"]): Decimal(str(cell["value"]))
        for result in model["results"] for cell in result["cells"]
        if cell.get("status") == "computed" and "value" in cell
    }
    gaps: list[str] = []
    if calendar_binding is None:
        gaps.append("annual columns unavailable: no bound fiscal calendar")
    for result in model["results"]:
        rr = result_rows[result["ref"]]
        level, row_style = _financial_row_presentation(result)
        financials.cell(
            rr,
            level + 1,
            f"{_financial_line_label(result)} ({_financial_display_unit(result)})"
            + (" — Not available" if result["status"] != "computed" else ""),
        )
        template_row_styles["financials"].append({
            "row": rr, "style": row_style, "level": level,
        })
        by_period = {c["period"]["end"]: c for c in result["cells"]}
        for period in periods:
            ci = period_columns[period]
            cell = by_period.get(period)
            formula = None
            model_cell_ref = None
            if period in history and result.get("driver_ref"):
                source = history_cells.get((result["driver_ref"], period))
                if source:
                    formula = f"={source}"
                    model_cell_ref = result["driver_ref"]
                    if (result["driver_ref"], period) in history_flow_periods:
                        result_flow_periods.add((result["ref"], period))
            elif period in history:
                components: list[str] = []
                structured_formula = structured_formulas.get(result["ref"])
                component_refs: list[str] = []
                if structured_formula is not None:
                    if structured_formula["operator"] == "sum":
                        line_refs = [item["line_ref"] for item in structured_formula["terms"]]
                        from .model_forecast_driver import _structure_result_refs
                        refs_by_line = _structure_result_refs(
                            model["financial_statement_structure"])
                        component_refs = [refs_by_line[item] for item in line_refs]
                        components = [result_cells.get((ref, period))
                                      for ref in component_refs]
                        if components and all(components):
                            pieces = []
                            for index, (term, coordinate) in enumerate(zip(
                                    structured_formula["terms"], components)):
                                prefix = ("-" if term["coefficient"] == "-1"
                                          else "+" if index else "")
                                pieces.append(prefix + coordinate)
                            formula = "=" + "".join(pieces)
                    else:
                        from .model_forecast_driver import _structure_result_refs
                        refs_by_line = _structure_result_refs(
                            model["financial_statement_structure"])
                        component_refs = [
                            refs_by_line[structured_formula["numerator_ref"]],
                            refs_by_line[structured_formula["denominator_ref"]],
                        ]
                        components = [result_cells.get((ref, period))
                                      for ref in component_refs]
                        if all(components):
                            formula = f"={components[0]}/{components[1]}"
                    if formula is not None:
                        model_cell_ref = "historical-structured-derived"
                        if all((ref, period) in result_flow_periods
                               for ref in component_refs):
                            result_flow_periods.add((result["ref"], period))
                elif result["ref"] == "result:gross_profit":
                    components = [result_cells.get(("result:revenue", period)),
                                  result_cells.get(("result:cost_of_revenue", period))]
                elif result["ref"] == "result:operating_income":
                    expense_refs = [item["ref"] for item in model["results"]
                                    if item["ref"].startswith("result:operating_expense:")]
                    components = [result_cells.get(("result:gross_profit", period)),
                                  *[result_cells.get((ref, period)) for ref in expense_refs]]
                elif result["ref"] == "result:net_income":
                    components = [result_cells.get(("result:operating_income", period)),
                                  result_cells.get(("result:income_tax_expense", period))]
                if structured_formula is None and components and all(components):
                    formula = f"={components[0]}-SUM({','.join(components[1:])})"
                    model_cell_ref = "historical-derived"
                    component_refs = (
                        ["result:revenue", "result:cost_of_revenue"]
                        if result["ref"] == "result:gross_profit"
                        else ["result:gross_profit", *expense_refs]
                        if result["ref"] == "result:operating_income"
                        else ["result:operating_income", "result:income_tax_expense"]
                    )
                    if all((ref, period) in result_flow_periods
                           for ref in component_refs):
                        result_flow_periods.add((result["ref"], period))
            elif cell is not None and cell["status"] == "computed":
                formula = _formula_for(result, cell, period, result_cells,
                                       assumption_cells, history_cells)
                model_cell_ref = cell["ref"]
                if formula is not None:
                    _verify_translated_value(
                        result, cell, result_values, assumption_values,
                        history_values)
                    if _duration_quarter_cell(cell, period):
                        result_flow_periods.add((result["ref"], period))
            if formula is None:
                if cell is not None:
                    reason = cell.get("reason") if cell["status"] != "computed" else result["formula"]
                    gaps.append(f"{result['ref']} {period}: unavailable or unsupported; {reason}")
                continue
            financials.cell(rr, ci, formula)
            result_cells[(result["ref"], period)] = f"'Financials'!{_col(ci)}{rr}"
            template_cell_styles["financials"].append({
                "range": financials.cell(rr, ci).coordinate,
                "style": "cross_sheet_formula" if "'Driver'!" in formula else "local_formula",
                "number_kind": _template_number_kind(result["unit"]),
            })
            formula_map.append({"cell": f"Financials!{_col(ci)}{rr}",
                                "model_cell_ref": model_cell_ref, "formula": formula,
                                "model_formula": result["formula"],
                                "model_label": result["label"]})
        for label, group in annual_groups:
            annual_i = annual_columns[label]
            quarter_cols = [period_columns[p] for p in group
                            if (result["ref"], p) in result_cells]
            target = financials.cell(rr, annual_i)
            flow_proven = all(
                (result["ref"], period) in result_flow_periods for period in group
            )
            structured_line = structured_lines.get(result["ref"])
            structured_role = (
                None if structured_line is None else structured_line.get("role")
            )
            annual_eps = annual_eps_results.get(label)
            if (
                annual_eps is not None and annual_eps.get("status") == "computed"
                and structured_role == "diluted_eps"
            ):
                # Written after every result row exists, so formula references
                # do not depend on structure line order.
                continue
            direct_key = {
                "diluted_eps_numerator": "direct_numerator",
                "diluted_weighted_average_shares": "direct_shares",
            }.get(str(structured_role))
            if annual_eps is not None and direct_key is not None:
                direct = annual_eps.get(direct_key) or {}
                direct_periods = direct.get("source_periods") or []
                if direct.get("status") == "computed" and len(direct_periods) == 1:
                    target.value = _template_value(
                        direct["value"], result["unit"], role=str(structured_role),
                    )
                    template_cell_styles["financials"].append({
                        "range": target.coordinate, "style": "hardcoded_input",
                        "number_kind": _template_number_kind(result["unit"]),
                    })
                    result_cells[(result["ref"], label)] = (
                        f"'Financials'!{target.coordinate}"
                    )
                    formula_map.append({
                        "cell": f"Financials!{target.coordinate}",
                        "model_cell_ref": (
                            "annual-structured-direct:"
                            + ",".join(direct_periods[0].get("source_accessions") or [])
                        ),
                        "formula": str(target.value),
                        "model_formula": (
                            "direct_annual_filed_value / 1000000"
                            if result["unit"].casefold() == "usd"
                            or structured_role == "diluted_weighted_average_shares"
                            else "direct_annual_filed_value"
                        ),
                        "model_label": result["label"],
                    })
                    continue
            annual_share = annual_share_forecasts.get(label)
            if (
                structured_role == "diluted_weighted_average_shares"
                and annual_share is not None
                and annual_share.get("status") == "computed"
            ):
                source_periods = annual_share["source_periods"]
                coordinates = [
                    result_cells.get((result["ref"], str(cell["period_end"])))
                    for cell in source_periods
                ]
                if len(coordinates) == 4 and all(coordinates):
                    weights = [_inclusive_duration_days(cell) for cell in source_periods]
                    target.value = "=(" + "+".join(
                        f"{coordinate}*{weight}"
                        for coordinate, weight in zip(coordinates, weights)
                    ) + f")/{sum(weights)}"
                    template_cell_styles["financials"].append({
                        "range": target.coordinate, "style": "local_formula",
                        "number_kind": _template_number_kind(result["unit"]),
                    })
                    result_cells[(result["ref"], label)] = (
                        f"'Financials'!{target.coordinate}"
                    )
                    formula_map.append({
                        "cell": f"Financials!{target.coordinate}",
                        "model_cell_ref": "structured-day-weighted-quarter-shares",
                        "formula": target.value,
                        "model_formula": "day_weighted_quarters",
                        "model_label": result["label"],
                    })
                    continue
            annual_semantic_ok = (
                structured_line.get("annual_semantics") == "sum_quarters"
                if structured_line is not None else result["unit"] != "ratio"
            )
            if (
                flow_proven
                and len(quarter_cols) == 4
                and annual_semantic_ok
            ):
                target.value = f"=SUM({_col(quarter_cols[0])}{rr}:{_col(quarter_cols[-1])}{rr})"
                template_cell_styles["financials"].append({
                    "range": target.coordinate, "style": "local_formula",
                    "number_kind": _template_number_kind(result["unit"]),
                })
                result_cells[(result["ref"], label)] = (
                    f"'Financials'!{target.coordinate}"
                )
                formula_map.append({"cell": f"Financials!{target.coordinate}",
                                    "model_cell_ref": None, "formula": target.value,
                                    "model_formula": "annual_from_four_fiscal_quarters",
                                    "model_label": result["label"]})
            else:
                target.value = None
                target.comment = None
                if structured_line is not None and not annual_semantic_ok:
                    reason = (
                        f"structure annual semantics "
                        f"{structured_line.get('annual_semantics')} is not sum_quarters"
                    )
                elif result["unit"] == "ratio":
                    reason = (
                        "ratio lacks an explicit frozen numerator/denominator "
                        "formula contract"
                    )
                elif not flow_proven:
                    reason = "four typed duration-quarter cells are not proven"
                else:
                    reason = f"{len(quarter_cols)}/{len(group)} supported quarters"
                gaps.append(f"{result['ref']} {label}: annual unavailable; {reason}")

    annual_eps_outcomes = {**annual_eps_results, **annual_share_forecasts}
    if annual_eps_outcomes:
        eps_result = next(
            (item for item in model["results"] if item.get("role") == "diluted_eps"),
            None,
        )
        numerator_result = next(
            (item for item in model["results"]
             if item.get("role") == "diluted_eps_numerator"), None,
        )
        share_result = next(
            (item for item in model["results"]
             if item.get("role") == "diluted_weighted_average_shares"), None,
        )
        if eps_result is not None and numerator_result is not None and share_result is not None:
            eps_row = result_rows[eps_result["ref"]]
            for label, outcome in annual_eps_outcomes.items():
                if outcome.get("status") != "computed":
                    gaps.append(
                        f"{eps_result['ref']} {label}: annual unavailable; "
                        f"{outcome.get('reason')}"
                    )
                    continue
                numerator_cell = result_cells.get((numerator_result["ref"], label))
                share_cell = result_cells.get((share_result["ref"], label))
                if numerator_cell is None or share_cell is None:
                    gaps.append(
                        f"{eps_result['ref']} {label}: annual direct numerator/share "
                        "cells are unavailable"
                    )
                    continue
                target = financials.cell(eps_row, annual_columns[label])
                target.value = f"={numerator_cell}/{share_cell}"
                template_cell_styles["financials"].append({
                    "range": target.coordinate, "style": "local_formula",
                    "number_kind": _template_number_kind(eps_result["unit"], per_share=True),
                })
                result_cells[(eps_result["ref"], label)] = (
                    f"'Financials'!{target.coordinate}"
                )
                formula_map.append({
                    "cell": f"Financials!{target.coordinate}",
                    "model_cell_ref": "historical-structured-annual-diluted-eps",
                    "formula": target.value,
                    "model_formula": eps_result["formula"],
                    "model_label": eps_result["label"],
                })
    if len(annual_groups) >= 2:
        annual_positions = {
            label: index for index, (label, _) in enumerate(annual_groups)
        }
        for result in model["results"]:
            rr = result_rows[result["ref"]]
            available = [
                label for label, _ in annual_groups
                if (result["ref"], label) in result_cells
            ]
            if len(available) < 2:
                continue
            first_label, last_label = available[0], available[-1]
            years = annual_positions[last_label] - annual_positions[first_label]
            first = result_cells.get((result["ref"], first_label))
            last = result_cells.get((result["ref"], last_label))
            if first is None or last is None or years <= 0:
                continue
            local_first = first.rsplit("!", 1)[-1]
            local_last = last.rsplit("!", 1)[-1]
            target = financials.cell(rr, cagr_column)
            target.value = (
                f'=IF(AND({local_first}>0,{local_last}>0),'
                f'({local_last}/{local_first})^(1/{years})-1,"")'
            )
            template_cell_styles["financials"].append({
                "range": target.coordinate, "style": "local_formula",
                "number_kind": "percentage",
            })
            formula_map.append({
                "cell": f"Financials!{target.coordinate}", "model_cell_ref": None,
                "formula": target.value, "model_formula": "annual_cagr",
                "model_label": result["label"],
            })

    for ci, value in enumerate(["Metric", "Value", "Role", "Method / gap"], 1):
        valuation_ws.cell(4, ci, value)
    template_row_styles["valuation"].append({
        "row": 4, "style": "valuation_section", "level": 0,
    })
    if valuation_scenario is not None:
        scenario = dict(valuation_scenario)
        allowed = {"scenario_ref", "as_of", "multiple_kind", "multiple",
                   "net_cash", "diluted_shares", "required_return",
                   "year_fraction", "content_hash"}
        if set(scenario) != allowed:
            raise FundWorkbookExportError("valuation scenario has an invalid closed shape")
        asserted = scenario.pop("content_hash")
        if asserted != content_hash(scenario):
            raise FundWorkbookExportError("valuation scenario content hash is invalid")
        if scenario["multiple_kind"] not in {"ev_revenue", "pe_net_income"}:
            raise FundWorkbookExportError("valuation scenario multiple_kind is unsupported")
        for key in ("multiple", "diluted_shares", "required_return", "year_fraction"):
            if _number(scenario[key]) <= 0:
                raise FundWorkbookExportError(f"valuation scenario {key} must be positive")
        rows = [
            ("Selected multiple", _number(scenario["multiple"]), "Input"),
            ("Net cash (USD MM)",
             _template_value(scenario["net_cash"], "USD"), "Input"),
            ("Diluted shares (MM)",
             _number(scenario["diluted_shares"]) / FUND_MONETARY_DISPLAY_SCALE,
             "Input"),
            ("Required return", _number(scenario["required_return"]), "Input"),
            ("Year fraction", _number(scenario["year_fraction"]), "Input"),
        ]
        for offset, (label, value, role) in enumerate(rows, 5):
            valuation_ws.cell(offset, 1, label)
            valuation_ws.cell(offset, 2, value)
            valuation_ws.cell(offset, 3, role)
            template_cell_styles["valuation"].append({
                "range": valuation_ws.cell(offset, 2).coordinate,
                "style": "assumption_input",
                "number_kind": (
                    "percentage" if offset == 8 else "multiple" if offset in {5, 9}
                    else "amount_one_decimal" if offset == 6 else "amount"
                ),
            })
        target_ref = ("result:revenue" if scenario["multiple_kind"] == "ev_revenue"
                      else "result:net_income")
        complete = [(label, group) for label, group in annual_groups
                    if len(group) == 4 and all((target_ref, p) in result_cells for p in group)]
        if complete:
            label, _ = complete[-1]
            annual_col = annual_columns[label]
            model_row = result_rows[target_ref]
            base = f"'Financials'!{_col(annual_col)}{model_row}"
            valuation_ws.append(["Forecast (USD MM)", f"={base}", "Linked", target_ref])
            if scenario["multiple_kind"] == "ev_revenue":
                valuation_ws.append(["Equity value (MM)", "=B10*B5+B6", "Formula",
                                     "EV bridge"])
            else:
                valuation_ws.append(["Equity value (MM)", "=B10*B5", "Formula",
                                     "P/E bridge"])
            valuation_ws.append(["Future target price", "=B11/B7", "Formula",
                                 "Per share"])
            valuation_ws.append(["Discounted target price", "=B12/(1+B8)^B9",
                                 "Formula", "Discounted"])
            for row in range(10, 14):
                template_cell_styles["valuation"].append({
                    "range": valuation_ws.cell(row, 2).coordinate,
                    "style": "cross_sheet_formula" if row == 10 else "local_formula",
                    "number_kind": "per_share" if row in {12, 13} else "amount_one_decimal",
                })
                formula_map.append({"cell": f"Valuation!B{row}",
                                    "model_cell_ref": target_ref if row == 10 else None,
                                    "formula": valuation_ws.cell(row, 2).value,
                                    "model_formula": valuation_ws.cell(row, 4).value})
        else:
            valuation_ws.append(["Scenario output", None, "unavailable",
                                 "No complete supported fiscal-year forecast"])
    elif valuation:
        for metric in valuation.get("metrics") or []:
            status = metric.get("status")
            valuation_ws.append([metric.get("label") or metric.get("metric"),
                                 _number(metric["value"]) if status == "available" else None,
                                 "derived authority value" if status == "available" else "unavailable",
                                 metric.get("formula") if status == "available" else metric.get("reason")])
    else:
        valuation_ws.append(["Valuation", None, "unavailable", "No bound valuation authority version"])

    source_rows = [
        ["ForecastModel", model["id"], model["content_hash"], model["created_at"]],
        ["CompanyModelSpec", model["spec_ref"], model["spec_hash"], ""],
        ["ModelInput", model["inputs_hash"], model["inputs_hash"], ""],
        ["Formula", model["formula_ref"], model["formula_hash"], ""],
    ]
    if mission_binding:
        source_rows.append(["CoverageMission", mission_binding["ref"],
                            mission_binding["content_hash"],
                            mission_binding["created_at"]])
    if valuation:
        source_rows.append(["ValuationSnapshot", valuation["id"], valuation["content_hash"], valuation["as_of"]])
    if valuation_scenario:
        source_rows.append(["ValuationScenario", valuation_scenario["scenario_ref"],
                            valuation_scenario["content_hash"], valuation_scenario["as_of"]])
    if calendar_binding:
        source_rows.append(["FiscalCalendar", calendar_binding["calendar_ref"],
                            calendar_binding["content_hash"], calendar_binding["as_of"]])
    for ci, value in enumerate(["Authority", "Reference", "SHA-256 / binding", "As of"], 1):
        sources.cell(4, ci, value)
    for values in source_rows:
        sources.append(values)
    formula_digest = content_hash({"layout_version": LAYOUT_VERSION, "formulas": formula_map})
    for ci, value in enumerate(
            ["Cell", "Model cell", "Excel formula", "Internal formula",
             "Original model label"], 1):
        manifest.cell(4, ci, value)
    for item in formula_map:
        # Formula Map is an audit table. Keep the formula literal instead of
        # executing a second, context-free copy of it on this sheet.
        manifest.append([item["cell"], item["model_cell_ref"],
                         "'" + item["formula"], item["model_formula"],
                         item.get("model_label")])
    manifest.append([])
    manifest.append(["Layout version", LAYOUT_VERSION])
    manifest.append(["Monetary display scale", FUND_MONETARY_DISPLAY_SCALE])
    manifest.append([
        "Monetary display transform",
        "ISO-currency and diluted-share inputs divided by 1,000,000; "
        "raw authority retained by immutable references and hashes",
    ])
    manifest.append(["Formula map SHA-256", formula_digest])
    manifest.append(["Gaps", len(gaps)])
    for gap in gaps:
        manifest.append(["Gap", gap])
    for ws in (sources, manifest):
        _style_sheet(ws, Font, PatternFill)
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 54)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:4"
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if cell.row not in {1, 4} and cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=10, color=cell.font.color)
    for ws in (valuation_ws, financials, driver):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    template_plan = build_fund_xlsx_template_plan(
        sheet_names={"valuation": "Valuation", "financials": "Financials", "driver": "Driver"},
        annual_periods=displayed_annual_labels, quarterly_periods=quarter_labels,
        unit_labels={"financials": "(USD MM)", "driver": "(US$mm)"},
        row_styles=template_row_styles, cell_styles=template_cell_styles,
    )
    apply_fund_xlsx_template(wb, template_plan)
    # Audit sheets contain long immutable references, hashes, and formula text.
    # Wrap them at readable widths so printed/PDF copies retain the full value.
    from openpyxl.styles import Alignment
    audit_widths = {
        sources: {"A": 24, "B": 56, "C": 68, "D": 24},
        manifest: {"A": 26, "B": 54, "C": 62, "D": 62, "E": 42},
    }
    for ws, widths in audit_widths.items():
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        for row_cells in ws.iter_rows(min_row=4):
            has_long_value = False
            for cell in row_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                has_long_value = has_long_value or len(str(cell.value or "")) > 45
            if has_long_value:
                ws.row_dimensions[row_cells[0].row].height = 42
    output = Path(output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FundWorkbookExportError("output already exists; refusing to overwrite")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    with tempfile.NamedTemporaryFile(
            prefix=f".{output.name}.", suffix=".tmp", dir=output.parent,
            delete=False) as handle:
        temporary_output = Path(handle.name)
    try:
        os.chmod(temporary_output, 0o600)
        wb.save(temporary_output)
        os.replace(temporary_output, output)
    except BaseException:
        temporary_output.unlink(missing_ok=True)
        raise
    return {"path": str(output), "company_ref": model["company_ref"],
            "formula_count": len(formula_map), "formula_map_hash": formula_digest,
            "gaps": gaps}


def _verify_statement_filing(
    connection: sqlite3.Connection, filing: Mapping[str, Any],
) -> None:
    """Replay accepted statement-line wire versions from immutable rows."""

    rows = connection.execute(
        "SELECT * FROM coverage_mission_statement_lines "
        "WHERE ingest_id=? ORDER BY ordinal", (filing["ingest_id"],),
    ).fetchall()
    if len(rows) != int(filing["line_count"]):
        raise FundWorkbookExportError("annual filing authority line count is invalid")
    identity = {
        "company_ref": filing["company_ref"], "cik": filing["cik"],
        "accession": filing["accession"], "form": filing["form"],
        "line_count": int(filing["line_count"]),
    }
    expected_ingest_id = f"statement-ingest:{content_hash(identity)[:32]}"
    if filing["ingest_id"] != expected_ingest_id:
        raise FundWorkbookExportError("annual filing authority ingest identity is invalid")
    for ordinal, row in enumerate(rows):
        if (row["ingest_id"] != expected_ingest_id
                or row["ordinal"] != ordinal
                or row["line_id"] != f"{expected_ingest_id}#{ordinal}"):
            raise FundWorkbookExportError("annual filing authority line identity is invalid")
    common_fields = (
        "statement", "concept", "label", "level", "parent_concept",
        "is_breakdown", "dimension_axis", "dimension_member", "period_start",
        "period_end", "value", "unit", "balance",
    )
    line_hashes = []
    for fields in (
            common_fields,
            common_fields[:8] + ("dimension_count",) + common_fields[8:]):
        lines = []
        for row in rows:
            line = {field: row[field] for field in fields}
            line["is_breakdown"] = bool(line["is_breakdown"])
            lines.append(line)
        line_hashes.append(content_hash(lines))
    body = {
        "company_ref": filing["company_ref"], "cik": filing["cik"],
        "accession": filing["accession"], "form": filing["form"],
        "line_count": filing["line_count"], "entity_name": filing["entity_name"],
        "filed": filing["filed"], "report_date": filing["report_date"],
        "source_record_refs": filing["source_record_refs"],
        "governance_ref": filing["governance_ref"],
        "governance_hash": filing["governance_hash"],
    }
    matches = [line_hash for line_hash in line_hashes
               if content_hash({**body, "statement_lines_hash": line_hash})
               == filing["content_hash"]]
    if len(matches) != 1:
        raise FundWorkbookExportError("annual filing authority hash is invalid")


def export_company_workbook(
    core_db: Path, company_ref: str, output: Path,
    *, valuation_scenario: Mapping[str, Any] | None = None,
    calendar_binding: Mapping[str, Any] | None = None,
    mission_ref: str | None = None,
) -> dict[str, Any]:
    """Read one exact Core snapshot and export without opening it writable."""
    source_path = Path(core_db).expanduser().resolve()
    if not source_path.is_file():
        raise FundWorkbookExportError("Core database does not exist")
    if Path(output).expanduser().resolve() == source_path:
        raise FundWorkbookExportError("output must not replace the Core database")
    with tempfile.TemporaryDirectory(prefix="dalton-fund-export-") as temp_dir:
        copied = Path(temp_dir) / "core.sqlite"
        source = sqlite3.connect(f"file:{source_path.as_posix()}?mode=ro", uri=True)
        target = sqlite3.connect(copied)
        try:
            source.backup(target)
        finally:
            target.close()
            source.close()
        with DaltonStore(copied) as store:
            models = ForecastModelAuthority(store)
            missions = CoverageMissionAuthority(store)
            if mission_ref is None:
                active = []
                rows = store.connection.execute(
                    "SELECT mission_ref FROM coverage_mission_pointer "
                    "ORDER BY mission_ref").fetchall()
                for row in rows:
                    candidate = missions.active_mission(row["mission_ref"])
                    if company_ref in {
                            member["company_ref"] for member in candidate["universe"]}:
                        active.append(candidate)
                if len(active) != 1:
                    raise FundWorkbookExportError(
                        "company must belong to exactly one active mission; pass mission_ref")
                mission = active[0]
            else:
                mission = missions.active_mission(mission_ref)
                if company_ref not in {
                        member["company_ref"] for member in mission["universe"]}:
                    raise FundWorkbookExportError(
                        "company is outside the requested active mission")
            model = models.latest(company_ref)
            if model is None:
                raise FundWorkbookExportError("company has no current model")
            spec_row = store.connection.execute(
                "SELECT * FROM coverage_mission_company_model_specs WHERE spec_id=?",
                (model["spec_ref"],),
            ).fetchone()
            if spec_row is None:
                raise FundWorkbookExportError(
                    "forecast model's exact specification is unavailable")
            spec = missions._model_spec_row(spec_row)
            if spec["mission_version_ref"] != mission["id"]:
                raise FundWorkbookExportError(
                    "latest forecast model is not bound to the current mission version")
            inputs = build_model_inputs(missions, spec)
            member = next(
                item for item in mission["universe"]
                if item["company_ref"] == company_ref
            )
            annual_filings = [
                item for item in missions.statement_filings(company_ref)
                if item["form"] == "10-K"
            ]
            latest_annual = (
                max(annual_filings,
                    key=lambda item: (item["report_date"], item["accession"]))
                if annual_filings else None
            )
            if calendar_binding is None:
                if annual_filings:
                    filing = latest_annual
                    _verify_statement_filing(store.connection, filing)
                    base = {
                        "calendar_ref": filing["ingest_id"],
                        "source_hash": filing["content_hash"],
                        "as_of": filing["report_date"],
                        "fiscal_year_end_month": int(filing["report_date"][5:7]),
                    }
                    calendar_binding = {**base, "content_hash": content_hash(base)}
            valuation = ValuationSnapshotAuthority(store).latest_version(company_ref)
            return export_fund_workbook(
                output, model=model, spec=spec, inputs=inputs,
                valuation=valuation, valuation_scenario=valuation_scenario,
                calendar_binding=calendar_binding,
                mission_binding={"ref": mission["id"],
                                 "content_hash": mission["content_hash"],
                                 "created_at": mission["created_at"],
                                 "ticker": member["ticker"],
                                 "entity_name": (
                                     latest_annual["entity_name"]
                                     if latest_annual else None
                                 )})


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export a governed company model as XLSX")
    parser.add_argument("--state-dir", type=Path, required=True)
    parser.add_argument("--company-ref", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--valuation-scenario", type=Path)
    parser.add_argument("--calendar-binding", type=Path)
    parser.add_argument("--mission-ref")
    args = parser.parse_args(argv)
    output = args.output.expanduser().resolve()
    protected = {args.state_dir.expanduser().resolve() / "core.sqlite"}
    for candidate in (args.valuation_scenario, args.calendar_binding):
        if candidate:
            protected.add(candidate.expanduser().resolve())
    if output in protected:
        parser.error("output must not replace a database or input configuration")
    scenario = None
    if args.valuation_scenario:
        scenario = json.loads(args.valuation_scenario.read_text(encoding="utf-8"))
    calendar = None
    if args.calendar_binding:
        calendar = json.loads(args.calendar_binding.read_text(encoding="utf-8"))
    result = export_company_workbook(
        args.state_dir / "core.sqlite", args.company_ref, output,
        valuation_scenario=scenario, calendar_binding=calendar,
        mission_ref=args.mission_ref)
    print(canonical_json(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
