"""Deterministic, inert structure manifests for prior Office artifacts."""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from .store import content_hash

_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _digest(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def docx_artifact_manifest(raw: bytes) -> dict[str, Any]:
    """Describe DOCX body order and assets without opening embedded objects."""

    with zipfile.ZipFile(_BytesReader(raw)) as archive:
        names = set(archive.namelist())
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        relationships: dict[str, tuple[str, str, str | None]] = {}
        rel_name = "word/_rels/document.xml.rels"
        if rel_name in names:
            for rel in ElementTree.fromstring(archive.read(rel_name)):
                relationships[str(rel.attrib.get("Id"))] = (
                    str(rel.attrib.get("Type", "")).rsplit("/", 1)[-1],
                    str(rel.attrib.get("Target", "")), rel.attrib.get("TargetMode"),
                )
        body_node = document.find(f"{_W_NS}body")
        body: list[dict[str, Any]] = []
        for index, node in enumerate([] if body_node is None else list(body_node)):
            kind = node.tag.rsplit("}", 1)[-1]
            text = "".join(part.text or "" for part in node.iter(f"{_W_NS}t")).strip()
            refs = []
            for child in node.iter():
                for attr, value in child.attrib.items():
                    if attr.startswith(_R_NS) and value in relationships:
                        rel_type, target, mode = relationships[value]
                        external = mode == "External"
                        refs.append({"relationship_id": value, "relationship_type": rel_type,
                                     "target": None if external else target,
                                     "target_sha256": _digest(target.encode("utf-8")),
                                     "external": external})
            if kind in {"p", "tbl"}:
                body.append({"order": index, "kind": "paragraph" if kind == "p" else "table",
                             "text_sha256": _digest(text.encode("utf-8")),
                             "text_chars": len(text), "relationships": refs})
        assets = []
        for name in sorted(names):
            if not (name.startswith("word/media/") or
                    name.startswith("word/embeddings/")):
                continue
            payload = archive.read(name)
            assets.append({"part": name, "kind": "embedded_workbook" if
                           name.startswith("word/embeddings/") else "media",
                           "sha256": _digest(payload), "bytes": len(payload)})
    result = {"schema_version": "prior-office-artifact-0.1", "format": "docx",
              "body": body, "assets": assets}
    result["manifest_hash"] = content_hash(result)
    return result


def xlsx_artifact_manifest(path: str | Path) -> dict[str, Any]:
    """Workbook topology and inactive dependency identities, never values."""

    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=False, data_only=False,
                             keep_links=True)
    try:
        sheets = [{"name": sheet.title, "state": sheet.sheet_state,
                   "max_row": sheet.max_row, "max_column": sheet.max_column,
                   "freeze_panes": None if sheet.freeze_panes is None else str(sheet.freeze_panes)}
                  for sheet in workbook.worksheets]
        cells = []
        formula_identity = []
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            for row_index, row in enumerate(values):
                for column_index, value in enumerate(row):
                    formula = value if isinstance(value, str) and value.startswith("=") else ""
                    if not formula and not isinstance(value, (int, float)):
                        continue
                    cell = sheet.cell(row_index + 1, column_index + 1)
                    label = _nearest_label(values, row_index, column_index)
                    period = _period(values, row_index, column_index)
                    role = ("external_formula" if "_xll." in formula or "[" in formula
                            else "cross_sheet_formula" if formula and "!" in formula
                            else "formula" if formula else "hardcoded_input")
                    unit, basis = _unit(label, str(cell.number_format))
                    cells.append({"sheet": sheet.title, "cell": cell.coordinate,
                                  "label": label, "formula": formula, "unit": unit,
                                  "unit_basis": basis, "cell_role": role,
                                  "number_format": str(cell.number_format)[:200],
                                  "font_color": _color(cell), "period": period,
                                  "period_status": "forecast" if period and period.endswith("E") else "unknown",
                                  "source_type": "prior_human_formula" if formula else "prior_human_hardcode"})
                    if formula:
                        formula_identity.append({"sheet": sheet.title, "cell": cell.coordinate,
                                                 "formula": formula})
        result = {"schema_version": "prior-office-artifact-0.1", "format": "xlsx",
                  "sheets": sheets, "external_link_count": len(workbook._external_links),
                  "macro_present": bool(getattr(workbook, "vba_archive", None)),
                  "cells": cells,
                  "import_metadata": {"complete": True, "truncations": [],
                                      "formula_map_hash": content_hash(formula_identity)}}
    finally:
        workbook.close()
    result["manifest_hash"] = content_hash(result)
    return result


def _nearest_label(values: list[tuple[Any, ...]], row: int, column: int) -> str:
    for back in range(column - 1, -1, -1):
        value = values[row][back]
        if isinstance(value, str) and value.strip() and not value.startswith("="):
            return value.strip()[:200]
    return ""


def _period(values: list[tuple[Any, ...]], row: int, column: int) -> str | None:
    for up in range(row - 1, -1, -1):
        if column >= len(values[up]):
            continue
        value = values[up][column]
        if isinstance(value, (str, int)):
            text = str(value).strip()
            if re.search(r"(?:FY|CY|[1-4]Q)?\d{2,4}E?$", text):
                return text[:40]
    return None


def _unit(label: str, number_format: str) -> tuple[str | None, str]:
    if "%" in number_format:
        return "percent", "number_format"
    folded = label.casefold()
    if any(token in folded for token in ("$", "usd", "revenue", "income", "cost")):
        return "currency", "label"
    if any(token in folded for token in ("multiple", " p/e", " ev/", "ratio")):
        return "ratio", "label"
    return None, "unknown"


def _color(cell: Any) -> str | None:
    color = getattr(getattr(cell, "font", None), "color", None)
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str):
        return color.rgb
    if color.type == "theme" and isinstance(color.theme, int):
        return f"theme:{color.theme}"
    return None


class _BytesReader:
    def __init__(self, raw: bytes) -> None:
        import io
        self._buffer = io.BytesIO(raw)

    def __getattr__(self, name: str) -> Any:
        return getattr(self._buffer, name)


__all__ = ["docx_artifact_manifest", "xlsx_artifact_manifest"]
