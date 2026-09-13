"""Create a redacted, inert fund-workbook style candidate from local XLSX bytes."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import stat
import zipfile
from pathlib import Path
from typing import Any, Mapping, Sequence

from .fund_xlsx_template import STYLE_SCHEMA_VERSION

CANDIDATE_SCHEMA_VERSION = "fund-xlsx-template-import-candidate-0.1"
STYLE_NAMES = (
    "unit_header", "period_header", "major_section", "section", "label",
    "subtotal", "growth_label", "hardcoded_input", "cross_sheet_formula",
    "local_formula", "assumption_input", "valuation_section",
)
NUMBER_STYLE_NAMES = {
    "amount": "number_amount", "amount_one_decimal": "number_amount_one_decimal",
    "percentage": "number_percentage", "multiple": "number_multiple",
    "date_month": "number_date_month", "per_share": "number_per_share",
}
MAX_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
ALLOWED_FONTS = {"Arial", "Calibri"}
ALLOWED_NUMBER_FORMATS = {
    "_(#,##0_);\\(#,##0\\);_(?\\-?_);@",
    "_(#,##0.0_);\\(#,##0.0\\);_(?\\-?_);@",
    "0.0%", "#,##0.0\\x", "mmm-yy",
    "_(#,##0.00_);\\(#,##0.00\\);_(?\\-?_);@",
}


class FundXlsxTemplateImportError(ValueError):
    pass


def _need(value: Any, reason: str) -> None:
    if not value:
        raise FundXlsxTemplateImportError(reason)


def _canonical(value: Any) -> str:
    wire = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n"
    return hashlib.sha256(wire.encode()).hexdigest()


def _source_bytes(path: Path) -> bytes:
    path = path.expanduser().absolute()
    _need(path.is_file() and not path.is_symlink() and path == path.resolve(),
          "source workbook must be a direct regular file")
    held = path.lstat()
    _need(stat.S_ISREG(held.st_mode) and held.st_size <= MAX_ARCHIVE_BYTES,
          "source workbook is unavailable or too large")
    data = path.read_bytes()
    _need(path.lstat() == held, "source workbook changed while read")
    return data


def _safe_archive(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            rows = archive.infolist()
            _need(rows and sum(row.file_size for row in rows) <= MAX_EXPANDED_BYTES,
                  "workbook archive expands beyond the import limit")
            for row in rows:
                name = Path(row.filename)
                _need(not name.is_absolute() and ".." not in name.parts
                      and not (row.flag_bits & 0x1), "workbook archive is unsafe")
                _need(row.compress_size or row.file_size == 0,
                      "workbook archive has an invalid compressed member")
                if row.compress_size:
                    _need(row.file_size <= row.compress_size * 200,
                          "workbook archive compression ratio is unsafe")
            names = {row.filename for row in rows}
            _need(not any(name.startswith(("xl/externalLinks/", "xl/connections"))
                          or name.endswith("vbaProject.bin") for name in names),
                  "workbook contains unsupported external or executable content")
            return hashlib.sha256(archive.read("xl/theme/theme1.xml")).hexdigest()
    except KeyError as exc:
        raise FundXlsxTemplateImportError("workbook theme is missing") from exc
    except zipfile.BadZipFile as exc:
        raise FundXlsxTemplateImportError("source workbook is not a valid XLSX archive") from exc


def _color(value: Any) -> str | None:
    if value is None or value.type != "rgb" or not isinstance(value.rgb, str):
        return None
    rgb = value.rgb.upper()
    _need(len(rgb) == 8, "style color is not an explicit ARGB value")
    return rgb


def _font(value: Any) -> dict[str, Any]:
    _need(isinstance(value.name, str) and value.name and isinstance(value.sz, (int, float)),
          "style font must have an explicit name and size")
    _need(value.name in ALLOWED_FONTS and 6 <= float(value.sz) <= 24,
          "style font is outside the renderer allowlist")
    result: dict[str, Any] = {"name": value.name, "size": float(value.sz)}
    if value.b: result["bold"] = True
    if value.i: result["italic"] = True
    color = _color(value.color)
    if color is not None: result["color"] = color
    return result


def _style(value: Any) -> dict[str, Any]:
    result: dict[str, Any] = {"font": _font(value.font)}
    fill = _color(value.fill.fgColor)
    if value.fill.fill_type is not None:
        _need(value.fill.fill_type == "solid" and fill is not None,
              "only explicit solid fills are supported")
        result["fill"] = fill
    if value.alignment.horizontal is not None:
        _need(value.alignment.horizontal in {"left", "center", "right"},
              "style alignment is unsupported")
        result["horizontal"] = value.alignment.horizontal
    edges = {}
    for name in ("top", "bottom", "left", "right"):
        edge = getattr(value.border, name)
        if edge is not None and edge.style is not None:
            color = _color(edge.color)
            _need(edge.style == "hair" and color is not None,
                  "only explicit hairline borders are supported")
            edges[name] = {"style": "hair", "color": color}
    if edges: result["border"] = edges
    return result


def import_fund_xlsx_template_candidate(
    workbook_path: Path, *, sheets: Mapping[str, str], template_slug: str,
) -> dict[str, Any]:
    """Extract only renderer-supported presentation properties."""
    _need(set(sheets) == {"valuation", "financials", "driver"}
          and len(set(sheets.values())) == 3, "sheet role binding is incomplete or ambiguous")
    _need(template_slug.isascii() and template_slug.replace("-", "").isalnum()
          and template_slug == template_slug.lower(), "template slug is invalid")
    data = _source_bytes(workbook_path)
    theme_sha256 = _safe_archive(data)
    try:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    except Exception as exc:
        raise FundXlsxTemplateImportError("workbook could not be parsed safely") from exc
    try:
        _need(all(isinstance(name, str) and name in book.sheetnames for name in sheets.values()),
              "bound workbook sheet is missing")
        named = {style.name: style for style in book._named_styles}
        required = {*STYLE_NAMES, "model_default", "valuation_default",
                    *NUMBER_STYLE_NAMES.values()}
        _need(required <= set(named), "workbook lacks required unambiguous named styles")
        financials, driver, valuation = (book[sheets[k]] for k in
                                          ("financials", "driver", "valuation"))
        _need(financials.freeze_panes == driver.freeze_panes == "E2"
              and financials.sheet_view.showGridLines is False
              and driver.sheet_view.showGridLines is False
              and valuation.sheet_view.showGridLines is False,
              "workbook grid layout is unsupported")
        _need(financials.sheet_format.defaultRowHeight
              == driver.sheet_format.defaultRowHeight
              and isinstance(financials.sheet_format.defaultRowHeight, (int, float))
              and isinstance(valuation.sheet_format.defaultRowHeight, (int, float)),
              "workbook default row heights are unsupported")
        def width(sheet: Any, column: str) -> float:
            value = sheet.column_dimensions[column].width
            _need(isinstance(value, (int, float)) and 0 < value <= 100,
                  "workbook column width is unsupported")
            return float(value)
        style_contract = {
            "schema_version": STYLE_SCHEMA_VERSION,
            "template_ref": f"fund-xlsx-template-style:{template_slug}:candidate-0.1",
            "source_proof": {
                "workbook_sha256": hashlib.sha256(data).hexdigest(),
                "exact_source_style_sha256": "0" * 64,
                "theme_xml_sha256": theme_sha256,
            },
            "sheet_order": ["valuation", "financials", "driver"],
            "model_grid": {
                "period_row": 1, "first_annual_column": 5,
                "hierarchy_label_columns": [1, 2, 3, 4],
                "period_column_width": width(financials, "E"),
                "annual_support": {
                    "default_headers": ["CAGR"],
                    "financials_width": width(financials, "H"),
                    "driver_width": width(driver, "H"),
                    "financials_gutter_before_width": width(financials, "G"),
                    "financials_gutter_after_width": width(financials, "I"),
                    "driver_gutter_before_width": width(driver, "G"),
                    "driver_gutter_after_width": width(driver, "I"),
                },
                "default_row_height": float(financials.sheet_format.defaultRowHeight),
                "freeze_cell": "E2", "show_grid_lines": False,
                "font": _font(named["model_default"].font),
                "column_widths": {
                    "gutter_1": width(financials, "A"),
                    "gutter_2": width(financials, "B"),
                    "gutter_3": width(financials, "C"),
                    "financials_label": width(financials, "D"),
                    "driver_label": width(driver, "D"),
                },
            },
            "valuation_grid": {
                "default_row_height": float(valuation.sheet_format.defaultRowHeight),
                "show_grid_lines": False,
                "font": _font(named["valuation_default"].font),
                "column_widths": {column: width(valuation, column) for column in "ABCDE"},
            },
            "styles": {name: _style(named[name]) for name in STYLE_NAMES},
            "number_formats": {kind: named[name].number_format
                               for kind, name in NUMBER_STYLE_NAMES.items()},
        }
        _need(set(style_contract["number_formats"].values()) <= ALLOWED_NUMBER_FORMATS,
              "number format is outside the renderer allowlist")
        style_hash = _canonical({key: value for key, value in style_contract.items()
                                 if key != "source_proof"})
        style_contract["source_proof"]["exact_source_style_sha256"] = style_hash
        candidate = {
            "schema_version": CANDIDATE_SCHEMA_VERSION,
            "status": "candidate_pending_owner_review",
            "activation_authorized": False,
            "template_slug": template_slug,
            "source": {"workbook_sha256": hashlib.sha256(data).hexdigest(),
                       "archive_bytes": len(data), "sheet_count": len(book.sheetnames)},
            "style_contract": style_contract,
            "redaction": {"cell_values": "omitted", "formulas": "omitted",
                          "defined_names": "omitted", "comments": "omitted",
                          "hyperlinks": "omitted", "external_links": "rejected"},
        }
        candidate["content_hash"] = _canonical(candidate)
        return candidate
    finally:
        book.close()


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--valuation-sheet", required=True)
    parser.add_argument("--financials-sheet", required=True)
    parser.add_argument("--driver-sheet", required=True)
    parser.add_argument("--template-slug", required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    candidate = import_fund_xlsx_template_candidate(
        args.workbook, sheets={"valuation": args.valuation_sheet,
                               "financials": args.financials_sheet,
                               "driver": args.driver_sheet},
        template_slug=args.template_slug)
    payload = json.dumps(candidate, ensure_ascii=False, sort_keys=True,
                         separators=(",", ":")) + "\n"
    if args.output is None:
        print(payload, end="")
    else:
        _need(not args.output.exists() and not args.output.is_symlink(), "output must be new")
        fd = os.open(args.output, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as stream: stream.write(payload)
    return 0


if __name__ == "__main__":
    try: raise SystemExit(main())
    except FundXlsxTemplateImportError as exc: raise SystemExit(f"STOP: {exc}") from exc
