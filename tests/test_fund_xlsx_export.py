from __future__ import annotations

import shutil
import sqlite3
import subprocess
import tempfile
import unittest
from pathlib import Path

from dalton_core.company_model_inputs import build_model_inputs
from dalton_core.fund_xlsx_export import (
    FundWorkbookExportError,
    _formula_for,
    _four_quarter_flow_cells,
    _verify_statement_filing,
    export_company_workbook,
    export_fund_workbook,
)
from dalton_core.model_forecast_driver import ForecastModelAuthority
from dalton_core.store import DaltonStore, content_hash
from tests.test_model_forecast_driver import ledger, model, spec


class FundXlsxExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = DaltonStore(":memory:")
        self.addCleanup(self.store.close)
        self.specification = spec(quarters=8)
        self.inputs = build_model_inputs(ledger(), self.specification)
        self.model = ForecastModelAuthority(self.store).publish(
            model(specification=self.specification))
        self.path = Path(self.temp.name) / "fund-model.xlsx"

    def scenario(self):
        value = {"scenario_ref": "valuation-scenario:test", "as_of": "2026-09-10",
                 "multiple_kind": "ev_revenue", "multiple": "2",
                 "net_cash": "100000000", "diluted_shares": "100000000",
                 "required_return": "0.10", "year_fraction": "1"}
        return {**value, "content_hash": content_hash(value)}

    def calendar(self):
        value = {"calendar_ref": "fiscal-calendar:acn:2026-09-10",
                 "source_hash": "1" * 64, "as_of": "2026-09-10",
                 "fiscal_year_end_month": 8}
        return {**value, "content_hash": content_hash(value)}

    def export(self, *, scenario=None):
        return export_fund_workbook(
            self.path, model=self.model, spec=self.specification,
            inputs=self.inputs, valuation_scenario=scenario,
            calendar_binding=self.calendar())

    def test_real_internal_model_exports_formula_chain_and_exact_bindings(self):
        from openpyxl import load_workbook

        result = self.export()
        book = load_workbook(self.path, data_only=False)
        self.assertEqual(book.sheetnames,
                         ["Valuation", "Financials", "Driver", "Sources", "Formula Map"])
        headers = {cell.value: cell.column for cell in book["Financials"][1]}
        revenue = book["Financials"].cell(3, headers["4Q26E"])
        self.assertTrue(revenue.value.startswith("='Driver'!"), revenue.value)
        self.assertIn("*(1+", revenue.value)
        assumption_row = next(
            row for row in range(3, book["Driver"].max_row + 1)
            if "Revenues — Quarterly growth (ratio)"
            in str(book["Driver"].cell(row, 3).value)
        )
        assumption = book["Driver"].cell(
            assumption_row, headers["4Q26E"]
        )
        self.assertEqual(assumption.font.color.rgb[-6:], "0000FF")
        self.assertIn(self.model["content_hash"],
                      [cell.value for cell in book["Sources"]["C"]])
        hashes = [book["Formula Map"].cell(row, 2).value
                  for row in range(1, book["Formula Map"].max_row + 1)
                  if book["Formula Map"].cell(row, 1).value == "Formula map SHA-256"]
        self.assertEqual(hashes, [result["formula_map_hash"]])

    def test_humanized_expense_label_keeps_concept_bound_formula_chain(self):
        from openpyxl import load_workbook

        missions = ledger()
        for line in missions.lines:
            if line["concept"] == "us-gaap:SellingGeneralAndAdministrativeExpense":
                line["label"] = "Selling, general and administrative"
        inputs = build_model_inputs(missions, self.specification)
        store = DaltonStore(":memory:")
        self.addCleanup(store.close)
        forecast = ForecastModelAuthority(store).publish(
            model(missions=missions, specification=self.specification))
        expense = next(
            item for item in forecast["results"]
            if item["ref"] ==
            "result:operating_expense:us-gaap:SellingGeneralAndAdministrativeExpense"
        )
        self.assertEqual(expense["label"], "Selling, general and administrative")
        self.assertEqual(
            expense["formula"],
            "SellingGeneralAndAdministrativeExpense[k] = revenue[k] * share[k]",
        )
        result = export_fund_workbook(
            self.path, model=forecast, spec=self.specification, inputs=inputs,
            calendar_binding=self.calendar())
        book = load_workbook(self.path, data_only=False)
        sheet = book["Financials"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        forecast_col = headers["4Q26E"]
        result_rows = {item["ref"]: 3 + index
                       for index, item in enumerate(forecast["results"])}
        expense_formula = sheet.cell(result_rows[expense["ref"]], forecast_col).value
        self.assertIsInstance(expense_formula, str)
        self.assertIn("*'Driver'!", expense_formula)
        for ref in ("result:operating_income", "result:income_tax_expense",
                    "result:net_income"):
            self.assertTrue(str(sheet.cell(result_rows[ref], forecast_col).value).startswith("="))
        self.assertFalse(any(
            expense["ref"] in gap and "unavailable or unsupported" in gap
            for gap in result["gaps"]
        ))

    def test_complete_year_is_formula_and_partial_year_stays_blank(self):
        from openpyxl import load_workbook

        result = self.export()
        book = load_workbook(self.path, data_only=False)
        sheet = book["Financials"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        self.assertIsNone(sheet.cell(3, headers["2025"]).value)
        self.assertTrue(sheet.cell(3, headers["2026A/E"]).value.startswith("=SUM("))
        self.assertTrue(sheet.cell(3, headers["2027E"]).value.startswith("=SUM("))
        self.assertTrue(any("FY2025A: annual unavailable" in gap for gap in result["gaps"]))

    def test_annual_sum_requires_four_typed_duration_quarters(self):
        ends = ["2025-11-30", "2026-02-28", "2026-05-31", "2026-08-31"]
        starts = ["2025-09-01", "2025-12-01", "2026-03-01", "2026-06-01"]
        cells = {
            end: {"period": {"start": start, "end": end,
                             "kind": "quarter", "calendar": "company:fiscal"}}
            for start, end in zip(starts, ends)
        }
        self.assertTrue(_four_quarter_flow_cells(cells, ends))
        instant = {key: {"period": dict(value["period"])}
                   for key, value in cells.items()}
        instant[ends[2]]["period"]["start"] = None
        self.assertFalse(_four_quarter_flow_cells(instant, ends))
        wrong_kind = {key: {"period": dict(value["period"])}
                      for key, value in cells.items()}
        wrong_kind[ends[2]]["period"]["kind"] = "instant"
        self.assertFalse(_four_quarter_flow_cells(wrong_kind, ends))

    def test_fund_layout_and_formula_roles_are_explicit(self):
        from openpyxl import load_workbook

        self.export(scenario=self.scenario())
        book = load_workbook(self.path, data_only=False)
        sheet = book["Financials"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers[4:8], [
            "2025", "2026A/E", "2027E", "2028E*",
        ])
        self.assertIsNone(headers[8])
        self.assertEqual(headers[9:12], ["CAGR", None, "4Q25"])
        self.assertEqual(headers[11:16], [
            "4Q25", "1Q26", "2Q26", "3Q26", "4Q26E",
        ])
        self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "3366FF")
        self.assertEqual(sheet["A1"].font.color.rgb[-6:], "FFFFFF")
        self.assertEqual(sheet.freeze_panes, "E2")
        self.assertEqual(sheet.column_dimensions["A"].width, 1.875)
        self.assertEqual(sheet.column_dimensions["D"].width, 25.125)
        self.assertEqual(sheet.column_dimensions["I"].width, 1.375)
        self.assertEqual(sheet.column_dimensions["K"].width, 1.125)
        self.assertEqual(sheet.sheet_format.defaultRowHeight, 11.25)
        self.assertFalse(sheet.sheet_view.showGridLines)

        # Cross-sheet links are green; local quarter and annual formulas are black.
        self.assertEqual(sheet["P3"].font.color.rgb[-6:], "008000")
        self.assertEqual(sheet["P5"].font.color.rgb[-6:], "000000")
        self.assertEqual(sheet["F3"].font.color.rgb[-6:], "000000")
        self.assertEqual(sheet["J3"].number_format, "0.0%")
        self.assertNotIn("#,##0.0", sheet["F3"].number_format)
        valuation = book["Valuation"]
        valuation_rows = {
            valuation.cell(row, 1).value: row
            for row in range(1, valuation.max_row + 1)
        }
        self.assertEqual(
            valuation.cell(valuation_rows["Required return"], 2).number_format,
            "0.0%",
        )
        self.assertIn(
            "#,##0.00",
            valuation.cell(valuation_rows["Discounted target price"], 2).number_format,
        )
        self.assertEqual(sheet["B3"].value, "Revenue")
        self.assertEqual(sheet["C4"].value, "Cost of revenue")
        self.assertTrue(sheet["B3"].font.bold)
        self.assertEqual(sheet["P3"].font.name, "Arial")
        self.assertEqual(sheet["P3"].font.sz, 8)
        self.assertEqual(sheet["A2"].fill.fgColor.rgb[-6:], "99CCFF")
        self.assertEqual(book["Driver"]["A2"].fill.fgColor.rgb[-6:], "993300")
        self.assertEqual(
            sheet["J3"].value,
            '=IF(AND(F3>0,G3>0),(G3/F3)^(1/1)-1,"")',
        )
        self.assertEqual(valuation["A5"].value, "TSO (MM) — N/A")
        self.assertIsNone(valuation["B5"].value)
        self.assertEqual(valuation["E5"].value, 100)
        driver = book["Driver"]
        revenue_assumption_rows = [
            row for row in range(3, driver.max_row + 1)
            if driver.cell(row, 3).value == "Revenues — Quarterly growth (ratio)"
        ]
        self.assertEqual(revenue_assumption_rows, [8])
        self.assertEqual(
            sum(driver.cell(8, column).value is not None
                for column in range(12, driver.max_column + 1)),
            8,
        )
        self.assertEqual(valuation["A1"].fill.fgColor.rgb[-6:], "3366FF")
        self.assertEqual(valuation["A2"].value, "Company")
        self.assertEqual(valuation["B2"].value, "Unavailable")
        self.assertEqual(valuation["A8"].value, "GAAP EPS (USD)")
        self.assertEqual(valuation["A9"].value, "House — N/A")
        self.assertEqual(valuation["A10"].value, "Street — N/A")
        self.assertEqual(valuation["A12"].value, "Total Revenue (USD MM)")
        self.assertEqual(valuation["A13"].value, "House")
        self.assertEqual(valuation["A12"].fill.fgColor.rgb[-6:], "99CCFF")
        self.assertFalse(any(
            str(cell.value).startswith("result:")
            for row in valuation.iter_rows() for cell in row
            if cell.value is not None
        ))

    def test_monetary_display_scales_cells_and_preserves_formula_dimensions(self):
        from openpyxl import load_workbook

        export_fund_workbook(
            self.path,
            model=self.model,
            spec=self.specification,
            inputs=self.inputs,
            calendar_binding=self.calendar(),
            mission_binding={
                "ref": "coverage-mission-version:test:1",
                "content_hash": "1" * 64,
                "created_at": "2026-09-10T00:00:00+00:00",
                "ticker": "ACN",
                "entity_name": None,
            },
        )
        book = load_workbook(self.path, data_only=False)
        financials = book["Financials"]
        headers = {cell.value: cell.column for cell in financials[1]}
        revenue = financials.cell(3, headers["4Q26E"])
        self.assertEqual(financials["B3"].value, "Revenue")
        self.assertTrue(revenue.value.startswith("='Driver'!"))
        self.assertNotIn(",,", revenue.number_format)
        self.assertEqual(book["Driver"]["L3"].value, 1_000)
        self.assertEqual(financials["A1"].value, "(USD MM)")
        formula_map = book["Formula Map"]
        scale = next(
            formula_map.cell(row, 2).value
            for row in range(1, formula_map.max_row + 1)
            if formula_map.cell(row, 1).value == "Monetary display scale"
        )
        self.assertEqual(scale, 1_000_000)
        transform = next(
            formula_map.cell(row, 2).value
            for row in range(1, formula_map.max_row + 1)
            if formula_map.cell(row, 1).value == "Monetary display transform"
        )
        self.assertIn("divided by 1,000,000", transform)

    def test_financial_lines_are_reader_labels_and_audit_keeps_originals(self):
        from openpyxl import load_workbook

        self.export()
        book = load_workbook(self.path, data_only=False)
        financials = book["Financials"]
        labels = [
            next((financials.cell(row, column).value for column in range(1, 5)
                  if financials.cell(row, column).value is not None), None)
            for row in range(3, financials.max_row + 1)
        ]
        self.assertIn("Revenue", labels)
        self.assertIn("Selling, general & administrative", labels)
        self.assertIn("Free cash flow — Not available", labels)
        self.assertEqual(financials["A2"].value, "Income statement")
        for label in labels:
            self.assertNotIn("formula output", label)
            self.assertNotIn("calculated", label)
            self.assertNotIn("SellingGeneralAndAdministrativeExpense", label)
            self.assertLessEqual(len(label), 54)

        formula_map = book["Formula Map"]
        headers = {cell.value: cell.column for cell in formula_map[4]}
        originals = [formula_map.cell(row, headers["Original model label"]).value
                     for row in range(5, formula_map.max_row + 1)]
        self.assertIn("SellingGeneralAndAdministrativeExpense", originals)
        formulas = [formula_map.cell(row, headers["Excel formula"]).value
                    for row in range(5, formula_map.max_row + 1)]
        self.assertTrue(any(str(value).startswith("'=") for value in formulas))

    def test_audit_sheets_wrap_complete_references_hashes_and_formulas(self):
        from openpyxl import load_workbook

        self.export()
        book = load_workbook(self.path, data_only=False)
        sources = book["Sources"]
        formula_map = book["Formula Map"]
        self.assertGreaterEqual(sources.column_dimensions["B"].width, 56)
        self.assertGreaterEqual(sources.column_dimensions["C"].width, 68)
        self.assertGreaterEqual(formula_map.column_dimensions["C"].width, 62)
        self.assertGreaterEqual(formula_map.column_dimensions["D"].width, 62)
        source_row = next(
            row for row in range(5, sources.max_row + 1)
            if sources.cell(row, 1).value == "ForecastModel"
        )
        formula_row = next(
            row for row in range(5, formula_map.max_row + 1)
            if str(formula_map.cell(row, 3).value).startswith("'=")
        )
        for sheet, row in ((sources, source_row), (formula_map, formula_row)):
            self.assertTrue(sheet.cell(row, 2).alignment.wrap_text)
            self.assertEqual(sheet.cell(row, 2).alignment.vertical, "top")
        self.assertTrue(any(
            (formula_map.row_dimensions[row].height or 0) >= 42
            for row in range(5, formula_map.max_row + 1)
        ))

    def test_without_calendar_binding_does_not_guess_annual_columns(self):
        from openpyxl import load_workbook

        result = export_fund_workbook(
            self.path, model=self.model, spec=self.specification, inputs=self.inputs)
        book = load_workbook(self.path, data_only=False)
        self.assertFalse(any(str(cell.value).startswith("FY")
                             for cell in book["Financials"][1]))
        self.assertIn("annual columns unavailable: no bound fiscal calendar",
                      result["gaps"])

    def test_tampered_calendar_binding_is_refused(self):
        calendar = self.calendar()
        calendar["fiscal_year_end_month"] = 12
        with self.assertRaisesRegex(FundWorkbookExportError, "calendar binding content hash"):
            export_fund_workbook(
                self.path, model=self.model, spec=self.specification,
                inputs=self.inputs, calendar_binding=calendar)
        self.assertFalse(self.path.exists())

    def test_changed_assumption_recalculates_quarter_annual_and_income_chain(self):
        soffice = shutil.which("soffice")
        if not soffice:
            self.skipTest("LibreOffice is unavailable")
        from openpyxl import load_workbook

        self.export(scenario=self.scenario())
        book = load_workbook(self.path, data_only=False)
        revenue_assumption = next(
            row for row in range(3, book["Driver"].max_row + 1)
            if "Revenues — Quarterly growth (ratio)"
            in str(book["Driver"].cell(row, 3).value))
        headers = {cell.value: cell.column for cell in book["Driver"][1]}
        book["Driver"].cell(revenue_assumption, headers["4Q26E"], 0.20)
        changed = Path(self.temp.name) / "changed" / "changed.xlsx"
        changed.parent.mkdir()
        book.save(changed)
        out = Path(self.temp.name) / "recalculated"
        out.mkdir()
        profile = Path(self.temp.name) / "lo-profile"
        subprocess.run([
            soffice, "--headless", f"-env:UserInstallation=file://{profile}",
            "--convert-to", "xlsx", "--outdir", str(out), str(changed),
        ], check=True, capture_output=True, text=True)
        calculated = load_workbook(out / "changed.xlsx", data_only=True)
        headers = {cell.value: cell.column for cell in calculated["Financials"][1]}
        revenue = calculated["Financials"].cell(3, headers["4Q26E"]).value
        operating_income = calculated["Financials"].cell(7, headers["4Q26E"]).value
        self.assertAlmostEqual(revenue, 1597.2, places=2)
        self.assertAlmostEqual(operating_income, 159.72, places=2)
        self.assertIsNotNone(calculated["Financials"].cell(3, headers["2027E"]).value)
        discounted_row = next(
            row for row in range(1, calculated["Valuation"].max_row + 1)
            if calculated["Valuation"].cell(row, 1).value == "Discounted target price"
        )
        self.assertAlmostEqual(calculated["Valuation"].cell(discounted_row, 2).value,
                               149.161194909091, places=5)

    def test_tampered_valuation_scenario_is_refused(self):
        scenario = self.scenario()
        scenario["multiple"] = "20"
        with self.assertRaisesRegex(FundWorkbookExportError, "content hash"):
            self.export(scenario=scenario)
        self.assertFalse(self.path.exists())

    def test_wrong_spec_or_input_hash_is_refused_before_file_creation(self):
        wrong_spec = dict(self.specification)
        wrong_spec["content_hash"] = "0" * 64
        with self.assertRaisesRegex(FundWorkbookExportError, "specification"):
            export_fund_workbook(self.path, model=self.model, spec=wrong_spec,
                                 inputs=self.inputs)
        self.assertFalse(self.path.exists())
        wrong_inputs = dict(self.inputs)
        wrong_inputs["company_ref"] = "company:wrong"
        with self.assertRaisesRegex(FundWorkbookExportError, "model inputs"):
            export_fund_workbook(self.path, model=self.model,
                                 spec=self.specification, inputs=wrong_inputs)
        self.assertFalse(self.path.exists())

    def test_existing_output_is_not_overwritten(self):
        self.path.write_bytes(b"original")
        with self.assertRaisesRegex(FundWorkbookExportError, "refusing to overwrite"):
            self.export()
        self.assertEqual(self.path.read_bytes(), b"original")

    def test_core_database_cannot_be_selected_as_output(self):
        core = Path(self.temp.name) / "core.sqlite"
        core.write_bytes(b"sqlite authority bytes")
        with self.assertRaisesRegex(FundWorkbookExportError, "Core database"):
            export_company_workbook(core, self.model["company_ref"], core)
        self.assertEqual(core.read_bytes(), b"sqlite authority bytes")

    def test_annual_filing_hash_replays_both_statement_wire_versions(self):
        connection = sqlite3.connect(":memory:")
        self.addCleanup(connection.close)
        connection.row_factory = sqlite3.Row
        connection.execute(
            "CREATE TABLE coverage_mission_statement_lines("
            "line_id TEXT,ingest_id TEXT,statement TEXT,ordinal INTEGER,"
            "concept TEXT,label TEXT,level INTEGER,parent_concept TEXT,"
            "is_breakdown INTEGER,dimension_axis TEXT,dimension_member TEXT,"
            "dimension_count INTEGER,period_start TEXT,period_end TEXT,"
            "value TEXT,unit TEXT,balance TEXT)")
        identity = {
            "company_ref": "company:test", "cik": "0000000001",
            "accession": "0000000001-25-000001", "form": "10-K",
            "line_count": 1,
        }
        ingest_id = f"statement-ingest:{content_hash(identity)[:32]}"
        stored = {
            "line_id": f"{ingest_id}#0", "ingest_id": ingest_id,
            "statement": "cash", "ordinal": 0, "concept": "us-gaap:Cash",
            "label": "Cash flow", "level": 0, "parent_concept": None,
            "is_breakdown": 0, "dimension_axis": None, "dimension_member": None,
            "dimension_count": 0, "period_start": "2025-01-01",
            "period_end": "2025-12-31", "value": "100", "unit": "usd",
            "balance": None,
        }
        connection.execute(
            "INSERT INTO coverage_mission_statement_lines VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            tuple(stored.values()))
        common = {key: stored[key] for key in (
            "statement", "concept", "label", "level", "parent_concept",
            "is_breakdown", "dimension_axis", "dimension_member", "period_start",
            "period_end", "value", "unit", "balance")}
        common["is_breakdown"] = False
        filing = {
            "ingest_id": stored["ingest_id"], **identity, "entity_name": "Test",
            "filed": "2026-01-01", "report_date": "2025-12-31",
            "source_record_refs": ["raw-sink:" + "1" * 64],
            "governance_ref": "governance:test", "governance_hash": "2" * 64,
        }
        for include_dimension_count in (False, True):
            line = dict(common)
            if include_dimension_count:
                line["dimension_count"] = 0
            filing["content_hash"] = content_hash({
                **{key: value for key, value in filing.items()
                   if key not in ("ingest_id", "content_hash")},
                "statement_lines_hash": content_hash([line]),
            })
            _verify_statement_filing(connection, filing)
        filing["report_date"] = "2025-11-30"
        with self.assertRaisesRegex(FundWorkbookExportError, "authority hash"):
            _verify_statement_filing(connection, filing)

    def test_annual_filing_rejects_forged_ingest_and_line_identities(self):
        connection = sqlite3.connect(":memory:")
        self.addCleanup(connection.close)
        connection.row_factory = sqlite3.Row
        connection.execute(
            "CREATE TABLE coverage_mission_statement_lines("
            "line_id TEXT,ingest_id TEXT,statement TEXT,ordinal INTEGER,"
            "concept TEXT,label TEXT,level INTEGER,parent_concept TEXT,"
            "is_breakdown INTEGER,dimension_axis TEXT,dimension_member TEXT,"
            "dimension_count INTEGER,period_start TEXT,period_end TEXT,"
            "value TEXT,unit TEXT,balance TEXT)")
        identity = {"company_ref": "company:test", "cik": "0000000001",
                    "accession": "0000000001-25-000001", "form": "10-K",
                    "line_count": 2}
        expected = f"statement-ingest:{content_hash(identity)[:32]}"
        base = ["cash", "us-gaap:Cash", "Cash flow", 0, None, 0, None, None,
                0, "2025-01-01", "2025-12-31", "100", "usd", None]
        for ordinal in range(2):
            connection.execute(
                "INSERT INTO coverage_mission_statement_lines VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (f"{expected}#{ordinal}", expected, base[0], ordinal, *base[1:]))
        lines = []
        for _ in range(2):
            lines.append({
                "statement": "cash", "concept": "us-gaap:Cash", "label": "Cash flow",
                "level": 0, "parent_concept": None, "is_breakdown": False,
                "dimension_axis": None, "dimension_member": None,
                "period_start": "2025-01-01", "period_end": "2025-12-31",
                "value": "100", "unit": "usd", "balance": None,
            })
        body = {**identity, "entity_name": "Test", "filed": "2026-01-01",
                "report_date": "2025-12-31",
                "source_record_refs": ["raw-sink:" + "1" * 64],
                "governance_ref": "governance:test", "governance_hash": "2" * 64}
        filing = {"ingest_id": expected, **body,
                  "content_hash": content_hash({**body,
                                                "statement_lines_hash": content_hash(lines)})}
        _verify_statement_filing(connection, filing)

        forged = "statement-ingest:" + "f" * 32
        connection.execute("UPDATE coverage_mission_statement_lines SET ingest_id=?,line_id=replace(line_id,?,?)",
                           (forged, expected, forged))
        with self.assertRaisesRegex(FundWorkbookExportError, "ingest identity"):
            _verify_statement_filing(connection, {**filing, "ingest_id": forged})
        connection.execute("UPDATE coverage_mission_statement_lines SET ingest_id=?,line_id=replace(line_id,?,?)",
                           (expected, forged, expected))

        connection.execute("UPDATE coverage_mission_statement_lines SET ordinal=7,line_id=? WHERE ordinal=1",
                           (f"{expected}#7",))
        with self.assertRaisesRegex(FundWorkbookExportError, "line identity"):
            _verify_statement_filing(connection, filing)
        connection.execute("UPDATE coverage_mission_statement_lines SET ordinal=1,line_id=? WHERE ordinal=7",
                           (f"{expected}#wrong",))
        with self.assertRaisesRegex(FundWorkbookExportError, "line identity"):
            _verify_statement_filing(connection, filing)

    def test_closed_cash_formulas_translate_without_display_label_inference(self):
        period = "2027-08-31"
        result_cells = {
            ("result:revenue", period): "'Financials'!B5",
            ("result:operating_cash_flow", period): "'Financials'!B12",
            ("result:capital_expenditure", period): "'Financials'!B13",
        }
        assumptions = {"assumption:cash-share": "'Driver'!B20"}
        share_cell = {
            "result_refs": [{"ref": "result:revenue", "period_end": period}],
            "assumption_refs": ["assumption:cash-share"],
        }
        ocf = {"ref": "result:operating_cash_flow", "label": "Localized label",
               "formula": "operating_cash_flow[k] = revenue[k] * share[k]"}
        self.assertEqual(
            _formula_for(ocf, share_cell, period, result_cells, assumptions, {}),
            "='Financials'!B5*'Driver'!B20",
        )
        fcf = {"ref": "result:free_cash_flow", "label": "Localized label",
               "formula": (
                   "free_cash_flow[k] = operating_cash_flow[k] "
                   "- capital_expenditure[k]")}
        fcf_cell = {
            "result_refs": [
                {"ref": "result:operating_cash_flow", "period_end": period},
                {"ref": "result:capital_expenditure", "period_end": period},
            ],
            "assumption_refs": [],
        }
        self.assertEqual(
            _formula_for(fcf, fcf_cell, period, result_cells, assumptions, {}),
            "='Financials'!B12-SUM('Financials'!B13)",
        )


if __name__ == "__main__":
    unittest.main()
