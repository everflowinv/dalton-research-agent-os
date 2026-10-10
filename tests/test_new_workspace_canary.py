"""Canary: a workspace created by the formal flow runs every lane, and keeps running.

The legacy Core reached a working governance and lane setup over eighteen
policy versions of hand edits; ws-7d, created through the setup flow, met the
missing pieces one refusal at a time.  This test is the rehearsal that would
have caught them: it creates a workspace exactly as the cockpit does, for a
mission that has nothing to do with US IT services (two fictional robotics
issuers, ``company:ticker:`` refs), and without signing anything by hand runs

    discovery -> extraction -> claim-support check -> ledger commit
    -> SEC company facts (10-Q, the 10-K pair, and FY - 9M for a 10-K of
       fiscal-year totals) -> claim index -> dossier -> event judgement,

then upgrades the mission through the real policy -> constitution -> mission
cascade and runs them again.  After the upgrade nothing may stall (an open
review left under the old version must reach the extraction queue) and nothing
may be read twice -- including when a search under the intermediate version
found the open document again and a second version was published over it
before any carry ran (the Guidepoint re-search shape, 2026-09-28).

A second case (2026-09-28) retires two Claims of the created workspace on a
support verdict "about another subject" and runs the industry reattribution
over them: an industry the rule has no vocabulary for, and a covered company,
are both refused, remembered, and nothing is reattributed.

Hermetic: every model is a fake, every connector a fake handle or a recorded
body, the SEC ticker resolver a fixture, and opening a socket fails the test.
"""

from __future__ import annotations

import json
import os
import socket
import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from dalton_core.store import DaltonStore, canonical_json, content_hash

OWNER = "human:canary-owner"
AUTOMATION = "automation:coverage-mission"
WDGT = "company:ticker:wdgt"
GZMO = "company:ticker:gzmo"
ISSUERS = {
    "WDGT": {"ticker": "WDGT", "cik": "9900001", "name": "Widget Robotics Inc."},
    "GZMO": {"ticker": "GZMO", "cik": "9900002", "name": "Gizmo Motion Corp."},
}
# The foundation's ceilings as the live setup writes them (ws-7d), including
# the probe ceiling the closed research-budget checks refuse.
BUDGET_CEILINGS = {"max_alphaengine_calls_24h": 50, "max_alphaengine_probe_calls_24h": 10,
                   "max_daily_cost_usd": 100.0, "max_daily_paid_calls": 100}


def _no_network(*args, **kwargs):  # pragma: no cover - only reached on a regression
    raise AssertionError(f"the canary opened a network connection: {args!r}")


def _write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, sort_keys=True), encoding="utf-8")


SEC_ACCESSION = "0009900001-26-000101"
# WDGT's 10-K for calendar 2025.  Like Accenture's, it reports its fourth
# quarter and the same quarter a year earlier as quarters, so the annual rule
# (COMPANY_FACTS_RULE_REFS["10-K"]) can answer it.
SEC_ANNUAL_ACCESSION = "0009900001-26-000050"


def annual_rows(accession: str) -> list[dict]:
    return [
        {"start": "2025-01-01", "end": "2025-12-31", "val": 440000000000, "accn": accession,
         "fy": 2025, "fp": "FY", "form": "10-K", "filed": "2026-02-10", "frame": "CY2025"},
        {"start": "2024-10-01", "end": "2024-12-31", "val": 104000000000, "accn": accession,
         "fy": 2025, "fp": "FY", "form": "10-K", "filed": "2026-02-10", "frame": "CY2024Q4"},
        {"start": "2025-10-01", "end": "2025-12-31", "val": 117000000000, "accn": accession,
         "fy": 2025, "fp": "FY", "form": "10-K", "filed": "2026-02-10", "frame": "CY2025Q4"},
    ]
# GZMO's 10-K for calendar 2025 reports fiscal years only, like AMZN's.  Its
# fourth quarter is derived as fiscal year less nine months from the filed
# statement rows of the 10-K and the year's three 10-Qs (FY - 9M), in millions:
# 2025: 414,000 - (100,000 + 102,000 + 104,000) = 108,000
# 2024: 372,000 - ( 90,000 +  92,000 +  94,000) =  96,000  -> +12.5%
GZMO_ANNUAL_ACCESSION = "0009900002-26-000050"
MILLION = 1_000_000


def annual_only_rows(accession: str) -> list[dict]:
    return [
        {"start": start, "end": end, "val": value * MILLION, "accn": accession,
         "fy": 2025, "fp": "FY", "form": "10-K", "filed": "2026-02-12", "frame": frame}
        for start, end, value, frame in (("2025-01-01", "2025-12-31", 414000, "CY2025"),
                                         ("2024-01-01", "2024-12-31", 372000, "CY2024"))]


def gzmo_statement_filings() -> dict[str, tuple[str, str, str, list[dict]]]:
    def line(start, end, value):
        return {"statement": "income",
                "concept": "us-gaap:RevenueFromContractWithCustomerExcludingAssessedTax",
                "label": "Revenue", "level": 0, "parent_concept": None, "is_breakdown": False,
                "dimension_axis": None, "dimension_member": None, "period_start": start,
                "period_end": end, "value": str(value * MILLION), "unit": "usd",
                "balance": "credit"}

    return {
        "0009900002-25-000011": ("10-Q", "2025-05-01", "2025-03-31", [
            line("2024-01-01", "2024-03-31", 90000), line("2025-01-01", "2025-03-31", 100000)]),
        "0009900002-25-000022": ("10-Q", "2025-08-01", "2025-06-30", [
            line("2024-04-01", "2024-06-30", 92000), line("2024-01-01", "2024-06-30", 182000),
            line("2025-04-01", "2025-06-30", 102000), line("2025-01-01", "2025-06-30", 202000)]),
        "0009900002-25-000033": ("10-Q", "2025-10-30", "2025-09-30", [
            line("2024-07-01", "2024-09-30", 94000), line("2024-01-01", "2024-09-30", 276000),
            line("2025-07-01", "2025-09-30", 104000), line("2025-01-01", "2025-09-30", 306000)]),
        GZMO_ANNUAL_ACCESSION: ("10-K", "2026-02-12", "2025-12-31", [
            line("2024-01-01", "2024-12-31", 372000), line("2025-01-01", "2025-12-31", 414000)]),
    }


DOC_A = "alphaengine-doc:990000000000001"
DOC_B = "alphaengine-doc:990000000000002"


def document_text(document_ref: str) -> str:
    # Both documents carry the same words, so one fixture quote serves both.
    del document_ref
    return ("HERMETIC FIXTURE ONLY. Widget Robotics Inc. (WDGT) management says factory "
            "automation orders remain strong; this is not a real research result.\n") * 40


class Selection:
    """The discovery-selection child, answering "read every candidate"."""

    def __init__(self) -> None:
        self.started: list[str] = []
        self.consumed: list[str] = []

    def currently_consumed(self, **kwargs):
        return []

    def current_selections(self, **kwargs):
        return {}

    def start(self, *, discovery_ref, view, **kwargs):
        self.started.append(discovery_ref)
        return {"id": f"discovery-selection:{len(self.started):024x}", "status": "succeeded",
                "discovery_ref": discovery_ref, "summary": {"status": "succeeded"},
                "effective_selected": [item["document_ref"] for item in view["candidates"]]}

    def mark_consumed(self, ticket) -> None:
        self.consumed.append(ticket["id"])


class Acquisitions:
    """The acquisition child: a real governed acquisition on a fake page handle.

    It also leaves the ticket directory the review plane reads a manifest from,
    exactly as the launched child does.
    """

    def __init__(self, harness, state: Path) -> None:
        self.h = harness
        self.state = state
        self.calls: list[str] = []
        self.tickets: dict[str, dict] = {}

    def start_bounded_probe(self, *, document_ref, caller_ref, max_pages=20):
        self.calls.append(document_ref)
        ticket = f"alphaengine-acquisition:{len(self.calls):024x}"
        self.tickets[ticket] = {"id": ticket, "status": "running", "document_ref": document_ref}
        return dict(self.tickets[ticket])

    def finish(self) -> None:
        for ticket in self.tickets.values():
            if ticket["status"] != "running":
                continue
            self.h.document_handle.text = document_text(ticket["document_ref"])
            self.h.document_handle.page_chars = 7000
            manifest = self.h.acquisition.acquire(
                self.h.acquisition.build_plan(ticket["document_ref"]))["manifest"]
            directory = self.state / "acquisitions" / ticket["id"].split(":")[1]
            directory.mkdir(parents=True, mode=0o700, exist_ok=True)
            for name, value in {
                "ticket.json": {"id": ticket["id"], "status": "succeeded",
                                "document_ref": ticket["document_ref"]},
                "summary.json": {"document_ref": ticket["document_ref"],
                                 "manifest_ref": manifest["id"],
                                 "manifest_hash": manifest["content_hash"],
                                 "manifest_status": "complete",
                                 "assembled_content_sha256": manifest["declared_content_sha256"]},
                "manifest.json": manifest,
            }.items():
                (directory / name).write_text(canonical_json(value)); (directory / name).chmod(0o600)
            ticket["status"] = "succeeded"

    def status(self, ticket_ref):
        return dict(self.tickets[ticket_ref])


class NewWorkspaceCanaryTests(unittest.TestCase):
    """One workspace, the whole research loop, a mission upgrade, the loop again."""

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(dir="/tmp")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        environ = patch.dict(os.environ, {}, clear=False)
        environ.start(); self.addCleanup(environ.stop)
        os.environ.pop("DALTON_WORKSPACE_MANIFEST", None)
        for target in ("socket.create_connection", "socket.socket.connect",
                       "urllib.request.urlopen", "http.client.HTTPConnection.connect"):
            guard = patch(target, _no_network)
            guard.start(); self.addCleanup(guard.stop)

    # -- stage 0: the formal creation flow --------------------------------------

    def create_workspace(self):
        from dalton_core.workspace import load_workspace_manifest
        from dalton_core.workspace_creation import create_blank_workspace
        from dalton_core.workspace_mission_setup import (
            draft_first_mission, publish_first_mission_to_store,
        )
        from dalton_core.workspace_runtime_setup import install

        release = self.root / "release"; release.mkdir()
        receipt = create_blank_workspace(
            self.root / "fleet", "canary", 18931, "release:sha256:" + "c" * 64, release,
            request_id="canary-request", display_name="Canary robotics",
            shared_readonly_paths=(release,))
        install(receipt["manifest_path"], actor_ref=OWNER, stage_host_lanes=False)
        workspace = load_workspace_manifest(receipt["manifest_path"])
        state = workspace.state_dir
        foundation = json.loads((state / "research-foundation.json").read_text())
        body = {key: value for key, value in foundation.items() if key != "content_hash"}
        body["mission_defaults"]["source_plan"] = [
            {"source_ref": ref, "role": "shared connected evidence source", "status": "connected"}
            for ref in ("source:alphaengine", "source:sales-notes", "source:sec-edgar")]
        body["mission_defaults"]["budget_ceilings"] = dict(BUDGET_CEILINGS)
        foundation = {**body, "content_hash": content_hash(body)}
        _write_json(state / "research-foundation.json", foundation)
        proposal = draft_first_mission(
            workspace, goal="Research the industrial robotics industry: $WDGT and $GZMO",
            method_foundation=foundation, industry="Industrial Robotics",
            companies=[{"company_ref": WDGT, "ticker": "WDGT"},
                       {"company_ref": GZMO, "ticker": "GZMO"}])
        self.assertEqual(proposal["setup_state"], "ready_for_confirmation", proposal["review_issues"])
        resolved: list[str] = []

        def resolver(ticker: str) -> dict[str, str]:
            resolved.append(ticker)
            return dict(ISSUERS[ticker])

        with DaltonStore(state / "core.sqlite") as store:
            mission = publish_first_mission_to_store(
                store, workspace, proposal=proposal, proposal_hash=proposal["content_hash"],
                actor_ref=OWNER, method_foundation=foundation, sec_ticker_resolver=resolver)
        self.assertEqual(sorted(resolved), ["GZMO", "WDGT"])
        return workspace, mission

    def assert_governance_baseline(self, workspace, mission) -> None:
        from dalton_core.coverage_mission import research_budget_shape_valid
        from dalton_core.research_auto_commit import policy_lists_document_rule
        from dalton_core.sec_company_facts_lane import check_core_governance_rules
        from dalton_core.workspace_health_parity import Environment, check_environment

        state = workspace.state_dir
        with DaltonStore(state / "core.sqlite") as store:
            check_core_governance_rules(store)  # ws-7d's policy-4 failed exactly here
            active = store.active_policy_version().to_dict()
            self.assertTrue(policy_lists_document_rule(active))
            # FY - 9M is signed with the rest of the baseline, by the flow.
            from dalton_core.research_auto_commit import SEC_FY_MINUS_9M_RULE_REF
            self.assertIn(SEC_FY_MINUS_9M_RULE_REF,
                          active["policy"]["research_candidate_auto_commit"]["rules"])
            self.assertEqual(active["policy"]["research_budget"], {
                "max_alphaengine_calls_24h": 50, "max_daily_cost_usd": 100.0,
                "max_daily_paid_calls": 100})
            self.assertTrue(active["independence_predicates"])
            from dalton_core.agenda import AgendaStore
            mandate = AgendaStore(store).mandate_version(mission["bindings"]["mandate_version"]["ref"])
        cap = mandate["constraints"]["research_budget"]
        self.assertTrue(research_budget_shape_valid(cap), cap)
        self.assertNotIn("max_alphaengine_probe_calls_24h", cap)
        env = Environment("canary", state, service_config=workspace.config_path)
        try:
            report = check_environment(env, include_host=False)
        finally:
            env.close()
        gaps = [row for row in report["rows"] if row["status"] == "gap"
                and row["section"] in {"governance", "authorization", "mission"}]
        self.assertEqual(gaps, [], json.dumps(gaps, indent=1, ensure_ascii=False))

    # -- stage 1: discovery ---------------------------------------------------------

    def discover(self, state: Path):
        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.macos_launchagent import _alphaengine_discovery_plan
        from dalton_core.mission_source_discovery import (
            MissionSourceDiscoveryCoordinator, load_discovery_plan,
        )
        from tests.test_mission_source_discovery import FakeSearchLauncher, SearchHarness

        harness = SearchHarness(state, [{"doc_id": DOC_A.split(":")[1]},
                                        {"doc_id": DOC_B.split(":")[1]}])
        self.addCleanup(harness.close)
        self.harness = harness
        self.missions = CoverageMissionAuthority(harness.core)
        # 2026-10-06: the approved file holds governed fields only (0.7); the
        # writer renders it with the code-owned forms, exactly as here.
        from dalton_core.writer_server import _rendered_alphaengine_plan

        approved = load_discovery_plan(_alphaengine_discovery_plan(state))
        self.assertEqual(approved["schema_version"], "0.7")
        self.assertTrue(all("query_variants" not in spec for spec in approved["specs"]))
        plan, query_forms = _rendered_alphaengine_plan(_alphaengine_discovery_plan(state))
        self.assertEqual(set(plan["companies"]), {WDGT, GZMO})
        self.search = FakeSearchLauncher(harness, self.missions, plan)
        self.acquisitions = Acquisitions(harness, state)
        self.selection = Selection()
        self.coordinator = MissionSourceDiscoveryCoordinator(
            store=harness.core, missions=self.missions, plan=plan,
            search_launcher=self.search, acquisition_launcher=self.acquisitions,
            selection_launcher=self.selection, spool_dir=state / "spool",
            clock=harness.clock, query_forms=query_forms)
        ticks = []
        for _ in range(12):
            ticks.append(self.coordinator.dispatch_once())
            self.acquisitions.finish()
        self.assert_first_search_is_current(plan, ticks)
        mission = self.missions.active_mission(plan["mission_ref"])
        reviews = self.missions.document_reviews(mission["id"])
        self.assertEqual(sorted(r["document_ref"] for r in reviews), [DOC_A, DOC_B],
                         json.dumps(ticks, indent=1, default=str)[-3000:])
        self.assertEqual(sorted(self.acquisitions.calls), [DOC_A, DOC_B])
        return mission, reviews

    def assert_first_search_is_current(self, plan, ticks) -> None:
        """2026-10-04: a new workspace searches on its first tick, in the current form.

        ws-7d approved the "{name} {quarter_label}" transcript query and then
        sat sixteen hours without trying it; a brand-new workspace must render
        the form the setup generator writes today and search on the first
        tick where a search is due, and its ledger must say where the tick
        went.
        """

        from dalton_core.mission_source_discovery import (
            build_discovery_parameters, discovery_coordinator_state, quarter_title_label,
        )
        from dalton_core.alphaengine_query_forms import code_query_variants

        for spec in plan["specs"]:
            self.assertEqual(spec["query_variants"],
                             code_query_variants(spec["spec_ref"], spec["document_type"]))
        self.assertEqual(plan["specs"][0]["spec_ref"], "earnings-call-transcripts")
        self.assertEqual(plan["specs"][0]["query_variants"][0]["query_template"],
                         "{name} {quarter_label}")
        first = ticks[0]
        self.assertEqual(first["discovery"]["status"], "launched",
                         json.dumps(first["discovery"], default=str)[:2000])
        start = self.search.starts[0]
        self.assertEqual((start["spec_ref"], start["variant_index"], start["cursor"]),
                         ("earnings-call-transcripts", 0, None))
        company_ref = start["authorization"]["company_ref"]
        compiled = build_discovery_parameters(
            plan, spec_ref="earnings-call-transcripts", company_ref=company_ref,
            as_of=start["as_of"], variant_index=0, missing_periods=start["missing_periods"])
        name = plan["companies"][company_ref]["name"]
        expected = " ".join(
            [name, *([quarter_title_label(start["missing_periods"][0])]
                     if start["missing_periods"] else [])])
        self.assertEqual(compiled["query"], expected)
        self.assertNotIn("earnings call transcript", compiled["query"])
        state = discovery_coordinator_state(first)
        for step in ("settle", "carry", "held", "acquire", "search"):
            self.assertIn(f"{step}_seconds", state)
        self.assertNotIn("search_blocked_reason", state)
        # The first tick's ledger says which form set the plan renders, once.
        self.assertEqual(state["query_forms"], "code")
        self.assertEqual(state["query_forms_hash"], plan["content_hash"])
        self.assertNotIn("ignored", state["query_forms_note"])
        self.assertNotIn("query_forms", discovery_coordinator_state(ticks[1]))

    # -- stages 2-4: extraction, claim-support check, ledger commit ------------------

    def prepare_extraction(self, state: Path) -> None:
        from dalton_core.extraction_backlog import DocumentProvenanceStore
        from dalton_core.model_router import ModelRouter
        from dalton_core.writer_server import HUMAN_GOVERNANCE_OPERATIONS, Principal, WriterServer
        from dalton_core.alphaengine_acquisition_launcher import AlphaEngineAcquisitionLauncher
        from tests.test_transcript_polish_model_worker import policy, profile

        h = self.harness
        provenance = DocumentProvenanceStore(h.core.connection, clock=h.clock)
        for document_ref in (DOC_A, DOC_B):
            provenance.record({
                "document_ref": document_ref, "source_ref": "source:alphaengine",
                "spec_ref": "earnings-call-transcripts", "provenance_tier": "management",
                "broker": None, "broker_key": "", "authors": None, "sources": None,
                "title": "Widget Robotics Inc. Q2 2026 Earnings Call Transcript",
                "named_companies": ["Widget Robotics Inc."], "published_at": None,
                "metadata_seen": True})
        router = ModelRouter(str(state / "canary-router.sqlite"), clock=h.clock)
        self.addCleanup(router.close)
        pr = profile()
        pr["provider"] = "hermetic-fixture"
        pr["cost"]["input_per_million_usd"] = pr["cost"]["output_per_million_usd"] = 0
        pr["availability"]["checked_at"] = h.clock().isoformat()
        pr["availability"]["valid_until"] = (h.clock() + timedelta(days=2)).isoformat()
        router.register_profile(pr); router.register_policy(policy())
        self.model_config = state / "canary-extraction-model-config.json"
        _write_json(self.model_config, {
            "routing_policy_ref": policy()["policy_version_ref"],
            "credential_slot_refs": [profile()["credential_slot_ref"]],
            "model_router_db": str(state / "canary-router.sqlite"),
            "broker_socket": str(state / "none.sock"), "broker_auth_key": str(state / "none.key"),
            "broker_client_id": "client:dalton-core", "expected_agent_id": "chem",
            "budget_db": str(state / "canary-budget.sqlite"),
            "budget_policy_ref": "thesis-impact-day-budget-policy:production:1"})
        launcher = AlphaEngineAcquisitionLauncher(state_dir=state,
                                                  governance_path=state / "unused-governance.json")
        writer = WriterServer(state / "unused.sqlite", state / "unused.sock", {
            "human": Principal("human", "fixture-token", HUMAN_GOVERNANCE_OPERATIONS, actor_ref=OWNER),
        }, acquisition_launcher=launcher)
        writer._store = h.core; writer._connectors = h.connectors
        writer._observability = h.observability; writer._coverage_mission = self.missions
        writer._transcript_spool = h.spool; writer._scheduler = h.scheduler
        self.writer = writer

    def quote_of(self, review) -> str:
        from dalton_core.document_extraction import DocumentExtractionService

        context = DocumentExtractionService(self.writer).view(
            review_id=review["review_id"], expected_review_hash=content_hash(review),
            offset=0, actor_ref=OWNER)["context"]
        return context["quotes"][0]["quote_id"]

    def extract(self, state: Path, *, quote_id: str, run: str, max_windows: int = 1):
        from dalton_core import document_extraction_cli
        from dalton_core.claim_support_verification import ClaimSupportVerifier
        from dalton_core.document_extraction_cli import run_extraction
        from tests.test_claim_support_verification import FakeModel, _reply

        fixture = state / f"canary-fixture-{run}.json"
        _write_json(fixture, {"schema_version": "0.1", "suggestions": [{
            "quote_id": quote_id,
            "normalized_statement": "Widget Robotics management says factory automation orders remain strong.",
            "metric_or_aspect": "aspect:order-intake", "period": "Q2 2026",
            "basis": "fixture management commentary"}]})
        verifier = FakeModel(*([_reply(("supported", "about_subject", None))] * 4))
        base = document_extraction_cli.ExtractionHost

        class Host(base):
            def __init__(inner, **kwargs):
                super().__init__(**kwargs)
                inner._claim_support_verifier = ClaimSupportVerifier(
                    store=inner.store, model_call=verifier, daily_cap_micros=100_000,
                    producer_family=lambda ref: "deepseek-v4", max_attempts=3)

        with patch.object(document_extraction_cli, "ExtractionHost", Host):
            summary = run_extraction(
                state_dir=state, model_config_path=self.model_config,
                summary_dir=state / "extractions" / f"canary-{run}", spool_dir=state / "spool",
                scheduler_db=state / "scheduler.sqlite", requested_by=None,
                max_windows=max_windows, max_numeric_windows=0, max_discovery_windows=0,
                connector_governance=None, web_fetch_governance=None,
                hermetic_fixture=fixture,
                candidate_staging=state / "research-review" / "candidate-staging.sqlite")
        return summary, verifier

    # -- stage 5: SEC company facts ----------------------------------------------------

    def sec_company_facts(self, state: Path, mission, *, run_key: str, ticker: str = "WDGT",
                          accession: str | None = None, day: int = 0, form: str = "10-Q",
                          annual_accession: str | None = None,
                          annual_only_accession: str | None = None):
        from dalton_core.sec_authority_harness import MutableClock
        from dalton_core.sec_company_facts_lane import Issuer, RehearsalGovernance, SecCompanyFactsLane
        from tests.test_research_plan_executor import _sec_company_facts_body
        from tests.test_sec_company_facts_lane import fake_adapter

        accession = accession or SEC_ACCESSION
        company_ref = {"WDGT": WDGT, "GZMO": GZMO}[ticker]
        body = json.loads(_sec_company_facts_body())
        body.update({"cik": int(ISSUERS[ticker]["cik"]),
                     "entityName": ISSUERS[ticker]["name"].upper()})
        rows = body["facts"]["us-gaap"][
            "RevenueFromContractWithCustomerExcludingAssessedTax"]["units"]["USD"]
        for fact in rows:
            fact["accn"] = accession
        if annual_accession is not None:
            rows.extend(annual_rows(annual_accession))
        if annual_only_accession is not None:
            rows.extend(annual_only_rows(annual_only_accession))
        clock = MutableClock()
        clock.advance(day * 86400)  # a new connector quota day
        issuer = Issuer(ticker, ISSUERS[ticker]["cik"], company_ref, ISSUERS[ticker]["name"])
        lane = SecCompanyFactsLane(
            state_dir=state, staging_path=state / "research-review" / "sec-candidate-staging.sqlite",
            governance=RehearsalGovernance(approved_by=OWNER), issuers=(issuer,),
            adapter=fake_adapter(clock, json.dumps(body, sort_keys=True).encode()), clock=clock)
        with lane:
            return lane.run_issuer(
                issuer, actor_ref=AUTOMATION, run_key=run_key,
                filed_from="2025-08-20", filed_to="2026-08-20", form=form,
                expected_accession=annual_accession if form == "10-K" else accession,
                mission_context={"mission_version_ref": mission["id"],
                                 "mission_version_hash": mission["content_hash"],
                                 "company_ref": company_ref})

    def assert_fourth_quarter_is_dispatched(self, state: Path, mission) -> None:
        from dalton_core.mission_sec_quarters import (
            MissionSecQuartersCoordinator,
            accession_in_hand,
        )

        def coordinator():
            return MissionSecQuartersCoordinator(
                store=self.harness.core, missions=self.missions, state_dir=state,
                checklist=lambda: [{"company_ref": WDGT, "ticker": "WDGT", "items": [
                    {"item_ref": "quarterly_financials", "have": 1, "required": 4}]}],
                clock=lambda: datetime(2026, 9, 28, tzinfo=timezone.utc))

        first = coordinator().dispatch_once()
        self.assertEqual(first["status"], "queued", json.dumps(first, indent=1, default=str))
        self.assertIn("2025-12-31", first["recent_quarters_missing"])
        [annual] = [q for q in first["queued"] if q["form"] == "10-K"]
        self.assertEqual((annual["accession"], annual["period"]),
                         (SEC_ANNUAL_ACCESSION, "2025-10-01..2025-12-31"))
        pending = {row["expected_accession"]: row
                   for row in self.missions.pending_sec_dispatches(limit=10)}
        self.assertEqual(pending[SEC_ANNUAL_ACCESSION]["form"], "10-K")
        claims = self.claim_count()
        q4 = self.sec_company_facts(state, mission, run_key="canary-q4", form="10-K",
                                    annual_accession=SEC_ANNUAL_ACCESSION, day=2)
        self.assertEqual(q4["status"], "committed", json.dumps(q4, indent=1, default=str)[-3000:])
        self.assertEqual((q4["form"], q4["candidate"]["period"]), ("10-K", "2025-10-01..2025-12-31"))
        self.assertEqual(q4["facts"]["growth_percent"], "12.50")
        self.assertEqual(self.claim_count(), claims + 1)
        # The drain's bookkeeping, as the writer does it after the lane ran.
        for accession, row in pending.items():
            ticket = f"sec-lane-run:canary-{accession[-6:]}"
            self.missions.mark_sec_dispatch_launched(row["dispatch_id"], ticket)
            self.missions.settle_sec_dispatch(
                row["dispatch_id"], outcome="finished", ticket_ref=ticket,
                detail="succeeded" if accession == SEC_ANNUAL_ACCESSION else "failed")
        # The bounded-planner path sees the accession answered ...
        self.assertEqual(accession_in_hand(self.harness.core.connection,
                                           SEC_ANNUAL_ACCESSION)["state"], "succeeded")
        # ... and the coordinator sees the quarter held.
        again = coordinator().dispatch_once()
        self.assertNotIn(SEC_ANNUAL_ACCESSION,
                         [q["accession"] for q in again.get("queued", [])], again)
        self.assertEqual(self.claim_count(), claims + 1)

    def ingest_statements(self, mission, company_ref: str, ticker: str, filings) -> None:
        """The financial-statements lane's record of these filings, as its child leaves it."""

        authorization = self.missions.authorize_sec_lane(
            company_ref=company_ref, ticker=ticker, actor_ref=AUTOMATION,
            mission_version_ref=mission["id"], mission_version_hash=mission["content_hash"])
        for form in ("10-Q", "10-K"):
            batch = [{"accession": accession, "form": kind, "filed": filed,
                      "report_date": report, "lines": lines}
                     for accession, (kind, filed, report, lines) in sorted(filings.items())
                     if kind == form]
            dispatch = self.missions.queue_statement_dispatch(authorization=authorization, form=form)
            self.missions.mark_statement_dispatch_launched(
                dispatch["dispatch_id"], f"sec-financials-run:canary-{form.lower()}")
            self.missions.record_statement_observation(
                dispatch_id=dispatch["dispatch_id"],
                observation={"schema_version": "0.1", "cik": ISSUERS[ticker]["cik"].zfill(10),
                             "entity_name": ISSUERS[ticker]["name"], "filings": batch,
                             "source_record_refs": ["raw-sink:" + ("9" if form == "10-K" else "8") * 64],
                             "next_cursor": None, "provider_status": 200},
                governance_ref="connector-governance:sec-financial-statements:v2",
                governance_hash="b" * 64)
            self.missions.settle_statement_dispatch(dispatch["dispatch_id"], outcome="succeeded")

    def assert_fourth_quarter_is_derived(self, state: Path, mission) -> None:
        """A 10-K of fiscal-year totals: the quarter lane derives Q4 as FY - 9M.

        No 10-K connector run (it could only fail) and no hand-signed rule: the
        workspace's baseline policy lists the FY - 9M rule, the statement lane's
        rows are in Core, and the coordinator stages the derivation and the
        Ledger admits it after rebuilding it from those rows.
        """

        from dalton_core.mission_sec_quarters import MissionSecQuartersCoordinator
        from dalton_core.research_auto_commit import SEC_FY_MINUS_9M_RULE_REF
        from dalton_core.research_verification import CandidateStagingStore

        self.ingest_statements(mission, GZMO, "GZMO", gzmo_statement_filings())
        staging = CandidateStagingStore(state / "research-review" / "candidate-staging.sqlite")
        self.addCleanup(staging.close)

        def coordinator():
            return MissionSecQuartersCoordinator(
                store=self.harness.core, missions=self.missions, state_dir=state,
                checklist=lambda: [{"company_ref": GZMO, "ticker": "GZMO", "items": [
                    {"item_ref": "quarterly_financials", "have": 1, "required": 4}]}],
                clock=lambda: datetime(2026, 9, 28, tzinfo=timezone.utc), staging=staging)

        claims = self.claim_count()
        first = coordinator().dispatch_once()
        self.assertIn(first["status"], {"queued", "committed"},
                      json.dumps(first, indent=1, default=str))
        [derived] = first["derived"]
        self.assertEqual(derived["status"], "committed", json.dumps(derived, indent=1))
        self.assertEqual((derived["accession"], derived["period"], derived["value"],
                          derived["rule_ref"]),
                         (GZMO_ANNUAL_ACCESSION, "2025-10-01..2025-12-31", "12.5",
                          SEC_FY_MINUS_9M_RULE_REF))
        self.assertEqual((derived["q4"], derived["prior_q4"]),
                         (str(108000 * MILLION), str(96000 * MILLION)))
        # The 10-K itself is never sent to the SEC lane.
        self.assertNotIn(GZMO_ANNUAL_ACCESSION, [q["accession"] for q in first["queued"]])
        self.assertNotIn(GZMO_ANNUAL_ACCESSION,
                         [row["expected_accession"]
                          for row in self.missions.pending_sec_dispatches(limit=20)])
        self.assertEqual(self.claim_count(), claims + 1)
        row = self.harness.core.connection.execute(
            "SELECT claim_json FROM claim_versions WHERE json_extract(claim_json,'$.subject_ref')=? "
            "AND json_extract(claim_json,'$.metric_or_aspect')='quarterly_revenue_yoy_growth' "
            "AND json_extract(claim_json,'$.period')='2025-10-01..2025-12-31'", (GZMO,)).fetchone()
        claim = json.loads(row[0])
        self.assertEqual(claim["basis"], "official-filing-xbrl-derived")
        self.assertIn("not a filed quarter", claim["normalized_statement"])
        # Held now: the next tick neither derives it again nor queues its 10-K.
        again = coordinator()
        self.assertIn("2025-10-01..2025-12-31", again._held_periods(GZMO))
        second = again.dispatch_once()
        self.assertFalse([item for item in second.get("derived", [])
                          if item["accession"] == GZMO_ANNUAL_ACCESSION], second)
        self.assertEqual(self.claim_count(), claims + 1)

    # -- stage 6: claim index and dossier -------------------------------------------

    def pass_screen(self, mission) -> None:
        from dalton_core.coverage_mission import CoverageMissionConflict

        for status in ("entered", "gate_passed"):
            if status == "entered":
                try:
                    self.missions.record_stage(
                        mission_version_ref=mission["id"],
                        mission_version_hash=mission["content_hash"],
                        company_ref=WDGT, stage_ref="initial_screen", status=status,
                        evidence_refs=[mission["id"]], rationale="canary: screened",
                        actor_ref=AUTOMATION, idempotency_key=f"canary:{mission['id']}:{status}")
                except CoverageMissionConflict:
                    pass  # the SEC lane's stage claim already entered it
                continue
            self.missions.record_stage(
                mission_version_ref=mission["id"], mission_version_hash=mission["content_hash"],
                company_ref=WDGT, stage_ref="initial_screen", status=status,
                evidence_refs=[mission["id"]], rationale="canary: screened",
                actor_ref=AUTOMATION, idempotency_key=f"canary:{mission['id']}:{status}")

    def index_and_dossier(self, state: Path, *, run: str):
        from dalton_core.claim_index_cli import run_claim_index
        from dalton_core.company_dossier_cli import run_dossier
        from tests.test_claim_index_lane import FakeModel as IndexModel
        from tests.test_dossier_lane import FakeModel as DossierModel, resolver

        config = state / "canary-empty-model-config.json"
        config.write_text("{}", encoding="utf-8")
        # One prose claim per batch; the numeric SEC claim is rule-tagged.
        index = []
        with patch("dalton_core.claim_index_cli.CockpitModel", IndexModel("1\tdemand_drivers\n")):
            for attempt in range(4):
                index.append(run_claim_index(
                    state_dir=state, model_config_path=config,
                    summary_dir=state / "claim-index-runs" / f"canary-{run}-{attempt}",
                    scheduler_db=None, company_ref=WDGT))
                if not index[-1].get("pending"):
                    break
        drafter = DossierModel()
        self.last_dossier_verifier = DossierModel(route="route:verify")
        dossier = run_dossier(
            state_dir=state, model_config_path=config,
            summary_dir=state / "company-dossier-runs" / f"canary-{run}",
            policy_path=state / "p12a-dossier-policy-v1.json",
            model_factory=lambda: drafter,
            verifier_model_factory=lambda: self.last_dossier_verifier,
            family_resolver=resolver)
        return index, dossier, drafter

    # -- stage 7: event judgement ------------------------------------------------------

    def record_news(self, mission, document_ref: str) -> dict:
        from dalton_core.research_event import ResearchEventAuthority, record_event

        occurred = (datetime.now(timezone.utc) - timedelta(days=1)).replace(microsecond=0)
        return record_event(
            ResearchEventAuthority(self.harness.core), company_ref=WDGT, kind="news",
            occurred_at=occurred.isoformat(),
            source_refs=["source:alphaengine", document_ref],
            payload={"document_ref": document_ref, "source_ref": "source:alphaengine",
                     "spec_ref": "earnings-call-transcripts", "discovery_ref": "d",
                     "title": None, "host": None},
            mission=mission, actor_ref=AUTOMATION)

    def judge(self, state: Path, *, run: str, triage=None):
        """One judgement run, with the cheap screen (event_triage) in front.

        ``triage`` is the screen's replies; by default it flags the group, so
        the judge sees it exactly as it did before the screen existed.
        """

        from dalton_core.event_judgement_cli import run_judgement
        from tests.test_event_judgement import (
            JUDGE_ROUTE, PASS, VERIFIER_ROUTE, FakeModel, decision, resolver,
        )
        from tests.test_event_triage import FLAG, TRIAGE_ROUTE

        judge_model = FakeModel([decision(), decision()], route=JUDGE_ROUTE)
        self.triage_model = FakeModel(list(triage if triage is not None else [FLAG, FLAG]),
                                      route=TRIAGE_ROUTE, cost_micros=1_500)
        summary = run_judgement(
            state_dir=state, summary_dir=state / "event-judgement-runs" / f"canary-{run}",
            policy_path=state / "tracking-policy.json", now=datetime.now(timezone.utc),
            judge_model=judge_model,
            verifier_model=FakeModel([PASS, PASS], route=VERIFIER_ROUTE),
            family_resolver=resolver(), triage_model=self.triage_model)
        return summary, judge_model

    def judged(self) -> int:
        from dalton_core.event_judgement import EventJudgementAuthority

        return EventJudgementAuthority(self.harness.core).judged_count(WDGT)

    # -- the upgrade ---------------------------------------------------------------------

    def upgrade_mission(self) -> dict:
        """Sign one more rule through the owner's real cascade, on this Core.

        The same policy -> constitution -> mission chain the owner scripts run
        (ws-7d's policy-5 was one), applied in place rather than on a copy.
        """

        from dalton_core.agenda import AgendaStore
        from dalton_core.research_constitution import ResearchConstitutionAuthority
        from dalton_core.research_plan import PLAN_AUTO_START_RULE_REF
        from scripts.publish_extraction_authority_chain import apply_chain
        from scripts.sign_auto_commit_rules import read_current
        from scripts.sign_research_plan_auto_start import build_chain

        store = self.harness.core
        current = read_current(store.connection)
        chain = build_chain(current, now=datetime.now(timezone.utc).isoformat(),
                            rules=[PLAN_AUTO_START_RULE_REF])
        agenda = AgendaStore(store)
        constitutions = ResearchConstitutionAuthority(store)

        def apply(operation, params):
            values = dict(params)
            if operation == "create_policy":
                return {"policy_version": store.create_policy(
                    values.pop("policy"), actor_ref=OWNER, **values)}
            if operation == "create_mandate":
                return agenda.create_mandate(values.pop("mandate_ref"), actor_ref=OWNER, **values)
            if operation == "publish_research_constitution":
                return constitutions.publish_constitution(
                    values.pop("constitution_ref"), actor_ref=OWNER, **values)
            if operation == "create_coverage_mission":
                return self.missions.create_mission(
                    values.pop("mission_ref"), actor_ref=OWNER, **values)
            raise AssertionError(operation)

        apply_chain(chain, apply)
        return self.missions.active_mission(current["mission_ref"])

    def republish_mission(self) -> dict:
        """A mission-only version over the one in force (P2's weekly-brief switch)."""

        from scripts.sign_research_plan_auto_start import BODY_FIELDS, _next_version_id

        current = self.missions.active_mission(self.missions.active_mission(
            self.search.plan["mission_ref"])["mission_ref"])
        number = int(current["version"]) + 1
        self.missions.create_mission(
            current["mission_ref"], actor_ref=OWNER,
            **{field: json.loads(json.dumps(current[field])) for field in BODY_FIELDS},
            version_id=_next_version_id(current["id"], number),
            prior_version_ref=current["id"],
            idempotency_key=f"{current['mission_ref']}:{number}:canary-republish")
        return self.missions.active_mission(current["mission_ref"])

    def research_without_reconciling(self, version) -> None:
        """One search under ``version`` recorded as a Guidepoint child records it.

        The row lands as the search saw it -- ``already_in_authority`` for a
        document Core holds -- and nothing settles it or opens a review.  The
        first version's search result is recorded again (the page a re-search
        would return), so no second governed call is needed.
        """

        [first] = self.missions.source_discoveries(
            version["id"], company_ref=WDGT, across_versions=True, limit=1000)[-1:]
        grant = self.missions.authorize_source_discovery(
            company_ref=WDGT, source_ref="source:alphaengine", requested_by=AUTOMATION,
            mission_version_ref=version["id"])
        self.missions.record_source_discovery(
            authorization=grant, discovery_plan_ref=first["discovery_plan_ref"],
            discovery_plan_hash=first["discovery_plan_hash"], spec_ref=first["spec_ref"],
            query_hash=first["query_hash"], parameters=first["parameters"],
            connector_invocation_ref=first["connector_invocation_ref"],
            connector_invocation_hash=first["connector_invocation_hash"],
            source_envelope_ref=first["source_envelope_ref"],
            source_envelope_hash=first["source_envelope_hash"],
            document_refs=first["document_refs"],
            in_authority_document_refs=first["document_refs"])

    def awaiting_in_active_version(self) -> int:
        # Exactly the extraction launcher's queue (``_awaiting_state``).
        return self.harness.core.connection.execute(
            "SELECT COUNT(*) FROM coverage_mission_document_reviews r "
            "JOIN coverage_mission_pointer p ON p.mission_version_id=r.mission_version_ref "
            "WHERE r.state='awaiting_human_extraction'").fetchone()[0]

    def claim_count(self) -> int:
        return self.harness.core.connection.execute(
            "SELECT COUNT(*) FROM claim_versions").fetchone()[0]

    def test_a_created_workspace_has_the_code_owned_sec_specs_and_the_release_ledger(self) -> None:
        import sqlite3

        from dalton_core.alphaengine_query_forms import effective_discovery_plan
        from dalton_core.macos_launchagent import _sec_discovery_plan
        from dalton_core.mission_source_discovery import load_discovery_plan
        from dalton_core.sec_plan_specs import SEC_PLAN_SPECS, sec_plan_report
        from dalton_core.workspace_health_parity import Environment, sec_plan_spec_rows

        workspace, _ = self.create_workspace()
        state = workspace.state_dir
        approved = load_discovery_plan(_sec_discovery_plan(state))
        # The file the generator wrote is already the code's set: no drift,
        # nothing re-rendered, so nothing re-searched for nothing.
        self.assertEqual([spec["spec_ref"] for spec in approved["specs"]],
                         [spec["spec_ref"] for spec in SEC_PLAN_SPECS])
        self.assertEqual(effective_discovery_plan(approved)["content_hash"], approved["content_hash"])
        report = sec_plan_report(approved)
        self.assertEqual({row["drift"] for row in report["specs"]}, {"none"})
        env = Environment("canary", state)
        self.addCleanup(lambda: env.core is not None and env.core.close())
        rows = sec_plan_spec_rows(env)
        self.assertEqual({row["status"] for row in rows}, {"ok"})
        connection = sqlite3.connect(f"file:{state / 'core.sqlite'}?mode=ro", uri=True)
        try:
            tables = {row[0] for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table'")}
        finally:
            connection.close()
        self.assertTrue({"earnings_release_reads", "preliminary_actuals"} <= tables)

    def test_an_old_style_plan_renders_the_same_forms_as_a_new_workspace(self) -> None:
        """2026-10-06 (owner decision A): one form set, whatever the approved file says.

        The new workspace's approved plan is governed fields only; an
        environment still on a plan file with its own forms (legacy's v2:
        three transcript phrasings, none the title form) keeps its approval
        and renders exactly the forms the new workspace renders.
        """

        from datetime import date

        from dalton_core.macos_launchagent import (
            ALPHAENGINE_PLAN_SELECTOR, _alphaengine_discovery_plan,
        )
        from dalton_core.mission_source_discovery import (
            build_discovery_parameters, build_discovery_plan, load_discovery_plan,
        )
        from dalton_core.workspace_health_parity import Environment, alphaengine_query_form_rows
        from dalton_core.writer_server import _rendered_alphaengine_plan

        workspace, _ = self.create_workspace()
        state = workspace.state_dir
        new_path = _alphaengine_discovery_plan(state)
        new_approved = load_discovery_plan(new_path)
        new_plan, new_note = _rendered_alphaengine_plan(new_path)
        self.assertEqual(new_note["query_forms"], "code")

        old_state = self.root / "old-style" / "state"
        (old_state / "discovery-plans").mkdir(parents=True)
        legacy_forms = {
            "earnings-call-transcripts": ["{name} {quarter} earnings call transcript",
                                          "{ticker} {quarter} results webcast transcript",
                                          "{name} quarterly earnings conference call"],
            "sell-side-reports": ["{name} equity research report", "{ticker} analyst report"],
        }
        old = build_discovery_plan(
            plan_id="discovery-plan:old-style:alphaengine:2-candidate",
            created_at="2026-09-11T00:00:00.000000+00:00",
            mission_ref=new_approved["mission_ref"], source_ref="source:alphaengine",
            max_calls_24h=130, companies=new_approved["companies"],
            specs=[{**spec, "query_variants": [
                {"query_template": form, "filters": {}}
                for form in legacy_forms[spec["spec_ref"]]]}
                for spec in new_approved["specs"]])
        self.assertEqual(old["schema_version"], "0.6")
        _write_json(old_state / "discovery-plans" / "old-v2.json", old)
        body = {"schema_version": "alphaengine-discovery-plan-selection-0.1",
                "id": "alphaengine-plan-selection:old-style:2", "status": "approved",
                "source_ref": "source:alphaengine", "plan_ref": old["id"],
                "plan_hash": old["content_hash"], "plan_path": "old-v2.json"}
        _write_json(old_state / "discovery-plans" / ALPHAENGINE_PLAN_SELECTOR,
                    {**body, "content_hash": content_hash(body)})
        old_plan, old_note = _rendered_alphaengine_plan(_alphaengine_discovery_plan(old_state))
        self.assertIn("ignored", old_note["query_forms_note"])
        self.assertEqual([spec["query_variants"] for spec in old_plan["specs"]],
                         [spec["query_variants"] for spec in new_plan["specs"]])
        self.assertEqual(old_plan["budget"], {"max_calls_24h": 130})
        for company_ref in new_plan["companies"]:
            for spec in new_plan["specs"]:
                kwargs = dict(spec_ref=spec["spec_ref"], company_ref=company_ref,
                              as_of=date(2026, 10, 6), missing_periods=["FY2026-Q2"])
                self.assertEqual(build_discovery_parameters(old_plan, **kwargs),
                                 build_discovery_parameters(new_plan, **kwargs))
        # Parity: the old file's forms are information, the new file's are ok.
        for directory, expected in ((old_state, "info"), (state, "ok")):
            env = Environment("canary", directory)
            try:
                statuses = {row["status"] for row in alphaengine_query_form_rows(env)
                            if row["check"] != "alphaengine.query_forms.names"}
            finally:
                env.close()
            self.assertEqual(statuses, {expected}, directory)

    def test_a_created_workspace_runs_the_whole_loop_and_survives_an_upgrade(self) -> None:
        workspace, mission = self.create_workspace()
        self.assertEqual({m["company_ref"] for m in mission["universe"]}, {WDGT, GZMO})
        self.assert_governance_baseline(workspace, mission)
        state = workspace.state_dir
        mission, reviews = self.discover(state)

        self.prepare_extraction(state)
        quote = self.quote_of(reviews[0])
        self.assertEqual(quote, self.quote_of(reviews[1]))
        claims_before = self.claim_count()
        summary, verifier = self.extract(state, quote_id=quote, run="v1")
        admitted = [item for item in summary["admitted"] if item["status"] == "admitted"]
        self.assertEqual(len(admitted), 1, json.dumps(summary, indent=1, default=str)[-4000:])
        self.assertEqual(admitted[0]["support_verdict"]["support"], "supported")  # 核验
        self.assertEqual(len(verifier.calls), 1)
        self.assertIn("Widget Robotics", verifier.calls[0]["prompt"])
        self.assertEqual(self.claim_count(), claims_before + 1)  # 入账
        states = {r["document_ref"]: r["state"] for r in self.missions.document_reviews(mission["id"])}
        closed = [ref for ref, value in states.items() if value != "awaiting_human_extraction"]
        still_open = [ref for ref, value in states.items() if value == "awaiting_human_extraction"]
        self.assertEqual((len(closed), len(still_open)), (1, 1), states)

        # Nothing signed by hand: the first mission's policy already carries
        # research_plan_auto_start and the company-facts rule.
        sec = self.sec_company_facts(state, mission, run_key="canary-v1",
                                     annual_accession=SEC_ANNUAL_ACCESSION)
        self.assertEqual(sec["status"], "committed", json.dumps(sec, indent=1, default=str)[-3000:])
        self.assertEqual(sec["mission_stage_claim"]["actor_ref"], AUTOMATION)
        # A replayed run is the same run: no second claim for the accession.
        claims = self.claim_count()
        replay_v1 = self.sec_company_facts(state, mission, run_key="canary-v1",
                                           annual_accession=SEC_ANNUAL_ACCESSION)
        self.assertEqual(replay_v1["status"], "duplicate", replay_v1)
        self.assertEqual(self.claim_count(), claims)
        # The fourth quarter lives in the 10-K: the quarterly coordinator
        # queues it by itself, the lane commits it under the annual rule, and
        # nothing -- neither path -- queues it again.
        self.assert_fourth_quarter_is_dispatched(state, mission)

        self.pass_screen(mission)
        index, dossier, _ = self.index_and_dossier(state, run="v1")
        self.assertNotIn("refused", [item["index_status"] for item in index], index)
        self.assertEqual(index[-1].get("pending", 0), 0, index)
        self.assertIn(dossier["dossier_status"], {"published", "partial_published"},
                      json.dumps(dossier, indent=1, default=str)[-3000:])
        self.assertTrue(dossier["version_ref"])

        self.record_news(mission, closed[0])
        judgement, judge_model = self.judge(state, run="v1")
        self.assertEqual((judgement["status"], judgement["judged"]), ("succeeded", 1),
                         json.dumps(judgement, indent=1, default=str)[-3000:])
        self.assertEqual(len(judge_model.prompts), 1)
        self.assertEqual(self.judged(), 1)
        # ADR-0010: the follow-up sweep runs on a real workspace Core, before
        # anything is paid for, and a NO_CHANGE ledger owes it nothing.
        sweep = judgement["followup_sweep"]
        self.assertEqual((sweep["status"], sweep["mode"], sweep["owed"], sweep["errors"]),
                         ("ok", "apply", [], []), sweep)

        # ---- mission upgrade: policy -> constitution -> mission ----
        claims_before_upgrade = self.claim_count()
        v2 = self.upgrade_mission()
        self.assertNotEqual(v2["id"], mission["id"])
        # The parity check sees the stall before any tick runs ...
        from dalton_core.workspace_health_parity import Environment, check_mission_reviews
        env = Environment("canary", state, service_config=workspace.config_path)
        try:
            [stranded] = [row for row in check_mission_reviews(env)
                          if row["check"] == "reviews.stranded_on_superseded_version"]
        finally:
            env.close()
        self.assertEqual(stranded["status"], "gap", stranded)
        # The stall ws-7d hit: the open review stayed on the old version and
        # the extraction queue under the new one was empty.
        self.assertEqual(self.awaiting_in_active_version(), 0)
        # 2026-09-28, the Guidepoint shape: before any carry runs, a search
        # under the new version finds the open document again and records it
        # already_in_authority (no reconciliation turns it into a review),
        # and the mission is upgraded a second time over it.
        intermediate = v2
        self.research_without_reconciling(intermediate)
        rows = {row["document_ref"]: row["status"]
                for row in self.missions.discovered_documents(intermediate["id"])}
        self.assertEqual(rows.get(still_open[0]), "already_in_authority", rows)
        v2 = self.republish_mission()
        self.assertNotIn(v2["id"], {mission["id"], intermediate["id"]})
        self.research_without_reconciling(v2)
        self.assertEqual(self.awaiting_in_active_version(), 0)
        # The extraction tick's first step (DocumentExtractionCoordinator
        # .dispatch_once -> _carry_awaiting -> carry_forward_awaiting_reviews).
        from dalton_core.document_extraction_launcher import DocumentExtractionCoordinator
        tick = object.__new__(DocumentExtractionCoordinator)
        tick.missions = self.missions
        carried = tick._carry_awaiting()
        self.assertEqual([item["status"] for item in carried], ["carried"], carried)
        self.assertEqual(carried[0]["from_version_ref"], mission["id"])
        self.assertEqual(self.awaiting_in_active_version(), 1)
        self.assertEqual(tick._carry_filtered().get("carryable"), None)
        env = Environment("canary", state, service_config=workspace.config_path)
        try:
            [stranded] = [row for row in check_mission_reviews(env)
                          if row["check"] == "reviews.stranded_on_superseded_version"]
        finally:
            env.close()
        self.assertEqual(stranded["status"], "ok", stranded)
        v2_reviews = self.missions.document_reviews(v2["id"])
        # The one open review, and the re-searched decided document carried
        # with its decision (the search under this version found it again).
        self.assertEqual([(r["document_ref"], r["state"]) for r in v2_reviews
                          if r["state"] == "awaiting_human_extraction"],
                         [(still_open[0], "awaiting_human_extraction")])
        self.assertEqual({r["document_ref"]: r["state"] for r in v2_reviews
                          if r["document_ref"] == closed[0]},
                         {closed[0]: states[closed[0]]})
        self.assertEqual(tick._carry_awaiting(), [])  # idempotent
        self.assert_governance_baseline(workspace, v2)  # ... and none after it

        # Discovery under v2 does not fetch what v1 already acquired.
        for _ in range(4):
            self.coordinator.dispatch_once()
            self.acquisitions.finish()
        self.assertEqual(sorted(self.acquisitions.calls), [DOC_A, DOC_B])

        # Extraction reads the carried review -- and only it.
        summary, verifier = self.extract(state, quote_id=quote, run="v2", max_windows=4)
        self.assertEqual(summary["partial"]["windows_succeeded"], 1,
                         json.dumps(summary, indent=1, default=str)[-3000:])
        self.assertEqual(summary["partial"]["reviews_complete"], 1)
        self.assertEqual(self.awaiting_in_active_version(), 0)
        v2_states = {r["document_ref"]: r["state"] for r in self.missions.document_reviews(v2["id"])}
        self.assertEqual(v2_states.get(closed[0], states[closed[0]]), states[closed[0]],
                         "a document finished under v1 was reopened")
        self.assertNotEqual(v2_states[still_open[0]], "awaiting_human_extraction")
        extraction_claims = self.claim_count() - claims_before_upgrade
        # The statement and citation were already checked under v1: the verdict
        # is reused, not bought again.
        self.assertEqual(len(verifier.calls), 0)

        # SEC company facts: a new filing commits under v2 (the lane's
        # governance precondition and grant resolve against the new version),
        # and the v1 accession still has exactly one claim.
        sec2 = self.sec_company_facts(state, v2, run_key="canary-v2", ticker="GZMO",
                                      accession="0009900002-26-000101", day=1,
                                      annual_only_accession=GZMO_ANNUAL_ACCESSION)
        self.assertEqual(sec2["status"], "committed", json.dumps(sec2, indent=1, default=str)[-3000:])
        self.assertEqual(self.claim_count() - claims_before_upgrade, extraction_claims + 1)
        # GZMO's 10-K reports only fiscal years: its fourth quarter is FY - 9M,
        # derived by the quarter lane under v2 from the statement rows in Core.
        self.assert_fourth_quarter_is_derived(state, v2)
        self.assertEqual(self.claim_count() - claims_before_upgrade, extraction_claims + 2)

        # V2 changes the actual governance bindings, so it permits one new
        # verified read. Unchanged citations do not create a formal version.
        def dossier_rows():
            return list(self.harness.core.connection.execute(
                "SELECT version_id,content_hash FROM company_dossier_versions ORDER BY version_id"))
        dossiers_before = dossier_rows()
        claims_before_dossier = self.claim_count()
        index2, dossier2, dossier_model = self.index_and_dossier(state, run="v2")
        self.assertNotIn("refused", [item["index_status"] for item in index2], index2)
        self.assertEqual((dossier2["status"], dossier2["dossier_status"]),
                         ("succeeded", "no_new_evidence"), dossier2)
        self.assertEqual(dossier2["verification"]["status"], "verified")
        self.assertTrue(dossier2["verification"]["independent"])
        self.assertTrue(dossier_model.prompts)
        self.assertEqual(len(self.last_dossier_verifier.prompts), 1)
        self.assertIsNone(dossier2["version_ref"])
        self.assertEqual(dossier2["formal_authority_writes"], 0)
        self.assertEqual(dossier_rows(), dossiers_before)
        self.assertEqual(self.claim_count(), claims_before_dossier)
        self.assertEqual(set(dossier2["input_consumption"]), set(dossier2["units_drafted"]))
        from dalton_core import research_product_consumption as consumption
        receipts = {}
        for unit, digest in dossier2["input_consumption"].items():
            binding = consumption.scope("dossier", WDGT, unit, [WDGT])
            receipt = consumption.latest(state, binding)
            self.assertEqual(receipt["content_hash"], digest)
            self.assertEqual(receipt["result"], "verified_no_new_evidence")
            self.assertEqual(receipt["schema_version"], "verified-unit-input-consumption:0.1")
            receipts[unit] = receipt
        receipt_bytes = (state / consumption.JOURNAL).read_bytes()
        _, unchanged, unchanged_model = self.index_and_dossier(state, run="v2-dossier-replay")
        self.assertEqual(unchanged["dossier_status"], "nothing_new", unchanged)
        self.assertEqual(unchanged_model.prompts, [])
        self.assertEqual(self.last_dossier_verifier.prompts, [])
        self.assertEqual(unchanged["cost_micros"], 0)
        self.assertEqual(dossier_rows(), dossiers_before)
        self.assertEqual(self.claim_count(), claims_before_dossier)
        self.assertEqual((state / consumption.JOURNAL).read_bytes(), receipt_bytes)
        for unit, receipt in receipts.items():
            self.assertEqual(consumption.latest(state,
                consumption.scope("dossier", WDGT, unit, [WDGT])), receipt)

        # Event judgement: the v1 event is not judged again; a v2 event is judged.
        again, judge_model = self.judge(state, run="v2-replay")
        self.assertEqual(again["judgement_status"], "nothing_unjudged", again)
        self.assertEqual(judge_model.prompts, [])
        self.record_news(v2, still_open[0])
        fresh, _ = self.judge(state, run="v2")
        self.assertEqual((fresh["status"], fresh["judged"]), ("succeeded", 1), fresh)
        self.assertEqual(self.judged(), 2)
        self.assertEqual(fresh["triage"]["routes"], {"escalated:triage": 1}, fresh["triage"])
        # The screen in front of the judge: a group it does not flag (and the
        # audit sample does not pick) is settled without a judgement, and
        # without the judge's call.
        from dalton_core.event_triage import triage_for
        from tests.test_event_triage import QUIET

        quiet = self.record_news(v2, "alphaengine-doc:canary-boilerplate")
        with patch("dalton_core.event_triage.AUDIT_SHARE", 0):
            screened, judge_model = self.judge(state, run="v2-screened", triage=[QUIET])
        self.assertEqual((screened["status"], screened["judgement_status"]),
                         ("succeeded", "screened_out"), screened)
        self.assertEqual(judge_model.prompts, [])
        self.assertEqual(self.triage_model.purposes, ["event_triage"])
        self.assertEqual(self.judged(), 2)
        self.assertEqual(triage_for(self.harness.core.connection, quiet["id"])["outcome"],
                         "screened_out")

        # ---- maintenance with the extraction queue empty ----
        # Nothing is open under v2 any more.  The claim-support checks and the
        # P13i subject re-check used to run only inside a drafting child, so an
        # empty queue stopped them; the lane now starts a support-only child
        # and that child runs all of them.
        self.assertEqual(self.awaiting_in_active_version(), 0)
        self.assert_maintenance_runs_with_an_empty_queue(state, v2)

        # ---- the discovery tick shares its budget ----
        self.assert_discovery_budget_is_shared(workspace)

        # A later real mission binding is a distinct input, and can be read
        # again even though the last formal dossier still binds v1.
        v3 = self.republish_mission()
        self.assertNotEqual(v3["content_hash"], v2["content_hash"])
        _, changed_binding, changed_model = self.index_and_dossier(state, run="v3-dossier")
        self.assertTrue(changed_model.prompts)
        self.assertTrue(self.last_dossier_verifier.prompts)
        self.assertEqual(changed_binding["verification"]["status"], "verified")
        self.assertEqual(changed_binding["dossier_status"], "no_new_evidence")
        for unit in changed_binding["input_consumption"]:
            receipt = consumption.latest(state, consumption.scope("dossier", WDGT, unit, [WDGT]))
            self.assertNotEqual(receipt["snapshot"]["input_hash"], receipts[unit]["snapshot"]["input_hash"])
        _, v3_replay, v3_model = self.index_and_dossier(state, run="v3-dossier-replay")
        self.assertEqual(v3_replay["dossier_status"], "nothing_new", v3_replay)
        self.assertEqual(v3_model.prompts, [])
        self.assertEqual(self.last_dossier_verifier.prompts, [])

    def assert_discovery_budget_is_shared(self, workspace) -> None:
        """A slow AlphaEngine coordinator must not starve SEC and web search.

        Legacy, 2026-09-25 to 09-29: the AlphaEngine reconciliation took most
        of the 20 s the three coordinators share, the other two reported
        ``deferred: tick budget exhausted`` every tick, and the tick ledger
        showed an idle lane. Here the created workspace's own writer op runs
        its real AlphaEngine coordinator, made to overrun the whole budget,
        with SEC and web search behind it.
        """

        from dalton_core import mission_source_discovery as msd
        from dalton_core.tick_ledger import TickLedger
        from dalton_core.workspace_health_parity import Environment, check_lanes

        state = workspace.state_dir
        clock = {"now": 5000.0}
        real = self.coordinator
        reached: dict[str, list[float]] = {"sec_filings_index": [], "web_search": []}

        class Overrunning:
            source_ref = real.source_ref
            plan = real.plan

            def dispatch_once(inner, deadline=None):
                result = real.dispatch_once(deadline)
                clock["now"] += msd.TICK_BUDGET_SECONDS + 5.0
                return result

        class Recorder:
            def __init__(inner, key):
                inner.key, inner.plan = key, real.plan
                inner.source_ref = {"sec_filings_index": msd.SEC_SOURCE_REF,
                                    "web_search": msd.WEB_SEARCH_SOURCE_REF}[key]

            def dispatch_once(inner, deadline=None):
                reached[inner.key].append(clock["now"])
                clock["now"] += 0.5
                return {"status": "idle", "source_ref": inner.source_ref,
                        "discovery": {"status": "idle"}, "acquisitions_launched": 0}

        writer = self.writer
        writer._source_discovery = Overrunning()
        writer._sec_filings_source_discovery = Recorder("sec_filings_index")
        writer._web_source_discovery = Recorder("web_search")
        results = []
        with patch.object(msd, "_monotonic", lambda: clock["now"]), \
                TickLedger(state / "tick-ledger.sqlite") as ledger:
            started = datetime(2026, 9, 29, 12, tzinfo=timezone.utc)
            for index in range(4):
                result = writer._op_dispatch_mission_source_discovery({})
                results.append(result)
                ledger.append_tick({"status": "ok", "mission_source_discovery": result},
                                   started_at=started + timedelta(minutes=5 * index))
                clock["now"] += 300.0
                self.acquisitions.finish()
        order = [r["coordinator_order"] for r in results]
        self.assertEqual(order[0], ["alphaengine", "sec_filings_index", "web_search"])
        self.assertEqual(results[0]["coordinators"]["web_search"]["status"], "deferred")
        # Starved on the first tick, first on the next: both are reached.
        self.assertEqual(order[1][:2], ["sec_filings_index", "web_search"])
        self.assertGreaterEqual(len(reached["sec_filings_index"]), 2, order)
        self.assertGreaterEqual(len(reached["web_search"]), 2, order)
        # Nothing new was fetched by the extra AlphaEngine ticks.
        self.assertEqual(sorted(self.acquisitions.calls), [DOC_A, DOC_B])
        # The ledger keeps each coordinator's word, and the parity check sees
        # no coordinator deferred three ticks running.
        env = Environment("canary", state, service_config=workspace.config_path)
        try:
            rows = {row["check"]: row for row in check_lanes(env)
                    if row["check"].startswith("mission_source_discovery.")}
        finally:
            env.close()
        self.assertEqual({key for key in rows if key.endswith(".deferred")},
                         {f"mission_source_discovery.{key}.deferred"
                          for key in ("alphaengine", "sec_filings_index", "web_search")})
        # The searches that ran all succeeded, and nothing held the queue.
        self.assertEqual(rows["mission_source_discovery.alphaengine.searches"]["status"], "ok", rows)
        self.assertFalse([key for key in rows if key.endswith(".selection")], rows)
        for key in ("sec_filings_index", "web_search"):
            self.assertEqual(rows[f"mission_source_discovery.{key}.deferred"]["status"], "ok",
                             rows)

    def assert_maintenance_runs_with_an_empty_queue(self, state: Path, mission) -> None:
        from dalton_core import document_extraction_cli
        from dalton_core.document_extraction_cli import run_extraction
        from dalton_core.document_extraction_launcher import DocumentExtractionCoordinator
        from tests.test_support_lane_independent import SupportFakeLauncher

        launcher_root = self.root / "support-launcher"
        launcher_root.mkdir()
        launcher = SupportFakeLauncher(launcher_root)
        tick = DocumentExtractionCoordinator(missions=self.missions, launcher=launcher).dispatch_once()
        self.assertEqual((tick["status"], tick.get("mode"), tick["awaiting"]),
                         ("launched", "support_only", 0), tick)
        self.assertTrue(launcher.starts[-1].get("support_only"))

        fixture = state / "canary-fixture-support-only.json"
        _write_json(fixture, {"schema_version": "0.1", "suggestions": []})
        ran: list[str] = []
        base = document_extraction_cli.ExtractionHost

        class Host(base):
            def __init__(inner, **kwargs):
                super().__init__(**kwargs)
                inner._claim_support_verifier = object()

        with patch.object(document_extraction_cli, "ExtractionHost", Host), \
                patch.object(document_extraction_cli, "_run_claim_support_recheck",
                             lambda host, db: ran.append("support_recheck") or {"status": "idle"}), \
                patch.object(document_extraction_cli, "_secondary_sweep",
                             side_effect=AssertionError("a support-only run reads no document")):
            summary = run_extraction(
                state_dir=state, model_config_path=self.model_config,
                summary_dir=state / "extractions" / "canary-support-only", spool_dir=state / "spool",
                scheduler_db=state / "scheduler.sqlite", requested_by=None,
                max_windows=1, max_numeric_windows=0, max_discovery_windows=0,
                connector_governance=None, web_fetch_governance=None,
                hermetic_fixture=fixture,
                candidate_staging=state / "research-review" / "candidate-staging.sqlite",
                support_only=True)
        self.assertEqual((summary["stop_reason"], summary["reviews_scanned"], summary["drafted"]),
                         ("support_only", 0, []), summary)
        self.assertEqual(ran, ["support_recheck"])
        [reevaluation] = summary["subject_reevaluation"]
        self.assertEqual((reevaluation["mission_version_ref"], reevaluation["stop_reason"]),
                         (mission["id"], None), reevaluation)

    # -- 2026-09-28: a support retirement about another subject -------------------
    #
    # The claim-review tick now also reads ``citation_support_rejected``
    # retirements whose verdict is supported and about another subject.  In a
    # workspace made by the formal flow the mission's industry is whatever the
    # owner typed ("Industrial Robotics"), for which the industry rule has no
    # vocabulary: nothing may be reattributed -- no industry is made up, the
    # retirement stands -- and the refusal is remembered, so the next tick
    # reads nothing again.  A covered company named as the other subject is
    # refused before anything else.

    def test_a_created_workspace_planner_sees_what_is_settled_and_reopened(self) -> None:
        """2026-10-01: legacy's 27 withdrawals moved nothing the planner hashed.

        A workspace created by the formal flow must carry the settlement in
        its very first planner state and coordinator signature: available,
        nothing settled, nothing reopened -- not an unavailable read on a Core
        that never had a directed read.
        """

        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.research_planner_cli import build_state
        from dalton_core.research_planner_launcher import ResearchPlannerCoordinator
        from dalton_core.research_state import planning_settlement_hash

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        with DaltonStore(state / "core.sqlite") as store:
            built = build_state(store, CoverageMissionAuthority(store), mission,
                                plans_dir=state / "discovery-plans",
                                as_of="2026-10-01T00:00:00+00:00")
            signature = ResearchPlannerCoordinator(store=store, launcher=None)._signature()
        settlement = built["document_research_settlement"]
        self.assertEqual(settlement["status"], "available", settlement)
        self.assertEqual((settlement["admissions"], settlement["settled"],
                          settlement["withdrawn"]), (0, 0, 0))
        self.assertEqual(len(planning_settlement_hash(built)), 64)
        self.assertEqual({c["company_ref"] for c in built["companies"]}, {WDGT, GZMO})
        for company in built["companies"]:
            self.assertEqual(company["document_research_reopened"], [])
        self.assertEqual(len(signature["document_research_settlement"]), 64)

    def test_a_created_workspace_walks_the_stage_ledger_and_says_what_each_gate_waits_on(self) -> None:
        """ADR-0011: the stage bridge and the memo lane on a Core nobody patched.

        Live, the industry_model gate waited from 2026-09-30 on every company
        for a reason nobody could read: the packaged module looked for a
        repository file beside site-packages.  On a brand-new workspace the
        same tick must run, enter the Deep Insight Gate when a draft exists,
        keep the human gate human, and name -- in Chinese -- the input each
        model gate is still missing.
        """

        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.investment_memo_cli import collect_frozen_input
        from dalton_core.mission_model_stage_lane import advance_once

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        with DaltonStore(state / "core.sqlite") as store:
            missions = CoverageMissionAuthority(store)
            mission = missions.mission(mission["id"])
            self.assertEqual(advance_once(missions, store.connection, mission)["status"], "idle")
            for status in ("entered", "gate_passed"):
                missions.record_stage(
                    mission_version_ref=mission["id"],
                    mission_version_hash=mission["content_hash"], company_ref=WDGT,
                    stage_ref="initial_screen", status=status, evidence_refs=[mission["id"]],
                    rationale="canary screen", actor_ref=AUTOMATION,
                    idempotency_key=f"canary-screen:{status}")
            # Passed the screen, no gate draft yet: nothing to enter.
            self.assertEqual(advance_once(missions, store.connection, mission)["status"], "idle")
            with patch("dalton_core.mission_model_stage_lane.latest_gate_draft",
                       return_value="deep-insight-gate-version:canary:1"):
                entered = advance_once(missions, store.connection, mission)
            self.assertEqual((entered["status"], entered["stage_ref"]),
                             ("entered", "deep_insight_gate"), entered)
            ladder = missions.current_stage_state(mission["mission_ref"], WDGT)
            self.assertEqual((ladder["current_stage"], ladder["current_status"]),
                             ("deep_insight_gate", "entered"))
            # Passing it stays human: the automation tick leaves it entered.
            with patch("dalton_core.mission_model_stage_lane.latest_gate_draft",
                       return_value="deep-insight-gate-version:canary:1"):
                self.assertEqual(
                    advance_once(missions, store.connection, mission)["status"], "idle")
            missions.record_stage(
                mission_version_ref=mission["id"], mission_version_hash=mission["content_hash"],
                company_ref=WDGT, stage_ref="deep_insight_gate", status="gate_passed",
                evidence_refs=["deep-insight-gate-version:canary:1"], rationale="owner approved",
                actor_ref=OWNER, idempotency_key="canary-gate:approve")
            waiting = advance_once(missions, store.connection, mission)
            self.assertEqual((waiting["status"], waiting["stage_ref"]),
                             ("waiting", "industry_model"), waiting)
            self.assertEqual(waiting["waiting:WDGT"], "行业模型缺：还没有行业框架")
            self.assertNotIn("读取依据失败", waiting["reason"])
            memo = collect_frozen_input(store)
        self.assertEqual(memo["status"], "idle", memo)
        self.assertIn("WDGT 停在行业模型（进行中），等深度模型构建通过审核", memo["reason"])

    def test_a_read_the_mission_moved_under_is_carried_with_its_conclusion(self) -> None:
        """2026-09-29, DXC's Guidepoint excerpt d50efc...: read twice across two upgrades.

        11:32 under v26: four Claims admitted, then the close failed because
        the mission had moved ("completion proof failed: source discovery must
        bind the active mission version"); the open review was carried to v27
        and v28 and read again at 12:31, three more Claims on exactly the same
        spans.  Here: the read and the failed close under v1; a carry that ran
        before the conclusion was written (v2, the DXC v27 shape); a second
        upgrade whose carry finds the conclusion and carries it (v3); and the
        other document -- the same bytes, the same cited span -- is the Claim
        already admitted, not a second one.
        """

        from dalton_core import document_admission_scope
        from dalton_core.coverage_mission import CoverageMissionConflict
        from dalton_core.document_extraction_launcher import DocumentExtractionCoordinator
        from dalton_core.document_read_completion import DocumentReadCompletionAuthority

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        mission, reviews = self.discover(state)
        self.prepare_extraction(state)
        quote = self.quote_of(reviews[0])
        before = self.claim_count()

        def mission_moved(*args, **kwargs):
            raise CoverageMissionConflict("source discovery must bind the active mission version")

        with patch.object(DocumentReadCompletionAuthority, "record", mission_moved):
            summary, verifier = self.extract(state, quote_id=quote, run="moved-v1")
        [admitted] = [item for item in summary["admitted"] if item["status"] == "admitted"]
        self.assertEqual(self.claim_count(), before + 1)
        self.assertEqual(len(verifier.calls), 1)
        [held] = summary["resolved_reviews"]
        self.assertIn("completion proof failed", held["reason"])
        by_review = {r["review_id"]: r["document_ref"] for r in self.missions.document_reviews(mission["id"])}
        read_doc = by_review[admitted["review_id"]]
        [other_doc] = [ref for ref in by_review.values() if ref != read_doc]
        self.assertEqual(self.awaiting_in_active_version(), 2)  # the close failed: both open
        outcome = document_admission_scope.admission_outcome(
            self.harness.core.connection, document_ref=read_doc, company_ref=WDGT)
        self.assertEqual((outcome["admitted"], outcome["claim_version_refs"]),
                         (1, [admitted["claim_version_ref"]]))

        tick = object.__new__(DocumentExtractionCoordinator)
        tick.missions = self.missions
        # v2: the carry ran before the conclusion could be read (DXC's v27).
        self.upgrade_mission()
        with patch.object(document_admission_scope, "admission_outcome", lambda *a, **k: None):
            carried = tick._carry_awaiting()
        self.assertEqual({item["carried_to"]["state"] for item in carried}, {"awaiting_human_extraction"})
        # v3: this carry finds it, and the read document is not opened again.
        v3 = self.republish_mission()
        carried = {item["document_ref"]: item["carried_to"] for item in tick._carry_awaiting()}
        self.assertEqual(carried[read_doc]["state"], "extraction_staged")
        self.assertEqual(carried[read_doc]["carried_conclusion_of"], admitted["review_id"])
        self.assertEqual(carried[other_doc]["state"], "awaiting_human_extraction")
        self.assertEqual(self.awaiting_in_active_version(), 1)
        staged = next(r for r in self.missions.document_reviews(v3["id"]) if r["document_ref"] == read_doc)
        self.assertEqual(staged["candidate_claim_version_ref"], admitted["candidate_claim_ref"])
        self.assertIn("is not read again", staged["rationale"])

        # v3 extraction: only the other document is read, and its statement
        # cites the same bytes at the same span as the Claim already admitted.
        summary, verifier = self.extract(state, quote_id=quote, run="moved-v3", max_windows=4)
        self.assertEqual(self.claim_count(), before + 1, json.dumps(summary, default=str)[-3000:])
        self.assertEqual(len(verifier.calls), 0)
        [duplicate] = summary["admitted"]
        self.assertEqual((duplicate["status"], duplicate["claim_version_ref"], duplicate["scope"]["basis"]),
                         ("duplicate", admitted["claim_version_ref"], "same_span"))
        self.assertEqual(summary["admission_scope"]["duplicates_of_ledger"], 1)
        self.assertEqual(self.awaiting_in_active_version(), 0)

    def test_an_overlong_companies_tag_neither_stalls_the_queue_nor_hides_the_company(self) -> None:
        """2026-09-29: one candidate tagged with 24-193 issuers held every
        AlphaEngine acquisition on both environments behind
        ``selection_pending: CandidateSelectionError: candidate companies is
        invalid``.  In a created workspace the long tag list is truncated with
        the covered company kept, a page the selector still refuses is
        quarantined without holding the pages behind it, and the owner can
        release it."""

        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.discovery_candidate_selection import (
            SELECTION_RULES_VERSION, CandidateSelectionError,
        )
        from dalton_core.discovery_selection_quarantine import DiscoverySelectionQuarantine
        from dalton_core.macos_launchagent import _alphaengine_discovery_plan
        from dalton_core.mission_source_discovery import (
            MissionSourceDiscoveryCoordinator, discovery_coordinator_state, load_discovery_plan,
        )
        from tests.test_mission_source_discovery import FakeSearchLauncher, SearchHarness

        workspace, _ = self.create_workspace()
        state = workspace.state_dir
        crowd = [f"Unrelated Issuer {index} Inc." for index in range(36)] + ["Widget Robotics Inc."]
        poisoned_page = [{"doc_id": "990000000000011"}, {"doc_id": "990000000000012"}]
        crowded_page = [{"doc_id": "990000000000021", "title": "Daily research highlights",
                         "companies": crowd},
                        {"doc_id": "990000000000022", "companies": ["Gizmo Motion Corp."]}]
        harness = SearchHarness(state, poisoned_page)
        self.addCleanup(harness.close)
        self.harness = harness
        missions = CoverageMissionAuthority(harness.core)
        plan = load_discovery_plan(_alphaengine_discovery_plan(state))
        # 2026-10-02: the generated plan names each company by its registered
        # name, not its ticker only (the ws-7d transcript stall).
        self.assertEqual(plan["companies"][WDGT]["name"], "Widget Robotics")
        self.assertEqual(plan["companies"][WDGT]["ticker"], "WDGT")

        class Refusing(Selection):
            """Refuses the first page it is shown, as ``selection_prompt`` does."""

            def __init__(self) -> None:
                super().__init__()
                self.state = state
                self.quarantine = DiscoverySelectionQuarantine(
                    state / "discovery-selections" / "quarantine")
                self.poisoned: set[str] = set()
                self.views: dict[str, dict] = {}

            def start(self, *, discovery_ref, view, **kwargs):
                if not self.started and not self.poisoned:
                    self.poisoned.add(discovery_ref)
                if discovery_ref in self.poisoned:
                    raise CandidateSelectionError("company selection identity is invalid")
                self.views[discovery_ref] = view
                return super().start(discovery_ref=discovery_ref, view=view, **kwargs)

        selection = Refusing()
        self.acquisitions = Acquisitions(harness, state)
        coordinator = MissionSourceDiscoveryCoordinator(
            store=harness.core, missions=missions, plan=plan,
            search_launcher=FakeSearchLauncher(harness, missions, plan),
            acquisition_launcher=self.acquisitions, selection_launcher=selection,
            spool_dir=state / "spool", clock=harness.clock)
        ticks = [coordinator.dispatch_once()]
        harness.handle.results = [dict(item) for item in crowded_page]
        for _ in range(8):
            ticks.append(coordinator.dispatch_once())
            self.acquisitions.finish()
        dump = json.dumps(ticks, indent=1, default=str)[-4000:]

        # The refused page is set aside, audibly, and nothing else waits on it.
        [poisoned] = sorted(selection.poisoned)
        record = selection.quarantine.record(poisoned)
        self.assertEqual(record["status"], "quarantined")
        self.assertEqual(record["rules_version"], SELECTION_RULES_VERSION)
        self.assertIn("company selection identity is invalid", record["reason"])
        self.assertEqual(selection.quarantine.active(rules_version=SELECTION_RULES_VERSION),
                         [poisoned])
        self.assertIn("alphaengine-doc:990000000000021", self.acquisitions.calls, dump)
        self.assertIn("alphaengine-doc:990000000000022", self.acquisitions.calls, dump)
        self.assertNotIn("alphaengine-doc:990000000000011", self.acquisitions.calls, dump)
        quarantined_tick = next(tick for tick in ticks if tick.get("selection_quarantined_now"))
        self.assertEqual(quarantined_tick["selection_quarantined_now"], [poisoned])
        self.assertEqual(discovery_coordinator_state(ticks[-1])["selection_quarantined"], 1)

        # The overlong tag list is truncated, and the covered company -- known
        # to the plan only as WDGT -- is the one kept first.
        crowded = [view["candidates"][0] for view in selection.views.values()
                   if view["candidates"][0]["document_ref"] == "alphaengine-doc:990000000000021"]
        self.assertTrue(crowded, dump)
        candidate = crowded[0]
        self.assertEqual(candidate["companies"][0], "Widget Robotics Inc.")
        self.assertEqual(candidate["companies_original_count"], len(crowd))
        self.assertTrue(candidate["companies_truncated"])
        self.assertEqual(candidate["companies_covered_kept"], 1)

        # The owner releases the page; it is read like any other.
        with self.assertRaises(Exception):
            selection.quarantine.release(poisoned, actor_ref=AUTOMATION, reason="not a human")
        selection.poisoned.clear()
        selection.quarantine.release(poisoned, actor_ref=OWNER, reason="selector fixed")
        self.assertEqual(selection.quarantine.active(rules_version=SELECTION_RULES_VERSION), [])
        for _ in range(4):
            coordinator.dispatch_once()
            self.acquisitions.finish()
        self.assertIn("alphaengine-doc:990000000000011", self.acquisitions.calls)
        events = [event["event"] for event in selection.quarantine.record(poisoned)["events"]]
        self.assertEqual(events, ["quarantined", "released"])

    def test_an_exhausted_selection_ticket_is_quarantined_and_the_queue_moves(self) -> None:
        """2026-10-02: one failed selection ticket (legacy ``database is
        locked``, ws-7d a reply in the wrong shape) answered ``selection_pending:
        exhausted`` for the same discovery on every tick and no AlphaEngine
        document was acquired for days.  An exhausted ticket is now
        quarantined with its ticket and reason, the discovery behind it is
        acquired in the same tick, and an owner release puts it back."""

        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.discovery_candidate_selection import SELECTION_RULES_VERSION
        from dalton_core.discovery_selection_quarantine import DiscoverySelectionQuarantine
        from dalton_core.macos_launchagent import _alphaengine_discovery_plan
        from dalton_core.mission_source_discovery import (
            MissionSourceDiscoveryCoordinator, discovery_coordinator_state, load_discovery_plan,
        )
        from tests.test_mission_source_discovery import FakeSearchLauncher, SearchHarness

        workspace, _ = self.create_workspace()
        state = workspace.state_dir
        stalled_page = [{"doc_id": "990000000000031"}, {"doc_id": "990000000000032"}]
        next_page = [{"doc_id": "990000000000041"}]
        harness = SearchHarness(state, stalled_page)
        self.addCleanup(harness.close)
        self.harness = harness
        missions = CoverageMissionAuthority(harness.core)
        plan = load_discovery_plan(_alphaengine_discovery_plan(state))
        reason = "CandidateSelectionError: selection has an invalid closed shape"

        class Exhausting(Selection):
            """The first discovery's selection failed every attempt it had."""

            def __init__(self) -> None:
                super().__init__()
                self.state = state
                self.quarantine = DiscoverySelectionQuarantine(
                    state / "discovery-selections" / "quarantine")
                self.stalled: set[str] = set()

            def start(self, *, discovery_ref, view, **kwargs):
                if not self.started and not self.stalled:
                    self.stalled.add(discovery_ref)
                if discovery_ref in self.stalled:
                    return {"id": "discovery-selection:" + "c" * 24, "status": "exhausted",
                            "discovery_ref": discovery_ref, "recovery_epoch": 2,
                            "attempts": 3, "failure_reason": reason,
                            "completed_at": "2026-10-01T11:50:42.587857+00:00",
                            "summary": {"status": "failed", "failure_reason": reason}}
                return super().start(discovery_ref=discovery_ref, view=view, **kwargs)

        selection = Exhausting()
        self.acquisitions = Acquisitions(harness, state)
        coordinator = MissionSourceDiscoveryCoordinator(
            store=harness.core, missions=missions, plan=plan,
            search_launcher=FakeSearchLauncher(harness, missions, plan),
            acquisition_launcher=self.acquisitions, selection_launcher=selection,
            spool_dir=state / "spool", clock=harness.clock)
        ticks = [coordinator.dispatch_once()]
        harness.handle.results = [dict(item) for item in next_page]
        for _ in range(8):
            ticks.append(coordinator.dispatch_once())
            self.acquisitions.finish()
        dump = json.dumps(ticks, indent=1, default=str)[-4000:]

        [stalled] = sorted(selection.stalled)
        record = selection.quarantine.record(stalled)
        self.assertEqual(record["status"], "quarantined", dump)
        self.assertEqual(record["error_type"], "DiscoverySelectionExhausted")
        self.assertEqual(record["rules_version"], SELECTION_RULES_VERSION)
        self.assertIn(reason, record["reason"])
        self.assertIn("after 3 attempt(s)", record["reason"])
        self.assertEqual(record["detail"]["ticket_ref"], "discovery-selection:" + "c" * 24)
        self.assertEqual(record["detail"]["failure_reason"], reason)
        self.assertEqual(record["detail"]["failure_type"], "CandidateSelectionError")
        self.assertEqual(record["detail"]["attempts"], 3)
        self.assertIn("alphaengine-doc:990000000000041", self.acquisitions.calls, dump)
        self.assertNotIn("alphaengine-doc:990000000000031", self.acquisitions.calls, dump)
        quarantined_tick = next(tick for tick in ticks if tick.get("selection_quarantined_now"))
        self.assertEqual(quarantined_tick["selection_quarantined_now"], [stalled])
        self.assertFalse(any(
            (tick["acquisition"].get("status"), tick["acquisition"].get("reason"))
            == ("selection_pending", "exhausted") for tick in ticks), dump)
        self.assertEqual(discovery_coordinator_state(ticks[-1])["selection_quarantined"], 1)

        # The owner releases it; the next attempt succeeds and it is read.
        selection.stalled.clear()
        selection.quarantine.release(stalled, actor_ref=OWNER, reason="selector retried")
        for _ in range(4):
            coordinator.dispatch_once()
            self.acquisitions.finish()
        self.assertIn("alphaengine-doc:990000000000031", self.acquisitions.calls)
        events = [event["event"] for event in selection.quarantine.record(stalled)["events"]]
        self.assertEqual(events, ["quarantined", "released"])

    def test_an_about_other_retirement_is_never_given_an_industry_the_rule_cannot_read(self) -> None:
        from dalton_core.claim_industry_reattribution import (
            ClaimIndustryReattributionAuthority, covering_missions, industry_reattributions,
        )
        from dalton_core.claim_retirement import ClaimRetirementAuthority, retired_claim_version_refs
        from dalton_core.claim_review import ClaimReviewDriver
        from dalton_core.claim_support_backfill import ClaimSupportBackfill
        from dalton_core.claim_support_verification import BACKFILL_PURPOSE, ClaimSupportVerifier
        from dalton_core.coverage_mission import CoverageMissionAuthority
        from tests.test_claim_retirement import ClaimRetirementHarness, _Spool
        from tests.test_claim_support_backfill import StatementModel

        workspace, mission = self.create_workspace()
        self.assertEqual(mission["industry_ref"], "industry:industrial-robotics")
        store = DaltonStore(workspace.state_dir / "core.sqlite")
        self.addCleanup(store.close)
        import dalton_core
        store.connection.executescript(
            (Path(dalton_core.__file__).parent / "transcript_correction_schema.sql").read_text(
                encoding="utf-8"))
        [covering] = {row["mission_version_ref"] for row in covering_missions(store.connection).values()}
        self.assertEqual(covering, mission["id"])

        # Two Claims filed under WDGT, through the Ledger's own citation chain.
        shim = type("Shim", (), {"store": store, "objects": {}, "_seq": 0})()
        industry = ("Industrial robotics vendors across the sector report that factory "
                    "automation orders from enterprise customers stayed strong in the quarter.")
        company = ("Gizmo Motion raised its full-year outlook on stronger orders from "
                   "automotive plants, the digest says.")
        claims = []
        for statement in (industry, company):
            source = "Weekly robotics digest. " + statement + " Other remarks followed."
            start = source.index(statement)
            claims.append(ClaimRetirementHarness.claim(
                shim, subject=WDGT, statement=statement, source=source,
                span=(start, start + len(statement))))
        missions = CoverageMissionAuthority(store)
        challenges = ClaimRetirementAuthority(store)
        reattributions = ClaimIndustryReattributionAuthority(store)
        driver = ClaimReviewDriver(store=store, missions=missions, challenges=challenges,
                                   spool=_Spool(shim.objects), needles={WDGT: ["widget"]},
                                   reattributions=reattributions)
        verifier = ClaimSupportVerifier(
            store=store, purpose=BACKFILL_PURPOSE, daily_cap_micros=500_000,
            model_call=StatementModel({
                industry: ("supported", "about_other", "Industrial Robotics sector"),
                company: ("supported", "about_other", "Gizmo Motion Corp. (GZMO)")}),
            spend_today=lambda _p, _d: 0, producer_family=lambda ref: "deepseek-v4")
        original = driver._citations
        driver._citations = lambda: {ref: {**item, "route_decision_ref": "route-decision:drafter"}
                                     for ref, item in original().items()}
        backfill = ClaimSupportBackfill(store=store, missions=missions, verifier=verifier,
                                        reader=driver, challenges=challenges, claim_sources=["test"])
        retired = backfill.run_once(max_items=10)
        self.assertEqual(set(retired["retired"]), {claim["ref"] for claim in claims}, retired)
        before = retired_claim_version_refs(store.connection)

        driver._citations = original
        first = driver.reattribute_industry_findings(principal=AUTOMATION, show=None)
        self.assertEqual((first["reattributed"], first["awaiting_rereview"]), ([], 0), first)
        refusals = {item["claim_version_ref"]: item["refusal"] for item in first["refused_examples"]}
        self.assertEqual(refusals, {claims[0]["ref"]: "no_industry_lexicon",
                                    claims[1]["ref"]: "other_subject_names_company:covered:"
                                    + GZMO + ",Gizmo,Motion,Corp"}, first)
        self.assertEqual(reattributions.reattributions(), [])
        self.assertEqual(industry_reattributions(store.connection), {})
        self.assertEqual(retired_claim_version_refs(store.connection), before)
        again = driver.reattribute_industry_findings(principal=AUTOMATION, show=None)
        self.assertEqual((again["examined"], again["already_reviewed"]), (0, 2), again)


class NewWorkspaceInheritsTheMatureEnvironmentTests(unittest.TestCase):
    """Canary (2026-09-30, review section 2): a created workspace starts at parity.

    A fake mature ("legacy") environment on a fake host -- 22 lane model
    configurations, a day-budget lineage ending at $500, the host's credential
    slots, the crowd tools / grants / a prior-research corpus inside its state,
    a claim-support settings file -- and the two runtime templates exported
    from it *before* the owner raised the document-draft output limit to
    16 000.  The workspace is then created through ``provision_runtime``, the
    exact sequence the cockpit's create button runs, and must come up with:
    the same credential slots, the same per-call budgets (the 16 000 included),
    the same day-budget head in every lane *and* in service.json, the crowd
    lane wired, its own (unshared) prior-research directory, every connector
    permission approved, the claim-support settings, the findata /
    ir-page-watch contracts, a weekly brief coordinator waiting for a plan,
    its own publication worker, a workspace-named SEC lane, a health report
    whose pre-mission governance rows are information, and nothing left for
    the repair script to do.
    """

    LEGACY_SLOTS = ["credential-slot:openclaw:deepseek", "credential-slot:openclaw:openai",
                    "credential-slot:openclaw:claude-cli"]
    DRAFT = "mission-document-draft-model-config.json"

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(dir="/tmp")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        environ = patch.dict(os.environ, {
            "DALTON_OPENCLAW_WORKSPACE": str(self.root / "no-openclaw")}, clear=False)
        environ.start(); self.addCleanup(environ.stop)
        for name in ("DALTON_WORKSPACE_MANIFEST", "DALTON_PRIOR_RESEARCH_DIR"):
            os.environ.pop(name, None)
        for target in ("socket.create_connection", "socket.socket.connect",
                       "urllib.request.urlopen", "http.client.HTTPConnection.connect"):
            guard = patch(target, _no_network)
            guard.start(); self.addCleanup(guard.stop)
        from dalton_core.workspace_template_freshness import clear_cache
        clear_cache()

    # -- the mature environment and its templates --------------------------------

    def build_legacy(self) -> None:
        from dalton_core.bootstrap import bootstrap
        from dalton_core.model_deployment import POLICY_REF, install_openclaw_catalog
        from dalton_core.thesis_impact_budget import ThesisImpactBudgetStore
        from dalton_core.workspace_model_setup import EXPECTED_CONFIG_NAMES

        host = self.root / "hostfiles"; host.mkdir()
        self.socket, self.key = host / "model.sock", host / "model.key"
        self.openclaw, self.tailscale = host / "openclaw.json", host / "tailscale"
        for path in (self.socket, self.key, self.openclaw, self.tailscale):
            path.write_text("{}")
        self.shared_policy = host / "model-call-budget-policy.json"
        body = {"schema_version": "dalton-shared-call-budget-policy-0.1",
                "default_max_cost_usd": 1.0, "purpose_max_cost_usd": {}, "revision": 1,
                "prior_hash": None, "updated_at": "2026-09-14T00:00:00+00:00",
                "actor_ref": "human:owner"}
        self.shared_policy.write_text(json.dumps({**body, "content_hash": content_hash(body)}))
        from dalton_core.shared_daily_budget import policy_wire, write_owner_only
        self.daily_policy = host / "shared-daily-budget.json"
        write_owner_only(self.daily_policy, policy_wire(
            max_daily_cost_usd=500.0, max_daily_paid_calls=100000,
            max_alphaengine_calls_24h=130, revision=1, prior_hash=None,
            updated_at="2026-09-18T00:00:00.000000+00:00", actor_ref="human:owner"))

        legacy = self.root / "legacy"
        self.legacy_state = legacy / "state" / "dalton-core"
        self.legacy_config = legacy / "config" / "service.json"
        self.legacy_state.mkdir(parents=True)
        now = datetime(2026, 9, 14, tzinfo=timezone.utc)
        install_openclaw_catalog(self.legacy_state / "model-router.sqlite", checked_at=now)
        self.head = "thesis-impact-day-budget-policy:configured:canary-500"
        with ThesisImpactBudgetStore(self.legacy_state / "thesis-impact-budget.sqlite",
                                     clock=lambda: now) as budget:
            prior = None
            for policy, cap in (("thesis-impact-day-budget-policy:production:1", 25),
                                ("thesis-impact-day-budget-policy:production:3", 180),
                                (self.head, 500)):
                budget.register_policy(policy_version_id=policy,
                                       day_cap_micros=cap * 1_000_000,
                                       prior_version_id=prior)
                prior = policy
        base = {
            "routing_policy_ref": POLICY_REF, "credential_slot_refs": list(self.LEGACY_SLOTS),
            "model_router_db": str(self.legacy_state / "model-router.sqlite"),
            "broker_socket": str(self.socket), "broker_auth_key": str(self.key),
            "broker_client_id": "client:dalton-core", "expected_agent_id": "chem",
            "budget_db": str(self.legacy_state / "thesis-impact-budget.sqlite"),
            "budget_policy_ref": self.head,
            "shared_call_budget_policy_path": str(self.shared_policy),
        }
        for name in EXPECTED_CONFIG_NAMES:
            _write_json(self.legacy_state / name, base)
        _write_json(self.legacy_state / self.DRAFT, {**base, "purpose_call_budgets": {
            "mission_directed_document_draft": {"max_cost_usd": 1.0}}})
        bootstrap(self.legacy_state, self.legacy_config)
        raw = json.loads(self.legacy_config.read_text())
        state = self.legacy_state
        raw.update({
            "bounded_planner": {"enabled": True, "interval_seconds": 300, "config": {
                "doctrine_pack_version_hash": "a" * 64,
                "doctrine_pack_version_ref": "doctrine-pack-version:default:1",
                "filed_window_days": 400, "max_probes_per_tick": 1,
                "max_response_bytes": 8388608,
                "observation_mandate_version_ref": "mandate-version:default:1",
                "planner_broker_auth_key": str(self.key),
                "planner_broker_client_id": "client:dalton-core",
                "planner_broker_socket": str(self.socket),
                "planner_call_budget": {"max_cost_usd": 1.0, "max_input_tokens": 80000,
                                        "max_output_tokens": 16000, "timeout_seconds": 300},
                "planner_credential_slot_refs": list(self.LEGACY_SLOTS),
                "planner_expected_agent_id": "chem", "planner_max_cost_usd": 0.5,
                "planner_model_router_db": str(state / "model-router.sqlite"),
                "planner_routing_policy_ref": POLICY_REF,
                "scheduler_db": str(state / "scheduler.sqlite"), "timeout_seconds": 60.0,
                "token_config": str(state / "writer-tokens.json"),
                "user_agent": "Dalton canary", "writer_socket": str(state / "run/writer.sock")}},
            "thesis_impact": {"enabled": False, "interval_seconds": 300, "config": {
                "assessment_routing_policy_ref": POLICY_REF,
                "broker_auth_key": str(self.key), "broker_client_id": "client:dalton-core",
                "broker_socket": str(self.socket),
                "budget_db": str(state / "thesis-impact-budget.sqlite"),
                # What template v6 carried: an older version than the lanes.
                "budget_policy_version_id": "thesis-impact-day-budget-policy:production:3",
                "company_thesis_refs": {}, "credential_slot_refs": list(self.LEGACY_SLOTS),
                "day_cap_micros": 25000000, "expected_agent_id": "chem", "max_targets": 25,
                "model_router_db": str(state / "model-router.sqlite"),
                "routing_policy_ref": POLICY_REF,
                "scheduler_db": str(state / "scheduler.sqlite"), "timeout_seconds": 180,
                "token_config": str(state / "writer-tokens.json"),
                "verifier_routing_policy_ref": POLICY_REF,
                "writer_socket": str(state / "run/writer.sock")}},
            "document_extraction": {"max_windows_per_tick": 30, "numeric_windows_per_tick": 10,
                                    "discovery_windows_per_tick": 10},
            "agenda": {"enabled": False, "interval_seconds": 3600, "config": {}},
            "weekly_brief": {"enabled": False, "interval_seconds": 300, "config": {}},
            "outbox": {"enabled": False, "interval_seconds": 60, "config": {}},
            "control": {"enabled": False, "config": {"cockpit": {
                "openclaw_config_path": str(self.openclaw), "state_dir": str(state)}}},
        })
        self.legacy_config.write_text(json.dumps(raw))
        _write_json(state / "document-research-config.json", {
            "schema_version": "document-research-config-0.1",
            "purpose": "mission_directed_document_research", "inventory_preview_chars": 600,
            "spool_dir": str(state / "transcript-spool"), "enabled_sources": ["source:sec-edgar"],
            "policy": {"schema_version": "document-research-policy-0.1",
                       "policy_ref": "document-research-policy:canary:1",
                       "allowed_purposes": ["mission_directed_document_research"],
                       "allowed_access_policy_refs": ["policy:access:public-web"],
                       "max_query_terms": 16, "max_query_term_chars": 240, "max_results": 24,
                       "max_read_chars": 100000, "max_question_chars": 12000,
                       "max_context_before_chars": 2000, "max_context_after_chars": 4000,
                       "content_hash": "a" * 64},
            "source_reading_limits": {"alphaengine_max_document_chars": 10000000,
                                      "public_web_max_source_chars": 10000000,
                                      "public_web_max_pdf_pages": 2000,
                                      "public_web_max_decompressed_bytes": 100000000}})
        # Host inputs the legacy install put inside its own state directory.
        for tool in ("xueqiu", "xreach"):
            path = state / "host-tools" / tool
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("#!/bin/sh\nexit 0\n"); os.chmod(path, 0o755)
            _write_json(state / "credential-grants" / f"{tool}.json", {"grant": tool})
        (state / "feeds" / "prior-research" / "0000000001").mkdir(parents=True)
        _write_json(state / "claim-support-verification.json", {"backfill_batches_per_run": 0})

    def export_templates(self) -> dict:
        import hashlib
        from dalton_core.workspace_model_setup import export_runtime_template
        from dalton_core.workspace_service_setup import export_service_template

        from dalton_core.model_router import ModelRouter
        from dalton_core.thesis_impact_budget import ThesisImpactBudgetStore

        templates = {"model": self.root / "hostfiles" / "model-runtime-v6.json",
                     "service": self.root / "hostfiles" / "service-runtime-v8.json"}
        # A live environment's writer keeps its WAL databases open; the
        # read-only export needs their sidecars to exist, as they do on a host.
        with ModelRouter(self.legacy_state / "model-router.sqlite"), \
                ThesisImpactBudgetStore(self.legacy_state / "thesis-impact-budget.sqlite"):
            export_runtime_template(self.legacy_state, templates["model"])
        export_service_template(self.legacy_config, templates["service"])
        for path in templates.values():
            os.chmod(path, 0o600)
        return {name: {"path": str(path.resolve()),
                       "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
                for name, path in templates.items()}

    def raise_draft_output_after_export(self) -> None:
        """What the owner did on 09-30, after the template had been exported."""

        path = self.legacy_state / self.DRAFT
        config = json.loads(path.read_text())
        config["purpose_call_budgets"] = {"mission_directed_document_draft": {
            "max_cost_usd": 1.0, "max_output_tokens": 16000}}
        _write_json(path, config)

    # -- the creation ------------------------------------------------------------

    def create(self, runtime_templates: dict):
        from dalton_core.workspace import load_workspace_manifest
        from dalton_core.workspace_creation import create_blank_workspace
        from dalton_core.workspace_manager import provision_runtime

        release = self.root / "release"; release.mkdir()
        host_root = self.root / "host"
        body = {"schema_version": "dalton-shared-connection-catalog-0.1", "models": [],
                "sources": [{"allowed_hosts": ["data.sec.gov", "www.sec.gov"],
                             "allowed_operations": ["list_filings", "get_company_facts"],
                             "auth_mode": "none",
                             "capability_id": "capability:dalton:connector:sec-edgar",
                             "connector_ref": "connector:sec-edgar", "credential_slot_refs": [],
                             "id": "connector-profile-template:sec:canary",
                             "transport": {"config_path": None,
                                           "endpoint_ref": "transport:public-http:0.1",
                                           "kind": "connector", "socket_path": None}}]}
        catalog = {**body, "content_hash": content_hash(body)}
        catalog_path = self.root / "hostfiles" / "catalog.json"
        catalog_path.write_text(json.dumps(catalog))
        manager_path = host_root / "manager.json"
        config = {
            "host_root": str(host_root), "release_path": str(release),
            "release_ref": "release:sha256:" + "d" * 64, "owner_login": "owner@example.com",
            "tailscale_host": "canary.example.ts.net",
            "tailscale_executable": str(self.tailscale),
            "launch_agents_dir": str(self.root / "agents"), "ports": [18961],
            "legacy_workspace": {"name": "现有研究环境", "workspace_id": "legacy",
                                 "status": "running", "url": "https://canary.example/",
                                 "state_dir": str(self.legacy_state),
                                 "config_path": str(self.legacy_config)},
            "runtime_templates": runtime_templates,
            "shared_call_budget_policy_path": str(self.shared_policy),
            "shared_daily_budget_policy_path": str(self.daily_policy),
        }
        _write_json(manager_path, config)
        os.chmod(manager_path, 0o600)
        from dalton_core.workspace_manager import _config
        config = _config(manager_path)  # the manager schema accepts the legacy paths
        shared = [release, self.socket, self.key, self.openclaw, self.tailscale,
                  self.shared_policy, self.daily_policy, manager_path, catalog_path,
                  *(Path(item["path"]) for item in runtime_templates.values())]
        receipt = create_blank_workspace(
            host_root, "ws-canaryinherit", 18961, config["release_ref"], release,
            request_id="canary-inherit", display_name="Canary inheritance",
            shared_readonly_paths=tuple(str(Path(path).resolve()) for path in shared),
            connection_catalog={"path": str(catalog_path.resolve()),
                                "content_hash": catalog["content_hash"]})
        manifest = Path(receipt["manifest_path"])
        ready = provision_runtime(config, manifest, login="owner@example.com",
                                  manager_config_path=manager_path.resolve())
        return load_workspace_manifest(manifest), ready, config

    # -- the test ---------------------------------------------------------------------

    def test_a_created_workspace_has_parity_with_the_mature_environment(self) -> None:
        from dalton_core.claim_support_verification import SETTINGS_FILENAME
        from dalton_core.sec_company_facts_lane import lane_slug_for_state
        from dalton_core.service import ServiceConfig
        from dalton_core.workspace_health_parity import Environment, check_governance
        from dalton_core.workspace_host_scheme import audit_host_scheme, source_credential_slots
        from dalton_core.workspace_lane_parity import plan_parity_actions, resolve_host_sources
        from dalton_core.workspace_process import render_publication_worker

        self.build_legacy()
        templates = self.export_templates()
        self.raise_draft_output_after_export()
        workspace, ready, config = self.create(templates)
        state = workspace.state_dir

        # a3: credential slots -- every list names every host slot.
        for slot in self.LEGACY_SLOTS:
            self.assertIn(slot, source_credential_slots(state))
        rows = {row["key"]: row for row in audit_host_scheme(
            state, source_state_dir=self.legacy_state)}
        self.assertEqual(rows["credential_slots"]["status"], "ok", rows["credential_slots"])
        self.assertEqual(rows["shared_daily_budget"]["status"], "ok", rows["shared_daily_budget"])

        # a2: the per-call budgets, including the one raised after the export.
        self.assertEqual(rows["call_budgets"]["status"], "ok", rows["call_budgets"])
        draft = json.loads((state / self.DRAFT).read_text())
        self.assertEqual(draft["purpose_call_budgets"]["mission_directed_document_draft"]
                         ["max_output_tokens"], 16000)
        self.assertEqual(ready["host_scheme"]["call_budgets"]["status"], "aligned")

        # b3: one day cap -- the lineage head -- in every lane and in service.json.
        pins = {json.loads(path.read_text())["budget_policy_ref"]
                for path in state.glob("*model-config.json")}
        self.assertEqual(pins, {self.head})
        service = json.loads(workspace.config_path.read_text())
        self.assertEqual(service["thesis_impact"]["config"]["budget_policy_version_id"],
                         self.head)
        self.assertEqual(len(list(state.glob("*model-config.json"))), 22)

        # c2: crowd is wired from the legacy state; prior research is this
        # workspace's own empty directory (owner instruction, 2026-09-30).
        for name in ("host-tools/xueqiu", "host-tools/xreach",
                     "credential-grants/xueqiu.json", "credential-grants/xreach.json"):
            self.assertTrue((state / name).exists(), name)
        prior = state / "feeds" / "prior-research"
        self.assertTrue(prior.is_dir() and not prior.is_symlink())
        self.assertEqual(list(prior.iterdir()), [])
        self.assertFalse(prior.resolve().is_relative_to(self.legacy_state.resolve()))
        host_lanes = ready["runtime"]["host_lane_inputs"]
        self.assertIn("source:prior-research", host_lanes["source_refs"])
        self.assertIn("source:xueqiu", host_lanes["source_refs"])
        self.assertNotIn("feeds/prior-research", host_lanes["borrowed_from_source_state"])
        # Every connector permission is approved at creation, the HKEX
        # buy-back tape and cn-hk-findata included.
        approved = {path.name for path in (state / "connector-governance").iterdir()
                    if json.loads(path.read_text()).get("status") == "approved"}
        self.assertIn("hkex-filings-daily-buyback-tape-v1.json", approved)
        self.assertEqual(len([name for name in approved if name.startswith("cn-hk-findata")]), 6)
        self.assertEqual(len([name for name in approved if name.startswith("ir-page-watch")]), 2)

        # c3: the findata and ir-page-watch contracts are approved by this owner.
        governance = {path.name for path in (state / "connector-governance").iterdir()}
        self.assertTrue(any(name.startswith("ir-page-watch") for name in governance), governance)

        # #21 (2026-10-07): approvals alone do not start the IR-page watcher.
        # The pages are the owner's to declare, so a new workspace has no
        # ``ir-pages.json``, the ownership lane emits no declaration flag, and
        # the lane's tick says why in Chinese instead of reading anything.
        from types import SimpleNamespace as _NS

        from dalton_core.ir_page_watch_core import IR_PAGES_ABSENT_REASON
        from dalton_core.lane_registry import LaunchAgentContext
        from dalton_core.mission_ownership_lane import (
            MissionOwnershipLaneCoordinator, argv_fragment as ownership_argv,
            ir_watch_absent_reason,
        )

        self.assertFalse((state / "ir-pages.json").exists())
        self.assertNotIn("--ir-page-declaration",
                         ownership_argv(LaunchAgentContext(state=state)))
        reason = ir_watch_absent_reason(
            _NS(), _NS(ir_declaration_path=None, ir_governance_dir=None))
        self.assertEqual(reason, IR_PAGES_ABSENT_REASON)
        lane = MissionOwnershipLaneCoordinator(
            connection=None, state_dir=state, launcher=None, mission=lambda: None,
            ir_watch=None, ir_watch_reason=reason)
        watched = lane._watch_ir_pages({})
        self.assertEqual(watched["status"], "unconfigured")
        self.assertIn("ir-pages.json", watched["reason"])

        # F14 (2026-10-07): every lane this workspace runs reaches the failure
        # ledger -- its own budget or the writer's observer -- so the cockpit
        # failure panel is not blind to any of them, with nothing to install.
        from dalton_core.cockpit_plane import REGISTRY_LANE_LABELS, _failure_ledger_coverage
        from dalton_core.lane_registry import failure_ledger_coverage

        coverage = failure_ledger_coverage()
        self.assertTrue(coverage)
        self.assertTrue(all(row["mode"] in {"own", "observed"} for row in coverage.values()))
        self.assertTrue(all(row["ledger_lane"] in REGISTRY_LANE_LABELS
                            for row in coverage.values()))
        self.assertEqual(_failure_ledger_coverage()["uncovered"], [])
        self.assertTrue(any("findata" in name or name.startswith("cn-hk") for name in governance),
                        governance)

        # The claim-support settings are the host's.
        self.assertEqual(json.loads((state / SETTINGS_FILENAME).read_text()),
                         {"backfill_batches_per_run": 0})

        # d2: a weekly brief coordinator is configured and waiting for a plan.
        self.assertTrue(service["weekly_brief"]["enabled"])
        self.assertEqual(service["weekly_brief"]["config"]["plan_path"],
                         str(state / "weekly-brief-plan.json"))
        ServiceConfig.from_file(workspace.config_path)

        # h2: its own publication worker, every path in its own state.
        rendered = render_publication_worker(workspace, self.root / "agents")
        worker = json.loads((state / "research-publication-worker-config.json").read_text())
        for key in ("core_db", "model_config", "work_dir", "output_directory"):
            self.assertTrue(Path(worker[key]).is_relative_to(state), key)
        self.assertTrue(any("publication-worker" in str(path) for path in rendered.values()))

        # e2: the SEC lane is named after this workspace.
        self.assertEqual(lane_slug_for_state(state), "ws-canaryinherit-sec-lane")

        # d1: before the first mission the baseline rows are information.
        env = Environment("canary", state, service_config=workspace.config_path)
        try:
            governance_rows = check_governance(env)
        finally:
            env.close()
        self.assertEqual([row["check"] for row in governance_rows if row["status"] == "pending"],
                         ["mission"])

        # Item 4: the template was exported before the draft-output change; that
        # field is realigned at creation, so the template still reads fresh.
        self.assertEqual(ready["template_freshness"]["status"], "fresh",
                         ready["template_freshness"])

        # And nothing is left for the repair script to do.
        leftover = plan_parity_actions(
            state, actor_ref=OWNER,
            host_sources=resolve_host_sources(source_state_dir=self.legacy_state),
            source_state_dir=self.legacy_state)
        self.assertEqual([(action.kind, action.target) for action in leftover], [])


class NewWorkspaceOwnerLeverTests(unittest.TestCase):
    """2026-10-06 (vision-gap #11/#12): the owner's levers work in a created workspace.

    Same code as every other environment and nothing configured for them: the
    receipts table appears on first use, the controller, the lane child
    launcher, the day ledger and the admission authorities obey the fold, the
    inspection reads it, and a coverage-list change publishes a mission
    version and the plans that follow from it.
    """

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace

    def writer(self, state: Path):
        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.writer_server import LaneShortCircuit
        from dalton_core.macos_launchagent import (
            _alphaengine_discovery_plan, _sec_discovery_plan,
        )
        from dalton_core.research_question_backlog import ResearchQuestionBacklog
        from dalton_core.writer_server import WriterServer

        store = DaltonStore(state / "core.sqlite")
        self.addCleanup(store.close)
        server = object.__new__(WriterServer)
        server.db_path = state / "core.sqlite"
        server._store = store
        server._coverage_mission = CoverageMissionAuthority(store)
        server._backlog = ResearchQuestionBacklog(store)
        server._lane_launchers = {}
        server.lane_holds = LaneShortCircuit()
        server._lane_deadline = None
        server._sec_company_resolver_identity = None
        server.sec_ticker_resolver = lambda ticker: dict(ISSUERS[ticker])
        server._discovery_plan_path = str(_alphaengine_discovery_plan(state))
        server._web_search_plan_path = None
        server._sec_filings_plan_path = str(_sec_discovery_plan(state))
        server._search_launcher = server._web_search_launcher = None
        server._sec_filings_launcher = None
        return server

    def test_pause_stop_prioritise_and_change_the_list_in_a_new_workspace(self) -> None:
        from dalton_core.coverage_list_change import coverage_selected_plan
        from dalton_core.lane_child_launcher import LaneChildOwnerPaused
        from dalton_core.mission_source_discovery import (
            ALPHAENGINE_SOURCE_REF, SEC_SOURCE_REF, load_discovery_plan,
        )
        from dalton_core.owner_levers import (
            OwnerLeverHold, assert_admission_open, lever_state, lever_state_for_state_dir,
            spawn_refusal, tick_summary,
        )
        from dalton_core.thesis_impact_budget import (
            OwnerEmergencyStopRefused, ThesisImpactBudgetStore,
        )
        from dalton_core.tick_ledger import TickLedger
        from dalton_core.workspace_health_parity import Environment, check_owner_levers

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        ISSUERS["NEWC"] = {"ticker": "NEWC", "cik": "9900003", "name": "Newco Actuators Inc."}
        self.addCleanup(ISSUERS.pop, "NEWC", None)
        server = self.writer(state)

        class Lane:
            def __init__(self, operation, handler):
                self.operation, self.handler, self.init_kwarg = operation, handler, None

        seen: list[list[str]] = []

        def handler(srv, _params):
            active = srv.coverage_mission.active_mission(mission["mission_ref"])
            seen.append([member["company_ref"] for member in active["universe"]])
            return {"status": "idle"}

        # Nothing pulled yet: a new workspace runs exactly as before.
        self.assertEqual(lever_state_for_state_dir(state)["receipt_count"], 0)
        with patch("dalton_core.writer_server.lane_for_operation", return_value=None):
            self.assertEqual(server._run_lane(Lane("dispatch_mission_tracking", handler), {})["status"], "idle")
        self.assertEqual(sorted(seen[-1]), [GZMO, WDGT])

        # Company pause with a TTL: the lane does not see WDGT, no child starts for it.
        until = (datetime.now(timezone.utc) + timedelta(hours=2)).isoformat()
        out = server._op_apply_owner_lever({
            "lever": "pause", "scope_kind": "company", "scope_ref": WDGT,
            "reason": "earnings next week", "until": until, "actor_ref": OWNER})
        self.assertEqual(out["receipt"]["outcome"], "applied")
        with patch("dalton_core.writer_server.lane_for_operation", return_value=None):
            server._run_lane(Lane("dispatch_mission_tracking", handler), {})
        self.assertEqual(seen[-1], [GZMO])
        # ...but the authority itself still reads the stored record.
        self.assertEqual(len(server.coverage_mission.active_mission(mission["mission_ref"])["universe"]), 2)
        self.assertIn("paused by owner until", spawn_refusal(state, {"company_ref": WDGT}))
        self.assertIsNone(spawn_refusal(state, {"company_ref": GZMO}))
        with self.assertRaisesRegex(OwnerLeverHold, "paused by owner"):
            assert_admission_open(server.store.connection, WDGT)
        assert_admission_open(server.store.connection, GZMO)
        # The TTL expires by itself: three hours later nothing is paused.
        later = lever_state(server.store.connection, now=datetime.now(timezone.utc) + timedelta(hours=3))
        self.assertEqual(later["companies"], {})
        self.assertEqual(later["expired"][0]["scope_ref"], WDGT)

        # Lane pause: the lane answers paused, the tick ledger says so.
        server._op_apply_owner_lever({
            "lever": "pause", "scope_kind": "lane", "scope_ref": "dispatch_mission_tracking",
            "reason": "noisy", "actor_ref": OWNER})
        with patch("dalton_core.writer_server.lane_for_operation", return_value=None):
            paused = server._run_lane(Lane("dispatch_mission_tracking", handler), {})
        self.assertTrue(paused["status"].startswith("paused:paused by owner until resumed"))

        # Emergency stop: every lane, every paid call, every admission refused.
        server._op_apply_owner_lever({"lever": "emergency_stop", "reason": "runaway spend",
                                      "actor_ref": OWNER})
        with patch("dalton_core.writer_server.lane_for_operation", return_value=None):
            stopped = server._run_lane(Lane("dispatch_mission_statements", handler), {})
        self.assertEqual(stopped["reason_code"], "owner_emergency_stop")
        with ThesisImpactBudgetStore(state / "thesis-impact-budget.sqlite") as budget:
            with self.assertRaisesRegex(OwnerEmergencyStopRefused, "owner_emergency_stop"):
                budget.admit(policy_version_id="p", day="2026-10-06", work_order_ref="w",
                             attempt_number=1, phase="model", route_decision_ref="r",
                             reserved_micros=1)
        with self.assertRaises(OwnerLeverHold):
            assert_admission_open(server.store.connection, GZMO)
        self.assertIn("no new lane children", spawn_refusal(state, {"company_ref": GZMO}))
        summary = tick_summary(server._op_owner_lever_state({}))
        self.assertTrue(summary["emergency_stop"])
        with TickLedger(state / "tick-ledger.sqlite") as ledger:
            ledger.append_tick({"status": "idle", "active_loop_count": 0, "probes_executed": 0,
                                "executed": [], "skipped": [], "owner_levers": summary},
                               started_at=datetime.now(timezone.utc))
        ticks = sqlite3.connect(state / "tick-ledger.sqlite")
        status = ticks.execute("SELECT status FROM tick_ledger_lanes WHERE driver_key='owner_levers'").fetchone()[0]
        ticks.close()
        self.assertIn("emergency stop by owner until resumed", status)
        # The inspection reads it, and a restart does not lose it.
        env = Environment("canary", state, service_config=workspace.config_path)
        try:
            rows = {row["check"]: row for row in check_owner_levers(env)}
        finally:
            env.close()
        self.assertEqual(rows["emergency_stop"]["status"], "warn")
        server.store.close()
        server = self.writer(state)
        self.assertIsNotNone(server._owner_lever_state()["emergency"])
        server._op_apply_owner_lever({"lever": "resume_all", "reason": "fixed", "actor_ref": OWNER})
        self.assertIsNone(server._owner_lever_state()["emergency"])
        with self.assertRaises(Exception):
            server._op_apply_owner_lever({"lever": "pause", "scope_kind": "company",
                                          "scope_ref": WDGT, "reason": "x",
                                          "actor_ref": "automation:coverage-mission"})

        # Stop a research line: an unknown question is refused, and the refusal is a receipt.
        stop = server._op_stop_research_line({"line_ref": "research-question:none", "reason": "dead end",
                                              "actor_ref": OWNER})
        self.assertEqual(stop["status"], "refused")
        self.assertEqual(stop["receipt"]["lever"], "stop_line")

        # Prioritise: a TrackingCadenceVersion per adjustable source, by the owner.
        first = server._op_prioritize_company({"company_ref": GZMO, "reason": "catalyst",
                                               "actor_ref": OWNER})
        self.assertEqual(first["status"], "published")
        # 2026-10-07: the same lever re-times discovery for this company only
        # -- plan intervals ÷4, never under a day -- with no plan file touched.
        from dalton_core.tracking_cadence import (
            discovery_cadence_divisor, scaled_interval_days,
        )

        self.assertEqual(discovery_cadence_divisor(server.store, GZMO)["divisor"], 4)
        self.assertEqual(discovery_cadence_divisor(server.store, WDGT)["divisor"], 1)
        self.assertEqual(scaled_interval_days(30, 4), 7)
        self.assertEqual(scaled_interval_days(2, 4), 1)
        sources = {item["source_key"] for item in first["published"]}
        self.assertIn("alphaengine", sources)
        self.assertNotIn("sec", sources)  # fixed cadences are never touched
        rows = server.store.connection.execute(
            "SELECT actor_ref, change_reason FROM tracking_cadence_versions WHERE company_ref=?",
            (GZMO,)).fetchall()
        self.assertTrue(rows and all(tuple(r) == (OWNER, "human_revision") for r in rows))
        back = server._op_prioritize_company({"company_ref": GZMO, "mode": "normal",
                                              "reason": "done", "actor_ref": OWNER})
        self.assertEqual(back["status"], "published")
        self.assertEqual(discovery_cadence_divisor(server.store, GZMO)["divisor"], 1)

        # Coverage list: draft (SEC issuer, plans, budget), then publish.
        old_alpha = load_discovery_plan(server._discovery_plan_path)
        draft = server._op_draft_coverage_change({
            "mission_ref": mission["mission_ref"], "add": ["NEWC"], "remove": [WDGT],
            "actor_ref": OWNER})
        added = draft["additions"][0]
        self.assertEqual((added["cik"], added["name"], added["company_ref"]),
                         ("0009900003", "Newco Actuators Inc.", "company:ticker:newc"))
        by_source = {change["source_ref"]: change for change in draft["plans"]}
        self.assertEqual(by_source[ALPHAENGINE_SOURCE_REF]["added"], ["company:ticker:newc"])
        self.assertEqual(by_source[SEC_SOURCE_REF]["removed"], [WDGT])
        self.assertTrue(draft["budget"]["caps_unchanged"])
        published = server._op_publish_coverage_change({
            "mission_ref": mission["mission_ref"], "expected_mission_hash": mission["content_hash"],
            "add": [{"ticker": "NEWC", "cik": added["cik"], "name": added["name"],
                     "coverage_tier": "B"}],
            "remove": [WDGT], "reason": "swap WDGT for NEWC", "actor_ref": OWNER,
            "idempotency_key": "canary-coverage"})
        self.assertEqual(published["plan_errors"], [])
        active = server.coverage_mission.active_mission(mission["mission_ref"])
        self.assertEqual(active["version"], mission["version"] + 1)
        self.assertEqual(sorted(m["company_ref"] for m in active["universe"]),
                         ["company:ticker:gzmo", "company:ticker:newc"])
        alpha = load_discovery_plan(coverage_selected_plan(state, ALPHAENGINE_SOURCE_REF))
        sec = load_discovery_plan(coverage_selected_plan(state, SEC_SOURCE_REF))
        self.assertEqual(alpha["companies"]["company:ticker:newc"]["ticker"], "NEWC")
        self.assertNotIn(WDGT, alpha["companies"])
        self.assertEqual(alpha["specs"], old_alpha["specs"])  # query forms untouched
        self.assertEqual(sec["companies"]["company:ticker:newc"], {"cik": "0009900003"})
        receipt = published["receipt"]
        self.assertEqual((receipt["lever"], receipt["actor_ref"]), ("coverage_change", OWNER))
        self.assertEqual(receipt["refs"]["mission_version_ref"], active["id"])
        # A stale page cannot publish over it.
        from dalton_core.errors import IdempotencyConflict
        with self.assertRaises(IdempotencyConflict):
            server._op_publish_coverage_change({
                "mission_ref": mission["mission_ref"],
                "expected_mission_hash": mission["content_hash"], "add": [],
                "remove": [GZMO], "reason": "x", "actor_ref": OWNER})
        # Every lever left a receipt the inspection can read.
        levers = [r["lever"] for r in server.owner_levers.receipts(limit=50)]
        for lever in ("pause", "emergency_stop", "resume_all", "stop_line", "prioritize",
                      "coverage_change"):
            self.assertIn(lever, levers)


class NewWorkspaceScenarioTests(unittest.TestCase):
    """ADR-0012: structured scenarios in a created workspace, with nothing configured.

    The mission a new workspace publishes already grants the three scopes the
    scenario lanes use (dossier, observation, market_event); the tables
    appear on first use; the proposal lane is on because the workspace holds
    a model configuration; and an owner decision, an observation tick and a
    probability revision all run through the same code as legacy.
    """

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace
    writer = NewWorkspaceOwnerLeverTests.writer

    def test_scenarios_run_in_a_new_workspace_with_nothing_extra(self) -> None:
        from types import SimpleNamespace

        from dalton_core.company_scenario import CompanyScenarioAuthority, in_force
        from dalton_core.mission_scenario_lane import (
            build_proposal_launcher, dispatch_observation,
        )
        from dalton_core.scenario_proposal import eligible_companies
        from dalton_core.tracking_cadence import TrackingCadenceAuthority

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        for scope in ("dossier", "observation", "market_event"):
            self.assertIn(scope, mission["autonomy"]["may_write"])
        server = self.writer(state)
        server.lane_state = {}
        # Nothing yet: the observation lane is idle and reads no table it lacks.
        self.assertEqual(dispatch_observation(server, {})["status"], "idle")
        self.assertEqual(eligible_companies(server.store.connection, mission,
                                            now=datetime.now(timezone.utc),
                                            min_repropose_days=7), [])
        launcher = build_proposal_launcher(SimpleNamespace(db=str(state / "core.sqlite")))
        if (state / "initial-screen-model-config.json").is_file() or \
                (state / "dossier-model-config.json").is_file():
            self.assertIsNotNone(launcher)
            launcher.close()
        # A proposal (as the lane's child would write it), decided by the owner.
        from tests.finance_scenario_fixture import seed_held_model
        from tests.test_market_price_authority import bar
        from tests.test_valuation_snapshot import roles, history
        from dalton_core.market_price import MarketPriceSeriesAuthority
        from dalton_core.valuation_snapshot import ValuationSnapshotAuthority
        from dalton_core.scenario_rules import partition_conditions
        today = datetime.now(timezone.utc).date()
        model, _specification, _inputs = seed_held_model(server.store, server._coverage_mission, mission,
            company_ref=WDGT, ticker="WDGT", cik="0009900001", entity_name="Widget Robotics Inc.")
        price = MarketPriceSeriesAuthority(server.store).publish_series(company_ref=WDGT, ticker="WDGT", currency="USD",
            bars=[bar(today.isoformat(), "200", open="198", high="202", low="197")],
            invocation_ref="connector-invocation:synthetic-price", artifact_hash="a" * 64,
            governance_ref="governance:synthetic", governance_hash="b" * 64,
            requested_start=today.isoformat(), requested_end=today.isoformat(), captured_at=today.isoformat()+"T00:00:00+00:00")
        p = {"version_ref":price["id"],"version_hash":price["content_hash"],"bar_date":today.isoformat(),"close":"200","currency":"USD",
            "invocation_ref":"connector-invocation:synthetic-price","artifact_hash":"a"*64}
        shares = {k:v for k,v in p.items() if k in ("version_ref","version_hash","invocation_ref","artifact_hash")}
        shares.update(as_of=today.isoformat(), shares_outstanding="1000000")
        valuation = ValuationSnapshotAuthority(server.store).publish_snapshot(company_ref=WDGT, price=p, shares=shares,
            fundamental_windows=[{"as_of":"2025-12-31","roles":roles()}], price_history=history(last=today.isoformat()),
            share_history=[{"as_of":"2020-01-01","shares_outstanding":"1000000"}])
        partition={"observed_by":"sec_actual","metric_ref":"concept:us-gaap:Revenues","transform":"yoy_pct","cutpoints":["5","15"],"source_refs":[model["id"]]}
        def scenario(key, probability, delta):
            return {"key":key,"name":key,"narrative":"Synthetic dated quantity scenario","probability":probability,
                "evidence_refs":[model["id"]],"partition":partition,"triggers":partition_conditions(key,partition,today=today),
                "probability_basis":{"kind":"template_prior","because":"Explicit synthetic template prior","source_refs":[model["id"]],"as_of":today.isoformat()},
                "driver_deltas":[{"driver_ref":"economic:volume","mode":"relative_level","delta":delta,"because":"Synthetic quantity variation","refs":[model["id"]]}]}

        proposal = CompanyScenarioAuthority(server.store).propose_set(
            company_ref=WDGT, scenarios=[scenario("bear", "0.25", "-.1"), scenario("base", "0.5", "0"), scenario("bull", "0.25", ".1")],
            anchors={"model_version_ref":model["id"],"model_version_hash":model["content_hash"],"valuation_snapshot_ref":valuation["id"],"price_as_of":today.isoformat()}, proposer=None, mission=mission,
            actor_ref=mission["autonomy"]["automation_principal"])
        decided = server._op_decide_company_scenario_set({
            "set_ref": proposal["id"], "set_hash": proposal["content_hash"],
            "decision": "accept", "reason": "looks right", "actor_ref": OWNER})
        self.assertEqual(decided["decision"], "accept")
        observed = dispatch_observation(server, {})
        # The middle interval has both a lower and an upper condition.
        self.assertEqual((observed["status"], observed.get("appended")), ("observed", 4), observed)
        live = in_force(server.store.connection, WDGT)
        bear = next(item for item in live if item["key"] == "bear")
        from dalton_core.company_scenario import CompanyScenarioValidationError
        blank_changes=[{"scenario_ref":row["scenario_ref"],"version_ref":row["id"],"version_hash":row["content_hash"],
            "probability":{"bear":"0.7","base":"0.2","bull":"0.1"}[row["key"]]} for row in live]
        with self.assertRaisesRegex(CompanyScenarioValidationError,"evidence rationale"):
            CompanyScenarioAuthority(server.store).revise_group(company_ref=WDGT,changes=blank_changes,
                reason="  ",actor_ref=OWNER,mission=mission)
        self.assertEqual(in_force(server.store.connection,WDGT),live)
        revised = server._op_revise_company_scenario_group({
            "company_ref":WDGT,"changes":[{"scenario_ref":row["scenario_ref"],"version_ref":row["id"],"version_hash":row["content_hash"],
                "probability":{"bear":"0.7","base":"0.2","bull":"0.1"}[row["key"]]} for row in live],
            "reason": "bookings weak", "actor_ref": OWNER})
        self.assertEqual(next(row["probability"] for row in revised["versions"] if row["key"]=="bear"), "0.70")
        self.assertTrue(revised["cadence"]["published"])
        cadence = TrackingCadenceAuthority(server.store).latest(WDGT, "alphaengine")
        self.assertEqual((cadence["decision"], cadence["actor_ref"]),
                         ("scenario_probability", OWNER))
        current=in_force(server.store.connection,WDGT)
        reset=CompanyScenarioAuthority(server.store).revise_group(company_ref=WDGT,changes=[{
            "scenario_ref":row["scenario_ref"],"version_ref":row["id"],"version_hash":row["content_hash"],
            "probability":{"bear":".25","base":".5","bull":".25"}[row["key"]]} for row in current],
            reason="Return to the explicit template prior while awaiting new evidence",actor_ref=OWNER,mission=mission)
        self.assertEqual(len(reset["versions"]),3)
        self.assertTrue(all(row["probability_basis"]["kind"]=="template_prior" for row in reset["versions"]))


class NewWorkspaceMarketProxyAndDisclosureTests(unittest.TestCase):
    """2026-10-07 slices #15-#17: a created workspace needs nothing to get them.

    Market-proxy derivation, the SEC-filing disclosure reader (10-K Item 5
    buybacks, Item 5 / 9B trading arrangements) and the proxy tagging rule run
    on the same code as legacy: the first mission already grants every scope
    they write under, the seeded proxy file leaves derivation on, and on a
    workspace with nothing fetched yet each one reads nothing and fails nothing.
    """

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace

    def test_proxies_and_filing_disclosures_need_no_workspace_setup(self) -> None:
        from dalton_core.acquired_sec_filings import acquired_filing_documents
        from dalton_core.market_proxy_catalog import derive_mappings
        from dalton_core.market_proxy_claim import load_mapping_config
        from dalton_core.tracking_lane_cli import company_events

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        for scope in ("market_price", "claim", "evidence", "claim_index",
                      "market_event", "observation"):
            self.assertIn(scope, mission["autonomy"]["may_write"])
        config = state / "market-proxy-mappings.json"
        if config.is_file():
            self.assertTrue(load_mapping_config(config)["derive_from_model_drivers"])
        with DaltonStore(state / "core.sqlite") as store:
            refs = [item["company_ref"] for item in mission["universe"]]
            self.assertEqual(derive_mappings(store.connection, refs), [])
            for ref in refs:
                found = acquired_filing_documents(store.connection, company_ref=ref,
                                                  state_dir=state)
                self.assertEqual(found["documents"], [])
                events = company_events(store, mission, company_ref=ref,
                                        now=datetime.now(timezone.utc))
                self.assertFalse([event for event in events if event["kind"] in
                                  ("insider_trading_plan", "buyback_disclosure")])


class NextVersionIdTests(unittest.TestCase):
    """The owner scripts could not upgrade a created workspace's first mission."""

    def test_a_hash_named_first_version_gets_the_writers_numbered_name(self) -> None:
        from scripts.sign_auto_commit_rules import _next_version_id

        # ws-7d's first version; the cascade used to refuse it outright.
        self.assertEqual(
            _next_version_id("coverage-mission-version:ws-7d:f85cd96a7f406dd704e8ca5c", 2),
            "coverage-mission-version:ws-7d:2")
        # A hash that happens to end in digits used to become "...a52".
        self.assertEqual(
            _next_version_id("coverage-mission-version:canary:3fde112926c0291adddab215", 2),
            "coverage-mission-version:canary:2")
        self.assertEqual(_next_version_id("policy-17", 18), "policy-18")
        self.assertEqual(_next_version_id("constitution-version:first-mission:ab12:5", 6),
                         "constitution-version:first-mission:ab12:6")


if __name__ == "__main__":
    unittest.main()


class NewWorkspaceOwnerDeliveryTests(unittest.TestCase):
    """Vision-gap #13: a created workspace inherits the delivery contract state.

    Seeded ``proposed`` with no recipient, exactly as install.sh seeds it on
    every other environment; nothing is sent until this workspace's owner
    approves it with a chat id, and then the same code delivers -- here, the
    owner's own emergency stop -- once, through the host tool (a fake), with a
    receipt.
    """

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace

    def test_a_created_workspace_waits_for_its_owner_then_delivers_once(self) -> None:
        from dalton_core import owner_delivery as od
        from dalton_core.lane_registry import LaunchAgentContext
        from dalton_core.mission_owner_delivery_lane import (
            OwnerDeliveryLauncher, argv_fragment,
        )
        from dalton_core.needs_human import collect
        from dalton_core.owner_levers import OwnerLeverAuthority

        workspace, _mission = self.create_workspace()
        state = workspace.state_dir
        seeded = json.loads(od.contract_path(state).read_text())
        self.assertEqual(seeded, od.build_contract_record())
        self.assertEqual(oct(od.contract_path(state).stat().st_mode & 0o777), "0o600")
        self.assertEqual(argv_fragment(LaunchAgentContext(state=state)),
                         ["--owner-delivery-state-dir", str(state)])
        # Unapproved: the lane starts no child, the pass sends nothing, and the
        # to-do list mentions it without calling it blocking.
        spawned = []
        launcher = OwnerDeliveryLauncher(state_dir=state, popen=lambda *a, **k: spawned.append(a))
        self.assertEqual(launcher.dispatch_once()["status"], "not_granted")
        self.assertEqual(spawned, [])
        quiet = [item for item in collect(state_dir=state, environment="canary")["items"]
                 if "飞书推送" in item["title"]]
        self.assertEqual(len(quiet), 1)
        self.assertFalse(quiet[0]["actionable"])

        class Sender:
            sent: list = []

            def send(self, chat_id, body):
                self.sent.append((chat_id, body))
                return "om_canary"

        sender = Sender()
        self.assertEqual(od.run_delivery_pass(state, sender=sender)["status"], "not_granted")
        # The owner approves this environment (the ops script's apply).
        approved_at = datetime.now(timezone.utc)
        with patch("sys.stdout"):
            self.assertEqual(od.main(["approve", "--state-dir", str(state), "--chat-id",
                                      "oc_canary0000000000", "--approved-by", OWNER,
                                      "--apply"]), 0)
        with DaltonStore(state / "core.sqlite") as store:
            OwnerLeverAuthority(store).record(
                lever="emergency_stop", scope_kind="global", scope_ref="global",
                reason="canary stop", actor_ref=OWNER,
                now=approved_at + timedelta(minutes=1))
        later = approved_at + timedelta(minutes=2)
        first = od.run_delivery_pass(state, sender=sender, now=later)
        self.assertEqual(first["status"], "delivered", first)
        self.assertEqual(len(sender.sent), 1)
        self.assertEqual(sender.sent[0][0], "oc_canary0000000000")
        self.assertIn("紧急停机已生效", sender.sent[0][1])
        od.run_delivery_pass(state, sender=sender, now=later + timedelta(hours=1))
        self.assertEqual(len(sender.sent), 1)
        receipts = od.ledger_summary(state)
        self.assertEqual(receipts["counts"], {"attempting": 1, "delivered": 1})


class NextVersionIdTests(unittest.TestCase):
    """The owner scripts could not upgrade a created workspace's first mission."""

    def test_a_hash_named_first_version_gets_the_writers_numbered_name(self) -> None:
        from scripts.sign_auto_commit_rules import _next_version_id

        # ws-7d's first version; the cascade used to refuse it outright.
        self.assertEqual(
            _next_version_id("coverage-mission-version:ws-7d:f85cd96a7f406dd704e8ca5c", 2),
            "coverage-mission-version:ws-7d:2")
        # A hash that happens to end in digits used to become "...a52".
        self.assertEqual(
            _next_version_id("coverage-mission-version:canary:3fde112926c0291adddab215", 2),
            "coverage-mission-version:canary:2")
        self.assertEqual(_next_version_id("policy-17", 18), "policy-18")
        self.assertEqual(_next_version_id("constitution-version:first-mission:ab12:5", 6),
                         "constitution-version:first-mission:ab12:6")


if __name__ == "__main__":
    unittest.main()


class NewWorkspaceConvictionV2Tests(unittest.TestCase):
    """ADR-0013: a created workspace starts on Conviction v2 with no patches."""

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace
    writer = NewWorkspaceOwnerLeverTests.writer

    def test_a_created_workspace_drafts_conviction_calls_under_v2(self) -> None:
        """ADR-0013: Conviction v2 with no policy published and nothing patched.

        The defaults are in code, so a workspace created by the formal flow
        already drafts under v2: the foundation binds the v2 default, the
        mission grants the scope and the checkpoint, the child's precheck asks
        only for an active thesis and a current valuation snapshot -- and says
        in Chinese which one is missing -- and the owner can publish the next
        version through the writer, which an automation principal cannot.
        """

        from types import SimpleNamespace

        from dalton_core.conviction_call_cli import (
            active_gate, gate_inputs_v2, precheck_for, run_conviction_call,
        )
        from dalton_core.conviction_policy import (
            DEFAULT_POLICY, DEFAULT_POLICY_HASH, effective_policy,
        )
        from dalton_core.coverage_mission import CoverageMissionAuthority
        from dalton_core.writer_server import (
            HUMAN_GOVERNANCE_OPERATIONS, MISSION_AUTOMATION_OPERATIONS,
        )

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        foundation = json.loads((state / "research-foundation.json").read_text())
        self.assertEqual(foundation["methods"]["conviction_policy"]["content_hash"],
                         DEFAULT_POLICY_HASH)
        with DaltonStore(state / "core.sqlite") as store:
            self.assertEqual(active_gate(store.connection), "v2")
            policy = effective_policy(store.connection)
            self.assertEqual(policy["source"], "code_default")
            current = CoverageMissionAuthority(store).mission(mission["id"])
            self.assertIn("conviction_call", current["autonomy"]["may_write"])
            self.assertIn("conviction_call", current["autonomy"]["human_checkpoints"])
            gate = precheck_for(gate_inputs_v2(store, WDGT,
                                               industry_ref=current["industry_ref"]),
                                WDGT, policy)
        self.assertFalse(gate["eligible"])
        self.assertEqual(gate["reason_labels"], ["没有生效的投资论点", "还没有估值快照"])
        summary = run_conviction_call(state_dir=state, company_ref=WDGT,
                                      summary_dir=self.root / "conviction-run", dry_run=True)
        self.assertEqual((summary["gate_version"], summary["call_status"]),
                         ("v2", "not_eligible"), summary)
        self.assertIn("还没有估值快照", summary["failure_reason"])
        self.assertIn("publish_conviction_policy", HUMAN_GOVERNANCE_OPERATIONS)
        self.assertNotIn("publish_conviction_policy", MISSION_AUTOMATION_OPERATIONS)
        server = self.writer(state)
        body = json.loads(json.dumps(dict(DEFAULT_POLICY)))
        body["max_drafts_per_day"] = 1
        out = server._op_publish_conviction_policy({
            "policy": body, "reason": "canary: one draft a day", "actor_ref": OWNER,
            "prior_version_ref": None, "idempotency_key": "canary-policy"})
        self.assertEqual(out["status"], "published")
        self.assertEqual(effective_policy(server.store.connection)["version"], 1)
        from dalton_core.mission_conviction_lane import argv_fragment

        model_config = state / "initial-screen-model-config.json"
        self.assertEqual(bool(argv_fragment(SimpleNamespace(state=state))),
                         model_config.is_file())


class NextVersionIdTests(unittest.TestCase):
    """The owner scripts could not upgrade a created workspace's first mission."""

    def test_a_hash_named_first_version_gets_the_writers_numbered_name(self) -> None:
        from scripts.sign_auto_commit_rules import _next_version_id

        # ws-7d's first version; the cascade used to refuse it outright.
        self.assertEqual(
            _next_version_id("coverage-mission-version:ws-7d:f85cd96a7f406dd704e8ca5c", 2),
            "coverage-mission-version:ws-7d:2")
        # A hash that happens to end in digits used to become "...a52".
        self.assertEqual(
            _next_version_id("coverage-mission-version:canary:3fde112926c0291adddab215", 2),
            "coverage-mission-version:canary:2")
        self.assertEqual(_next_version_id("policy-17", 18), "policy-18")
        self.assertEqual(_next_version_id("constitution-version:first-mission:ab12:5", 6),
                         "constitution-version:first-mission:ab12:6")


class NewWorkspaceFundTemplateTests(unittest.TestCase):
    """Vision-gap #4: a created workspace exports with the default fund template.

    No patch, no configuration: a Core that has never seen the registry
    resolves every company to the built-in template; the writer seeds the
    built-in row on first use; the owner can register a template and choose
    it for a company through the same doors as every other environment; and
    the export records which template it used and why.
    """

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace
    writer = NewWorkspaceOwnerLeverTests.writer

    def test_a_created_workspace_exports_with_the_default_template(self) -> None:
        from openpyxl import load_workbook

        from dalton_core.company_model_inputs import build_model_inputs
        from dalton_core.fund_xlsx_export import export_fund_workbook
        from dalton_core.fund_xlsx_template_registry import (
            BUILTIN_VERSION_REF, builtin_version, export_binding, resolve_template,
            template_versions,
        )
        from dalton_core.model_forecast_driver import ForecastModelAuthority
        from tests.test_fund_xlsx_template_registry import candidate
        from tests.test_model_forecast_driver import ledger, model, spec

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        member = mission["universe"][0]["company_ref"]
        # Read-only, before anything opened the registry: the built-in.
        reader = sqlite3.connect(f"file:{state / 'core.sqlite'}?mode=ro", uri=True)
        self.addCleanup(reader.close)
        self.assertEqual([BUILTIN_VERSION_REF], [r["id"] for r in template_versions(reader)])
        resolved = resolve_template(reader, member)
        self.assertEqual(("builtin", BUILTIN_VERSION_REF),
                         (resolved["basis"], resolved["template"]["id"]))
        # Export a model with what the new workspace resolved.
        specification = spec(quarters=8)
        inputs = build_model_inputs(ledger(), specification)
        server = self.writer(state)
        forecast = ForecastModelAuthority(server.store).publish(model(specification=specification))
        calendar = {"calendar_ref": "fiscal-calendar:canary", "source_hash": "1" * 64,
                    "as_of": "2026-08-31", "fiscal_year_end_month": 8}
        calendar["content_hash"] = content_hash(calendar)
        out = Path(self.temp.name) / "canary-model.xlsx"
        result = export_fund_workbook(out, model=forecast, spec=specification, inputs=inputs,
                                      calendar_binding=calendar,
                                      template=export_binding(resolved))
        self.assertEqual((BUILTIN_VERSION_REF, "builtin"),
                         (result["template_version_ref"], result["template_basis"]))
        book = load_workbook(out)
        self.assertEqual(["Valuation", "Financials", "Driver", "Sources", "Formula Map"],
                         book.sheetnames)
        # The writer seeds the built-in row the first time it opens the registry.
        self.assertEqual(builtin_version(), server.fund_xlsx_templates.versions()[0])
        rows = server.store.connection.execute(
            "SELECT version_id FROM fund_xlsx_template_versions").fetchall()
        self.assertEqual([BUILTIN_VERSION_REF], [row[0] for row in rows])
        # The owner registers a template and picks it for one company.
        registered = server._op_register_fund_xlsx_template({
            "candidate": candidate(Path(self.temp.name).resolve()), "display_name": "canary",
            "reason": "first upload", "actor_ref": OWNER})["template"]
        server._op_select_fund_xlsx_template({
            "scope_kind": "company", "scope_ref": member, "reason": "try it",
            "actor_ref": OWNER, "template_version_ref": registered["id"],
            "template_version_hash": registered["content_hash"]})
        chosen = resolve_template(server.store.connection, member)
        self.assertEqual(("company", registered["id"]),
                         (chosen["basis"], chosen["template"]["id"]))
        others = [item["company_ref"] for item in mission["universe"][1:]]
        for other in others:
            self.assertEqual("builtin", resolve_template(server.store.connection, other)["basis"])


if __name__ == "__main__":
    unittest.main()


class NewWorkspaceReturnAttributionTests(unittest.TestCase):
    """ADR-0014: income bridges, per-share series and return attribution, nothing configured.

    The mission a created workspace publishes grants ``valuation``; the lane is
    in process (no launcher, no model config); the tables appear on first
    write; and the bridges a filed statement yields in the created workspace
    are the same bytes the pure builders give anywhere else.
    """

    setUp = NewWorkspaceCanaryTests.setUp
    create_workspace = NewWorkspaceCanaryTests.create_workspace
    writer = NewWorkspaceOwnerLeverTests.writer
    ingest_statements = NewWorkspaceCanaryTests.ingest_statements

    def test_return_attribution_runs_in_a_new_workspace_with_nothing_extra(self) -> None:
        from dalton_core.cockpit_return_attribution import card_views
        from dalton_core.lane_registry import lane_for_operation, load_lanes
        from dalton_core.mission_return_attribution_lane import dispatch
        from dalton_core.return_attribution import (
            compute_company, latest_bridges, latest_series, table_exists,
        )
        from tests.test_return_attribution import calendar_company_filings

        workspace, mission = self.create_workspace()
        state = workspace.state_dir
        self.assertIn("valuation", mission["autonomy"]["may_write"])
        load_lanes()
        spec = lane_for_operation("dispatch_return_attribution")
        self.assertEqual((spec.order, spec.init_kwarg, spec.launcher_factory), (98, None, None))
        server = self.writer(state)
        server.lane_state = {}
        self.missions = server.coverage_mission
        # Nothing filed: each company is visited once and nothing is written.
        for _ in range(2):
            out = dispatch(server, {})
            self.assertEqual(out["status"], "unchanged", out)
            self.assertIn("还没有申报报表行", out["reason"])
        if table_exists(server.store.connection, "income_bridge_versions"):
            self.assertEqual(server.store.connection.execute(
                "SELECT COUNT(*) FROM income_bridge_versions").fetchone()[0], 0)
        # WDGT's statements, as the financial-statements lane records them.
        numbered = {}
        for index, (_accession, form, filed, report, lines) in enumerate(calendar_company_filings()):
            numbered[f"0009900001-26-{index + 1:06d}"] = (form, filed, report, [
                {"statement": line.get("statement", "income"), "concept": line["concept"],
                 "label": line["concept"].split(":")[1], "level": 1,
                 "parent_concept": line["parent"], "is_breakdown": False, "dimension_axis": None,
                 "dimension_member": None, "period_start": line["start"], "period_end": line["end"],
                 "value": line["value"], "unit": line["unit"], "balance": line["balance"]}
                for line in lines])
        self.ingest_statements(mission, WDGT, "WDGT", numbered)
        later = datetime.now(timezone.utc) + timedelta(minutes=6)
        outcomes = [dispatch(server, {}, now=later + timedelta(seconds=step)) for step in range(2)]
        published = next(item for item in outcomes if item.get("company_ref") == WDGT)
        self.assertEqual(published["status"], "published", outcomes)
        bridges = latest_bridges(server.store.connection, WDGT)
        self.assertTrue(bridges)
        self.assertEqual({item["status"] for item in bridges}, {"complete"})
        self.assertTrue(all(item["actor_ref"] == AUTOMATION for item in bridges))
        series = latest_series(server.store.connection, WDGT)
        self.assertEqual(series["missing"], [])
        self.assertIn("FY2025 Q4", [item["fiscal_label"] for item in series["quarters_not_stated"]])
        # Same code everywhere: the stored bodies are what the pure builders give.
        recomputed = compute_company(server.store.connection, WDGT)["bridges"]
        stored = [{key: item[key] for key in recomputed[0]} for item in bridges]
        self.assertEqual(sorted(json.dumps(item, sort_keys=True) for item in stored),
                         sorted(json.dumps(item, sort_keys=True) for item in recomputed))
        view = card_views(server.store.connection, [WDGT, GZMO], mission.get("industry_ref"))
        self.assertIn(WDGT, view)
        self.assertNotIn(GZMO, view)
        self.assertIn("个完整勾稽", view[WDGT]["headline"])
