from __future__ import annotations

import json
import unittest
from unittest import mock

from dalton_core.fund_xlsx_template import (
    FundXlsxTemplateError,
    apply_fund_xlsx_template,
    build_fund_xlsx_template_plan,
    load_fund_xlsx_template_style,
)


class FundXlsxTemplateTests(unittest.TestCase):
    def plan(self):
        return build_fund_xlsx_template_plan(
            sheet_names={
                "valuation": "Valuation - DXC",
                "financials": "Financials - DXC",
                "driver": "Driver - DXC",
            },
            annual_periods=["FY2025A", "FY2026E"],
            quarterly_periods=["Q1 FY2026A", "Q2 FY2026E", "Q3 FY2026E"],
            annual_support_headers=["CAGR FY2024–FY2026E"],
            hidden_periods=["FY2025A"],
            unit_labels={"financials": "(USD MM)", "driver": "(USD MM)"},
            row_styles={
                "financials": [
                    {"row": 2, "style": "section", "level": 0},
                    {"row": 3, "style": "label", "level": 1},
                    {"row": 4, "style": "subtotal", "level": 0},
                ],
                "driver": [
                    {"row": 2, "style": "major_section", "level": 0},
                    {"row": 3, "style": "growth_label", "level": 2},
                ],
                "valuation": [
                    {"row": 2, "style": "valuation_section", "level": 0},
                ],
            },
            cell_styles={
                "financials": [
                    {"range": "E3:F3", "style": "cross_sheet_formula",
                     "number_kind": "amount"},
                    {"range": "J3:L3", "style": "local_formula",
                     "number_kind": "amount_one_decimal"},
                    {"range": "J4:L4", "style": "local_formula",
                     "number_kind": "amount"},
                ],
                "driver": [
                    {"range": "F3:F3", "style": "hardcoded_input",
                     "number_kind": "percentage"},
                    {"range": "J3:L3", "style": "assumption_input",
                     "number_kind": "percentage"},
                ],
                "valuation": [
                    {"range": "B3:B3", "style": "hardcoded_input",
                     "number_kind": "per_share"},
                ],
            },
        )

    def test_source_style_contract_is_exact_and_contains_no_company_financial_rows(self):
        style = load_fund_xlsx_template_style()
        self.assertEqual(style["source_proof"]["workbook_sha256"],
                         "545709e06eb0c1452cf74224d92e8ecb70bef4764f316e74595b2f41f87567b0")
        self.assertEqual(style["sheet_order"], ["valuation", "financials", "driver"])
        self.assertEqual(style["model_grid"]["hierarchy_label_columns"], [1, 2, 3, 4])
        wire = json.dumps(style)
        for source_value in ("Online stores", "Retail & Subscription", "2798509", "AMZN"):
            self.assertNotIn(source_value, wire)

    def test_resource_tamper_is_rejected_before_json_is_trusted(self):
        with mock.patch("dalton_core.fund_xlsx_template._resource_bytes",
                        return_value=b'{}\n'):
            with self.assertRaisesRegex(FundXlsxTemplateError, "bytes differ"):
                load_fund_xlsx_template_style()

    def test_reference_period_geometry_keeps_cagr_between_two_gutters(self):
        plan = build_fund_xlsx_template_plan(
            sheet_names={"valuation": "Valuation", "financials": "Financials",
                         "driver": "Driver"},
            annual_periods=[str(year) for year in range(2013, 2027)],
            quarterly_periods=["1Q17"],
            unit_labels={"financials": "(USD MM)", "driver": "(USD MM)"},
        )
        columns = plan["columns"]
        self.assertEqual((columns["annual"][0]["column"],
                          columns["annual"][-1]["column"]), (5, 18))
        self.assertEqual(columns["annual_support_gutter_before"], 19)
        self.assertEqual(columns["annual_support"], [{"label": "CAGR", "column": 20}])
        self.assertEqual(columns["annual_support_gutter_after"], 21)
        self.assertEqual(columns["quarterly"], [{"label": "1Q17", "column": 22}])

    def test_dynamic_apply_plan_preserves_values_formulas_and_applies_source_geometry(self):
        from openpyxl import Workbook

        book = Workbook()
        sources = book.active
        sources.title = "Sources"
        driver = book.create_sheet("Driver - DXC")
        financials = book.create_sheet("Financials - DXC")
        valuation = book.create_sheet("Valuation - DXC")

        financials["A2"] = "Income statement"
        financials["B3"] = "Subscription revenue"
        financials["A4"] = "Revenue"
        financials["E3"] = "='Driver - DXC'!E3"
        financials["F3"] = "='Driver - DXC'!F3"
        financials["J3"] = "=E3+1"
        financials["K3"] = "=F3+1"
        financials["L3"] = "=K3+1"
        financials["J4"] = "=SUM(J3:J3)"
        financials["K4"] = "=SUM(K3:K3)"
        financials["L4"] = "=SUM(L3:L3)"
        driver["A2"] = "Operating drivers"
        driver["C3"] = "Subscription growth"
        driver["F3"] = 0.08
        driver["J3"] = 0.09
        driver["K3"] = 0.10
        driver["L3"] = 0.11
        driver["D1"] = "A/E mixed; * partial"
        financials["D1"] = "A/E mixed; * partial"
        valuation["A2"] = "Valuation"
        valuation["A3"] = "Target price"
        valuation["B3"] = 98.25
        originals = {
            "financials": tuple(financials.cell(3, column).value for column in range(4, 13)),
            "driver": tuple(driver.cell(3, column).value for column in range(4, 13)),
            "valuation": (valuation["A3"].value, valuation["B3"].value),
        }

        apply_fund_xlsx_template(book, self.plan())

        self.assertEqual(book.sheetnames,
                         ["Valuation - DXC", "Financials - DXC", "Driver - DXC", "Sources"])
        self.assertEqual(tuple(financials.cell(3, column).value for column in range(4, 13)),
                         originals["financials"])
        self.assertEqual(tuple(driver.cell(3, column).value for column in range(4, 13)),
                         originals["driver"])
        self.assertEqual((valuation["A3"].value, valuation["B3"].value),
                         originals["valuation"])

        self.assertEqual([financials.cell(1, column).value
                          for column in (5, 6, 7, 8, 9, 10, 11, 12)],
                         ["FY2025A", "FY2026E", None, "CAGR FY2024–FY2026E", None,
                          "Q1 FY2026A", "Q2 FY2026E", "Q3 FY2026E"])
        self.assertEqual(financials.freeze_panes, "E2")
        self.assertFalse(financials.sheet_view.showGridLines)
        self.assertEqual(financials.sheet_format.defaultRowHeight, 11.25)
        self.assertEqual(financials.column_dimensions["A"].width, 1.875)
        self.assertEqual(financials.column_dimensions["B"].width, 2.125)
        self.assertEqual(financials.column_dimensions["C"].width, 2.875)
        self.assertEqual(financials.column_dimensions["D"].width, 24.375)
        self.assertEqual(
            sum(financials.column_dimensions[column].width for column in "ABCD"),
            31.25,
        )
        self.assertEqual(driver.column_dimensions["C"].width, 2.875)
        self.assertEqual(driver.column_dimensions["D"].width, 21.625)
        self.assertEqual(
            sum(driver.column_dimensions[column].width for column in "ABCD"),
            28.5,
        )
        self.assertEqual(financials.column_dimensions["G"].width, 1.375)
        self.assertEqual(financials.column_dimensions["H"].width, 6.375)
        self.assertEqual(financials.column_dimensions["I"].width, 1.125)
        self.assertEqual(driver.column_dimensions["G"].width, 1.125)
        self.assertEqual(driver.column_dimensions["H"].width, 6.625)
        self.assertEqual(driver.column_dimensions["I"].width, 0.875)
        self.assertTrue(financials.column_dimensions["E"].hidden)
        self.assertEqual(financials.row_dimensions[3].outlineLevel, 1)
        self.assertEqual(driver.row_dimensions[3].outlineLevel, 2)
        self.assertEqual(financials["B3"].value, "Subscription revenue")
        self.assertIsNone(financials["D3"].value)
        self.assertEqual(self.plan()["row_styles"]["financials"][1]["label_column"], 2)
        self.assertEqual(self.plan()["row_styles"]["driver"][1]["label_column"], 3)

        self.assertEqual(financials["A1"].fill.fgColor.rgb, "FF3366FF")
        self.assertIn("A1:C1", {str(value) for value in financials.merged_cells.ranges})
        self.assertIn("A1:C1", {str(value) for value in driver.merged_cells.ranges})
        self.assertEqual(financials["A1"].value, "(USD MM)")
        self.assertEqual(driver["A1"].value, "(USD MM)")
        self.assertEqual(financials["D1"].value, "A/E mixed; * partial")
        self.assertEqual(driver["D1"].value, "A/E mixed; * partial")
        self.assertIsNone(financials["A1"].alignment.horizontal)
        self.assertEqual(financials["E1"].alignment.horizontal, "right")
        self.assertEqual(financials["E1"].font.color.rgb, "FFFFFFFF")
        self.assertEqual(financials["E3"].font.color.rgb, "FF008000")
        self.assertEqual(financials["J3"].font.color.rgb, "FF000000")
        self.assertTrue(financials["J4"].font.bold)
        self.assertEqual(driver["F3"].font.color.rgb, "FF0000FF")
        self.assertTrue(driver["F3"].font.italic)
        self.assertEqual(driver["J3"].fill.fgColor.rgb, "FFFFFFC8")
        self.assertEqual(driver["J3"].border.top.style, "hair")
        self.assertEqual(driver["J3"].number_format, "0.0%")
        self.assertEqual(financials["E3"].number_format,
                         "_(#,##0_);\\(#,##0\\);_(?\\-?_);@")
        self.assertEqual(valuation.column_dimensions["A"].width, 14.625)
        self.assertEqual(valuation.column_dimensions["B"].width, 14.375)
        self.assertEqual(valuation.column_dimensions["D"].width, 9.125)
        self.assertEqual(valuation.column_dimensions["E"].width, 9.125)
        self.assertEqual(valuation["A2"].fill.fgColor.rgb, "FF99CCFF")
        self.assertEqual(valuation["B3"].font.color.rgb, "FF0000FF")

    def test_plan_refuses_ambiguous_or_out_of_contract_mappings(self):
        arguments = {
            "sheet_names": {"valuation": "Valuation", "financials": "Financials",
                            "driver": "Driver"},
            "annual_periods": ["FY2025A"],
            "quarterly_periods": ["Q1 FY2026E"],
            "unit_labels": {"financials": "(USD MM)", "driver": "(USD MM)"},
        }
        with self.assertRaisesRegex(FundXlsxTemplateError, "hidden periods"):
            build_fund_xlsx_template_plan(**arguments, hidden_periods=["FY1900A"])
        with self.assertRaisesRegex(FundXlsxTemplateError, "unsupported"):
            build_fund_xlsx_template_plan(
                **arguments,
                row_styles={"financials": [
                    {"row": 2, "style": "copy_amzn_revenue_rows", "level": 0}
                ]},
            )
        with self.assertRaisesRegex(FundXlsxTemplateError, "closed shape"):
            build_fund_xlsx_template_plan(
                **arguments,
                cell_styles={"financials": [
                    {"range": "E2", "style": "hardcoded_input",
                     "number_kind": "amount", "value": 123}
                ]},
            )

        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "Valuation - DXC"
        book.create_sheet("Financials - DXC")
        book.create_sheet("Driver - DXC")
        tampered = self.plan()
        tampered["columns"]["hierarchy_label_columns"] = [1, 2, 3]
        with self.assertRaisesRegex(FundXlsxTemplateError, "plan hash"):
            apply_fund_xlsx_template(book, tampered)


if __name__ == "__main__":
    unittest.main()
