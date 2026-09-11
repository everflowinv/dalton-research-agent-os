from __future__ import annotations

import copy
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from dalton_core.company_model_annual_projection import build_annual_projection
from dalton_core.company_financial_statement_structure import (
    forecast_structure_binding,
    materialize_financial_statement_structure,
)
from dalton_core.fund_xlsx_export import (
    _financial_result_layout,
    export_fund_workbook,
)
from dalton_core.model_forecast_driver import (
    ForecastModelAuthority,
    actualize_model,
    build_structured_forecast_model,
)
from dalton_core.store import DaltonStore, content_hash
from tests.test_company_financial_statement_structure import ACCESSION
from tests.test_structured_cash_flow_companion import authorities


class FundXlsxCashFlowCompanionTests(unittest.TestCase):
    def export(self, *, cash: bool = True):
        inputs, spec, structure, replay, binding = authorities(cash=cash)
        body = build_structured_forecast_model(
            spec, inputs, structure=structure, replay=replay, binding=binding)
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        store = DaltonStore(str(Path(temporary.name) / "core.sqlite"))
        self.addCleanup(store.close)
        model = ForecastModelAuthority(store).publish(body)
        calendar = {
            "calendar_ref": "statement-filing:test", "source_hash": "a" * 64,
            "as_of": "2025-12-31", "fiscal_year_end_month": 12,
        }
        calendar["content_hash"] = content_hash(calendar)
        annual = build_annual_projection(
            model=model, inputs=inputs, calendar_binding=calendar)
        path = Path(temporary.name) / "cash-flow.xlsx"
        exported = export_fund_workbook(
            path, model=model, spec=spec, inputs=inputs,
            calendar_binding=calendar, annual_projection=annual,
        )
        return path, load_workbook(path, data_only=False), model, annual, exported

    def test_cash_companion_uses_distinct_financial_section_and_bound_results(self):
        _path, book, model, annual, exported = self.export()
        financials = book["Financials"]
        driver = book["Driver"]
        rows, section_row = _financial_result_layout(model)
        self.assertIsNotNone(section_row)
        self.assertEqual(financials.cell(section_row, 1).value, "Cash flow statement")
        self.assertEqual(
            financials.cell(section_row, 1).fill.fgColor.rgb[-6:], "99CCFF")
        self.assertIsNone(financials.cell(section_row - 1, 1).value)
        expected = (
            ("result:operating_cash_flow", "Operating cash flow", True),
            ("result:capital_expenditure", "Capital expenditures", False),
            ("result:free_cash_flow", "Free cash flow", True),
        )
        for result_ref, label, bold in expected:
            row = rows[result_ref]
            self.assertEqual(financials.cell(row, 2).value, label)
            self.assertEqual(financials.cell(row, 2).font.bold, bold)
        headers = {
            financials.cell(1, column).value: column
            for column in range(1, financials.max_column + 1)
        }
        forecast_column = headers["1Q26E"]
        ocf_formula = financials.cell(
            rows["result:operating_cash_flow"], forecast_column).value
        capex_formula = financials.cell(
            rows["result:capital_expenditure"], forecast_column).value
        fcf_formula = financials.cell(
            rows["result:free_cash_flow"], forecast_column).value
        self.assertTrue(str(ocf_formula).startswith("='Financials'!"))
        self.assertIn("*'Driver'!", ocf_formula)
        self.assertTrue(str(capex_formula).startswith("='Financials'!"))
        self.assertIn("*'Driver'!", capex_formula)
        self.assertEqual(
            fcf_formula,
            f"='Financials'!{financials.cell(rows['result:operating_cash_flow'], forecast_column).coordinate}"
            f"-SUM('Financials'!{financials.cell(rows['result:capital_expenditure'], forecast_column).coordinate})",
        )
        annual_column = headers["2025"]
        for result_ref, _label, _bold in expected:
            outcome = next(
                row for row in annual["periods"] if row["label"] == "FY2025A"
            )["line_outcomes"][result_ref]
            self.assertEqual(outcome["status"], "computed")
            self.assertTrue(str(financials.cell(
                rows[result_ref], annual_column).value).startswith("=SUM("))
        driver_labels = {
            driver.cell(row, 2).value for row in range(1, driver.max_row + 1)
        }
        self.assertIn("Operating cash flow — actual", driver_labels)
        self.assertIn("Capital expenditure — actual", driver_labels)
        assumption_labels = {
            driver.cell(row, 3).value for row in range(1, driver.max_row + 1)
        }
        self.assertIn(
            "Operating cash flow — Share of line (ratio)", assumption_labels)
        self.assertIn(
            "Capital expenditure — Share of line (ratio)", assumption_labels)
        self.assertEqual(exported["annual_projection_hash"], annual["content_hash"])

    def test_unavailable_cash_companion_stays_blank(self):
        _path, book, model, annual, exported = self.export(cash=False)
        financials = book["Financials"]
        driver = book["Driver"]
        rows, _section = _financial_result_layout(model)
        for result_ref in (
            "result:operating_cash_flow", "result:capital_expenditure",
            "result:free_cash_flow",
        ):
            result = next(item for item in model["results"]
                          if item["ref"] == result_ref)
            row = rows[result_ref]
            self.assertIn("Not available", str(financials.cell(row, 2).value))
            self.assertTrue(all(
                financials.cell(row, column).value is None
                for column in range(5, financials.max_column + 1)
            ))
            self.assertTrue(all(
                period["line_outcomes"][result_ref]["status"] == "unavailable"
                for period in annual["periods"]
            ))
            self.assertTrue(any(result_ref in gap for gap in exported["gaps"]))
        driver_text = "\n".join(
            str(driver.cell(row, column).value or "")
            for row in range(1, driver.max_row + 1)
            for column in (2, 3)
        )
        self.assertNotIn("Operating cash flow — actual", driver_text)
        self.assertNotIn("Capital expenditure — actual", driver_text)

    def test_v03_layout_has_no_cash_flow_section(self):
        inputs, spec, structure, replay, binding = authorities(cash=False)
        spec = dict(spec)
        spec["schema_version"] = "0.3"
        spec.pop("cash_flow_companion")
        body = build_structured_forecast_model(
            spec, inputs, structure=structure, replay=replay, binding=binding)
        rows, section_row = _financial_result_layout(body)
        self.assertIsNone(section_row)
        self.assertEqual(
            rows,
            {result["ref"]: 3 + index
             for index, result in enumerate(body["results"])},
        )

    def test_cash_actual_ratio_formulas_reference_relocated_financial_rows(self):
        inputs, spec, structure, replay, binding = authorities()
        body = build_structured_forecast_model(
            spec, inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            authority = ForecastModelAuthority(store)
            prior = authority.publish(body)
            current = copy.deepcopy(inputs)
            current["periods"].append("2026-03-31")
            actuals = {
                "revenue": 1400, "cost": 840, "opex": 280, "operating": 280,
                "interest_income": 10, "interest_expense": 20, "pretax": 270,
                "tax": 60, "net": 210, "nci": 5, "parent": 205,
                "eps_numerator": 205, "shares": 100, "eps": "2.05",
            }
            for line in current["filed_lines"]:
                line["cells"]["2026-03-31"] = {
                    "period_start": "2026-01-01",
                    "value": str(actuals[line["concept"]]),
                    "unit": next(iter(line["cells"].values()))["unit"],
                    "basis": "reported", "source_accessions": [ACCESSION],
                }
            for source, value in zip(current["cash_flow_inputs"], ("140", "28")):
                source["series"]["quarters"].append({
                    "period_start": "2026-01-01", "period_end": "2026-03-31",
                    "value": value, "unit": "usd", "basis": "reported",
                    "source_accessions": [ACCESSION], "source_forms": ["10-Q"],
                })
            structure, replay = materialize_financial_statement_structure(
                spec, current)
            binding = forecast_structure_binding(structure, replay, current)
            updated = actualize_model(
                prior, current, structure=structure, replay=replay, binding=binding)
            self.assertIsNotNone(updated)
            model = authority.publish(updated)
            calendar = {
                "calendar_ref": "statement-filing:test", "source_hash": "a" * 64,
                "as_of": "2025-12-31", "fiscal_year_end_month": 12,
            }
            calendar["content_hash"] = content_hash(calendar)
            annual = build_annual_projection(
                model=model, inputs=current, calendar_binding=calendar)
            path = Path(temporary) / "cash-flow-actualized.xlsx"
            export_fund_workbook(
                path, model=model, spec=spec, inputs=current,
                calendar_binding=calendar, annual_projection=annual,
            )
            book = load_workbook(path, data_only=False)
        financials, driver = book["Financials"], book["Driver"]
        result_rows, _section = _financial_result_layout(model)
        headers = {
            driver.cell(1, column).value: column
            for column in range(1, driver.max_column + 1)
        }
        actual_column = headers["1Q26"]
        column_letter = financials.cell(1, actual_column).column_letter
        assumption_rows = {
            driver.cell(row, 3).value: row
            for row in range(1, driver.max_row + 1)
        }
        revenue_row = result_rows["result:revenue"]
        for result_ref, label in (
            ("result:operating_cash_flow",
             "Operating cash flow — Share of line (ratio)"),
            ("result:capital_expenditure",
             "Capital expenditure — Share of line (ratio)"),
        ):
            cell = driver.cell(assumption_rows[label], actual_column)
            self.assertEqual(
                cell.value,
                f"='Financials'!{column_letter}{result_rows[result_ref]}/"
                f"'Financials'!{column_letter}{revenue_row}",
            )
            self.assertEqual(cell.font.color.rgb[-6:], "008000")
        self.assertEqual(
            financials.cell(
                result_rows["result:free_cash_flow"], actual_column).value,
            f"='Financials'!{column_letter}{result_rows['result:operating_cash_flow']}"
            f"-SUM('Financials'!{column_letter}{result_rows['result:capital_expenditure']})",
        )
        for role in ("diluted_weighted_average_shares", "diluted_eps"):
            result = next(item for item in model["results"]
                          if item.get("role") == role)
            label = " ".join(
                str(financials.cell(result_rows[result["ref"]], column).value or "")
                for column in range(1, 5)
            )
            self.assertIn("Forecast unavailable", label)


if __name__ == "__main__":
    unittest.main()
