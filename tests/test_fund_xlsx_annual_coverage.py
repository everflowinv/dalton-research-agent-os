from __future__ import annotations

import copy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from dalton_core.company_model_annual_projection import (
    annual_projection_result_refs, build_annual_projection,
)
from dalton_core.company_financial_statement_structure import (
    forecast_structure_binding, materialize_financial_statement_structure,
)
from dalton_core.fund_xlsx_export import (
    FundWorkbookExportError, export_fund_workbook,
    _formula_for, _verify_translated_value,
)
from decimal import Decimal
from dalton_core.model_forecast_driver import (
    ForecastModelAuthority, build_structured_forecast_model,
)
from dalton_core.store import DaltonStore, content_hash
from tests import test_company_financial_statement_structure as financial_fixture
from tests import test_structured_cash_flow_companion as cash_fixture


class AnnualWorkbookCoverageTests(unittest.TestCase):
    def authorities(self, *, forecast="carry_forward_mean"):
        inputs, spec, _structure, _replay, _binding = cash_fixture.authorities()
        concept = "us-gaap:PaymentsToAcquireBusinessesNetOfCashAcquired"
        inputs["filed_lines"].append(financial_fixture.input_line(
            concept, (3, 4, 5, 6), statement="cash"))
        line = financial_fixture.filed(
            "business-acquisitions", "acquisitions", concept, forecast=forecast)
        line["statement"] = "cash"
        spec["financial_statement_structure"]["lines"].append(line)
        structure, replay = materialize_financial_statement_structure(spec, inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            spec, inputs, structure=structure, replay=replay, binding=binding)
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        with DaltonStore(str(Path(temporary.name)/"core.sqlite")) as store:
            model = ForecastModelAuthority(store).publish(body)
        calendar = {"calendar_ref": "filing:another-company", "source_hash": "a"*64,
                    "as_of": "2025-12-31", "fiscal_year_end_month": 12}
        calendar["content_hash"] = content_hash(calendar)
        annual = build_annual_projection(
            model=model, inputs=inputs, calendar_binding=calendar)
        return model, spec, inputs, calendar, annual

    def test_uncovered_cash_line_preserves_quarters_and_leaves_annual_blank(self):
        model, spec, inputs, calendar, annual = self.authorities()
        original = copy.deepcopy(annual)
        result = next(r for r in model["results"] if r["role"] == "acquisitions")
        self.assertEqual(result["status"], "computed")
        self.assertNotIn(result["ref"], annual_projection_result_refs(model))
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/"partial.xlsx"
            exported = export_fund_workbook(
                path, model=model, spec=spec, inputs=inputs,
                calendar_binding=calendar, annual_projection=annual)
            book = load_workbook(path)
            self.addCleanup(book.close)
            sheet = book["Financials"]
            formulas = [
                {"cell": row[0].value, "formula": str(row[2].value).removeprefix("'")}
                for row in book["Formula Map"].iter_rows(min_row=5)
                if row[1].value in {c["ref"] for c in result["cells"]}
            ]
            self.assertEqual(len(formulas), len(result["cells"]))
            for item in formulas:
                self.assertEqual(sheet[item["cell"].split("!", 1)[1]].value,
                                 item["formula"])
            row = sheet[formulas[0]["cell"].split("!", 1)[1]].row
            annual_columns = [column for column in range(1, sheet.max_column+1)
                              if str(sheet.cell(1, column).value).isdigit()]
            self.assertTrue(annual_columns)
            self.assertTrue(all(sheet.cell(row, column).value is None
                                for column in annual_columns))
            self.assertTrue(any(result["ref"] in gap and "暂未生成年度汇总" in gap
                                for gap in exported["gaps"]))
        self.assertEqual(annual, original)

    def test_unavailable_nonincome_line_is_a_gap_without_invented_amount(self):
        model, spec, inputs, calendar, annual = self.authorities(forecast="unavailable")
        result = next(r for r in model["results"] if r["role"] == "acquisitions")
        self.assertEqual(result["status"], "unavailable")
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/"partial.xlsx"
            exported = export_fund_workbook(
                path, model=model, spec=spec, inputs=inputs,
                calendar_binding=calendar, annual_projection=annual)
            book = load_workbook(path)
            self.addCleanup(book.close)
            self.assertFalse(any(row[1].value in {c["ref"] for c in result["cells"]}
                                 for row in book["Formula Map"].iter_rows(min_row=5)))
        self.assertTrue(any(result["ref"] in gap for gap in exported["gaps"]))

    def test_promised_income_and_cash_companion_missing_outcome_still_refused(self):
        model, spec, inputs, calendar, annual = self.authorities()
        covered = annual_projection_result_refs(model)
        income_ref = next(r["ref"] for r in model["results"]
                          if r["role"] == "revenue")
        for missing in (income_ref, "result:operating_cash_flow"):
            self.assertIn(missing, covered)
            malformed = copy.deepcopy(annual)
            for period in malformed["periods"]:
                period["line_outcomes"].pop(missing, None)
            # Isolate the renderer's own promised-coverage guard. The normal
            # preceding validator also rejects this forged annual authority.
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as temporary:
                with self.assertRaises(FundWorkbookExportError):
                    export_fund_workbook(
                        Path(temporary)/"bad.xlsx", model=model, spec=spec, inputs=inputs,
                        calendar_binding=calendar, annual_projection=malformed)
                with patch("dalton_core.company_model_annual_projection.validate_annual_projection",
                           return_value=malformed):
                    with self.assertRaisesRegex(FundWorkbookExportError,
                                               "annual projection lacks result"):
                        export_fund_workbook(
                            Path(temporary)/"bad-renderer.xlsx", model=model, spec=spec,
                            inputs=inputs, calendar_binding=calendar,
                            annual_projection=malformed)

    def test_carried_mean_requires_the_single_exact_verified_amount(self):
        result = {"formula": "another-line[k] = trailing four-quarter mean (carried)",
                  "unit": "usd"}
        cell = {"ref": "result:another@period", "assumption_refs": ["assumption:one"],
                "result_refs": [], "input_cell_refs": [], "value": "-3.25"}
        measures = {"assumption:one": "carry_forward_mean"}
        units = {"assumption:one": "usd"}
        self.assertEqual(_formula_for(result, cell, "period", {},
                                     {"assumption:one": "'Driver'!A5"}, {}, measures,
                                     assumption_units=units),
                         "='Driver'!A5")
        _verify_translated_value(result, cell, {},
                                 {"assumption:one": Decimal("-3.25")}, {}, measures,
                                 assumption_units=units)
        with self.assertRaisesRegex(FundWorkbookExportError, "differs from result"):
            _verify_translated_value(result, cell, {},
                                     {"assumption:one": Decimal("3.25")}, {}, measures,
                                     assumption_units=units)
        for extra in ({"result_refs": [{"ref": "result:other", "period_end": "period"}]},
                      {"input_cell_refs": [{"concept": "other", "period_end": "period"}]},
                      {"assumption_refs": ["assumption:one", "assumption:two"]}):
            with self.subTest(extra=extra):
                with self.assertRaises(FundWorkbookExportError):
                    _formula_for(result, {**cell, **extra}, "period", {},
                                 {"assumption:one": "'Driver'!A5"}, {}, measures,
                                 assumption_units=units)
        self.assertIsNone(_formula_for(result, cell, "period", {},
                                      {"assumption:one": "'Driver'!A5"}, {},
                                      {"assumption:one": "unknown"}))

    def test_held_and_oneoff_amounts_are_exact_units_and_values(self):
        for measure, formula, value in (
                ("hold_last", "balance[k] = balance[last filed] (held)", Decimal("17")),
                ("zero_one_off", "item[k] = 0 (one-off)", Decimal("0"))):
            with self.subTest(measure=measure):
                result = {"formula": formula, "unit": "eur"}
                cell = {"ref": "result:other@period", "value": str(value),
                        "assumption_refs": ["a"], "result_refs": [], "input_cell_refs": []}
                options = {"assumption_units": {"a": "eur"}}
                self.assertEqual(_formula_for(result, cell, "period", {}, {"a": "D5"}, {},
                                             {"a": measure}, **options), "=D5")
                _verify_translated_value(result, cell, {}, {"a": value}, {},
                                         {"a": measure}, **options)
                for modified in ({**cell, "value": str(value+1)},
                                 {**cell, "result_refs": [{"ref": "extra", "period_end": "period"}]}):
                    with self.assertRaises(FundWorkbookExportError):
                        _verify_translated_value(result, modified, {}, {"a": value}, {},
                                                 {"a": measure}, **options)
                with self.assertRaises(FundWorkbookExportError):
                    _formula_for(result, cell, "period", {}, {"a": "D5"}, {},
                                 {"a": measure}, assumption_units={"a": "usd"})

    def test_rate_times_balance_checks_dependencies_units_and_value(self):
        for kind in ("input", "result"):
            with self.subTest(kind=kind):
                result = {"unit": "usd", "formula": (
                    "interest[k] = cash[last filed] * rate[k]" if kind == "input"
                    else "interest[k] = cash[k] * rate[k]")}
                cell = {"ref": "result:interest@period", "value": "2.5",
                        "assumption_refs": ["a"], "input_cell_refs": [], "result_refs": []}
                dependency = {"concept" if kind == "input" else "ref": "cash", "period_end": "period"}
                cell["input_cell_refs" if kind == "input" else "result_refs"] = [dependency]
                options = {"assumption_units": {"a": "ratio"},
                           "history_units": {("cash", "period"): "usd"},
                           "result_units": {"cash": "usd"}}
                self.assertEqual(_formula_for(
                    result, cell, "period", {("cash", "period"): "C4"}, {"a": "D5"},
                    {("cash", "period"): "C4"}, {"a": "rate_times_balance"}, **options),
                    "=C4*D5")
                _verify_translated_value(result, cell, {("cash", "period"): Decimal("100")},
                                         {"a": Decimal(".025")}, {("cash", "period"): Decimal("100")},
                                         {"a": "rate_times_balance"}, **options)
                for modified in ({**cell, "value": "25"}, {**cell, "assumption_refs": ["a", "extra"]}):
                    with self.assertRaises(FundWorkbookExportError):
                        _verify_translated_value(result, modified, {("cash", "period"): Decimal("100")},
                                                 {"a": Decimal(".025")}, {("cash", "period"): Decimal("100")},
                                                 {"a": "rate_times_balance"}, **options)
                with self.assertRaises(FundWorkbookExportError):
                    _formula_for(result, cell, "period", {("cash", "period"): "C4"}, {"a": "D5"},
                                 {("cash", "period"): "C4"}, {"a": "rate_times_balance"},
                                 **{**options, "assumption_units": {"a": "usd"}})
                with self.assertRaises(FundWorkbookExportError):
                    _formula_for({**result, "unit": "eur"}, cell, "period", {}, {"a": "D5"}, {},
                                 {"a": "rate_times_balance"}, **options)


if __name__ == "__main__":
    unittest.main()
