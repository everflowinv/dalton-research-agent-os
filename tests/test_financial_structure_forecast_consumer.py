from __future__ import annotations

import copy
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from dalton_core.company_financial_statement_structure import (
    forecast_structure_binding,
    materialize_financial_statement_structure,
    replay_historical_structure,
    validate_financial_statement_structure,
)
from dalton_core.model_forecast_driver import (
    ForecastModelAuthority,
    ForecastModelUnavailable,
    actualize_model,
    build_structured_forecast_model,
    build_structure_drivers,
    compute_structure_results,
    default_structure_assumptions,
    forecast_periods,
    revenue_anchor,
    revise_assumptions,
)
from dalton_core.forecast_sensitivity import build_projection, measure_series, recompute
from dalton_core.company_model_forecast import publish_forecast_lines
from dalton_core.model_forecast import (
    ModelForecastAuthority,
    STRUCTURED_DRIVER_FORMULA_HASH,
    STRUCTURED_DRIVER_FORMULA_REF,
)
from dalton_core.fund_xlsx_export import export_fund_workbook
from dalton_core.store import DaltonStore, content_hash
from tests.test_company_financial_statement_structure import (
    ACCESSION,
    company_spec,
    financial_inputs,
    proposal,
)


def cells(results):
    return {
        (row["role"], cell["period"]["end"]): cell
        for row in results for cell in row["cells"]
    }


class FinancialStructureForecastConsumerTests(unittest.TestCase):
    def authority(self):
        inputs = financial_inputs()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            if line["ref"] in {"interest-income", "interest-expense"}:
                line.update(forecast_method="share_of_line", forecast_base_ref="revenue")
            elif line["ref"] == "tax":
                line.update(forecast_method="share_of_line", forecast_base_ref="pretax")
            elif line["ref"] == "nci":
                line.update(forecast_method="share_of_line", forecast_base_ref="net")
            elif line["ref"] == "eps-numerator":
                line.update(forecast_method="share_of_line", forecast_base_ref="parent")
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        self.assertTrue(replay["ready_for_forecast"])
        return inputs, structure

    def compute(self, structure, inputs):
        drivers = build_structure_drivers(inputs, structure)
        periods = forecast_periods(revenue_anchor(drivers), 2)
        assumptions = default_structure_assumptions(drivers, periods, structure)
        return compute_structure_results(drivers, assumptions, periods, structure)

    def test_company_formula_controls_nonoperating_sign_and_tax_base(self):
        inputs, structure = self.authority()
        result = self.compute(structure, inputs)
        by_cell = cells(result)
        end = sorted(end for role, end in by_cell if role == "pretax_income")[0]
        operating = Decimal(by_cell[("operating_income", end)]["value"])
        interest_income = Decimal(by_cell[("interest_income", end)]["value"])
        interest_expense = Decimal(by_cell[("interest_expense", end)]["value"])
        pretax = Decimal(by_cell[("pretax_income", end)]["value"])
        tax = Decimal(by_cell[("income_tax_expense", end)]["value"])
        self.assertEqual(pretax, operating + interest_income - interest_expense)
        # The tax assumption explicitly names pretax as its company base.
        tax_result = next(row for row in result if row["role"] == "income_tax_expense")
        self.assertEqual(tax_result["cells"][0]["result_refs"], [
            {"ref": "result:pretax_income", "period_end": end},
        ])
        self.assertGreater(tax, 0)

        # A second company's validated DAG may classify the same magnitude as
        # an expense. The consumer follows the coefficient; it does not infer
        # the sign from a label or role.
        expense_structure = copy.deepcopy(structure)
        pretax_formula = next(item for item in expense_structure["formulas"]
                              if item["output_ref"] == "pretax")
        income_term = next(item for item in pretax_formula["terms"]
                           if item["line_ref"] == "interest-income")
        income_term["coefficient"] = "-1"
        expense = self.compute(expense_structure, inputs)
        expense_cell = cells(expense)[("pretax_income", end)]
        self.assertEqual(
            Decimal(expense_cell["value"]),
            operating - interest_income - interest_expense,
        )

    def test_missing_leaf_assumption_propagates_unavailable_without_zero_fill(self):
        inputs = financial_inputs()
        structure, replay = validate_financial_statement_structure(
            proposal(inputs), company_spec(), inputs)
        self.assertTrue(replay["ready_for_forecast"])
        results = self.compute(structure, inputs)
        by_cell = cells(results)
        end = sorted(end for role, end in by_cell if role == "operating_income")[0]
        self.assertEqual(by_cell[("operating_income", end)]["status"], "computed")
        self.assertEqual(by_cell[("interest_income", end)]["status"], "unavailable")
        pretax = by_cell[("pretax_income", end)]
        self.assertEqual(pretax["status"], "unavailable")
        self.assertIsNone(pretax["value"])
        self.assertIn("interest-income", pretax["reason"])
        self.assertEqual(by_cell[("net_income", end)]["status"], "unavailable")
        self.assertEqual(by_cell[("diluted_eps", end)]["status"], "unavailable")

    def test_v03_model_freezes_structure_replay_and_binding(self):
        inputs, structure = self.authority()
        # Recreate the public replay from the exact proposal authority.
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"] if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            record = ForecastModelAuthority(store).publish(body)
        self.assertEqual(record["schema_version"], "0.3")
        self.assertEqual(record["forecast_structure_binding"], binding)
        self.assertEqual(record["financial_statement_structure"], structure)

    def test_v03_actualization_replays_exact_dag_and_eps_inputs(self):
        inputs, structure = self.authority()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"]
                            if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        spec = {**company_spec(), "decided_by": "automation:test",
                "financial_statement_structure": {
                    "schema_version": "0.1", "lines": candidate["lines"],
                    "formulas": candidate["formulas"],
                }}
        structure, replay = materialize_financial_statement_structure(spec, inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            spec, inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            authority = ForecastModelAuthority(store)
            prior = authority.publish(body)

            current = copy.deepcopy(inputs)
            current["periods"].append("2026-03-31")
            actuals = {
                "revenue": 1400, "cost": 840, "opex": 280, "operating": 280,
                "interest_income": 10, "interest_expense": 20, "pretax": 270,
                "tax": 60, "net": 210, "nci": 5, "parent": 205,
                "eps_numerator": 205, "shares": 100, "eps": "2.05",
            }
            for line in current["filed_lines"]:
                line["cells"]["2026-03-31"] = {
                    "period_start": "2026-01-01",
                    "value": str(actuals[line["concept"]]),
                    "unit": next(iter(line["cells"].values()))["unit"],
                    "basis": "reported", "source_accessions": ["0000000000-26-000002"],
                }
            current_structure, current_replay = materialize_financial_statement_structure(
                spec, current)
            current_binding = forecast_structure_binding(
                current_structure, current_replay, current)
            self.assertNotEqual(current_structure["structure_ref"],
                                structure["structure_ref"])
            updated = actualize_model(
                prior, current, structure=current_structure, replay=current_replay,
                binding=current_binding)
            self.assertIsNotNone(updated)
            record = authority.publish(updated)

        by_cell = cells(record["results"])
        for role, value in (("operating_income", "280.00000000"),
                            ("pretax_income", "270.00000000"),
                            ("parent_net_income", "205.00000000"),
                            ("diluted_eps", "2.05000000")):
            self.assertEqual(by_cell[(role, "2026-03-31")]["value"], value)
            self.assertEqual(by_cell[(role, "2026-03-31")]["kind"], "actual")
        eps = by_cell[("diluted_eps", "2026-03-31")]
        self.assertEqual(eps["result_refs"], [
            {"ref": "result:diluted_eps_numerator", "period_end": "2026-03-31"},
            {"ref": "result:diluted_weighted_average_shares",
             "period_end": "2026-03-31"},
        ])

    def test_v03_actualization_refuses_changed_formula_topology(self):
        inputs, structure = self.authority()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"]
                            if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            prior = ForecastModelAuthority(store).publish(body)
        changed = copy.deepcopy(structure)
        changed["created_by"] = "automation:other-structure-decision"
        changed_body = dict(changed)
        changed_body.pop("content_hash")
        from dalton_core.store import content_hash
        changed["content_hash"] = content_hash(changed_body)
        changed_replay = replay_historical_structure(changed, inputs)
        changed_binding = forecast_structure_binding(changed, changed_replay, inputs)
        with self.assertRaisesRegex(ForecastModelUnavailable, "structure changed"):
            actualize_model(
                prior, inputs, structure=changed, replay=changed_replay,
                binding=changed_binding)

    def test_v03_sensitivity_recomputes_the_frozen_dag_and_exact_share_base(self):
        inputs, structure = self.authority()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"]
                            if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            record = ForecastModelAuthority(store).publish(body)

        tax_driver = next(item for item in record["drivers"]
                          if item["structure_line_ref"] == "tax")
        series = measure_series(record, tax_driver["ref"], "share_of_line")
        self.assertEqual(series["status"], "available")
        self.assertEqual(len(series["points"]), 4)
        # The tax line is divided by the structure's pretax line, not revenue
        # or the legacy operating-income shortcut.
        self.assertEqual(Decimal(series["points"][-1]["value"]), Decimal(55) / 250)

        interest_driver = next(item for item in record["drivers"]
                               if item["structure_line_ref"] == "interest-income")
        base_results, _ = recompute(record, None, None)
        moved_results, replaced = recompute(record, interest_driver["ref"], Decimal("0.02"))
        first = record["forecast_periods"][0]["end"]
        base_pretax = Decimal(cells(base_results)[("pretax_income", first)]["value"])
        moved_pretax = Decimal(cells(moved_results)[("pretax_income", first)]["value"])
        self.assertGreater(moved_pretax, base_pretax)
        self.assertEqual(len(replaced), len(record["forecast_periods"]))
        projection = build_projection(record)
        self.assertEqual(projection["formula_hash"], record["formula_hash"])
        self.assertGreaterEqual(len(projection["drivers"]), 3)

    def test_v03_revision_recomputes_same_structure_and_preserves_authority(self):
        inputs, structure = self.authority()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"]
                            if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            authority = ForecastModelAuthority(store)
            prior = authority.publish(body)
            tax_driver = next(item for item in prior["drivers"]
                              if item["structure_line_ref"] == "tax")
            end = prior["forecast_periods"][0]["end"]
            revised = revise_assumptions(
                prior, [{"driver": tax_driver["ref"], "period": end,
                         "value": "0.219", "because": "new filed tax guidance",
                         "refs": [{"kind": "filing", "ref": None, "concept": "tax",
                                   "period_end": "2025-12-31",
                                   "accession": ACCESSION}]}],
                change_reason="assumption_review",
                evidence_refs=[{"kind": "filing", "ref": None, "concept": "tax",
                                "period_end": "2025-12-31", "accession": ACCESSION}],
                actor_ref="automation:test")
            record = authority.publish(revised)
        before = Decimal(cells(prior["results"])[("net_income", end)]["value"])
        after = Decimal(cells(record["results"])[("net_income", end)]["value"])
        self.assertLess(after, before)
        self.assertEqual(record["financial_statement_structure"], structure)
        self.assertEqual(record["formula_hash"], prior["formula_hash"])

    def test_v03_published_line_keeps_exact_structure_formula_authority(self):
        inputs, structure = self.authority()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"]
                            if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            record = ForecastModelAuthority(store).publish(body)
            lines = ModelForecastAuthority(store)
            published = publish_forecast_lines(lines, record)
            held = lines.line(published[0]["version_ref"])
        self.assertEqual(held["formula_ref"], STRUCTURED_DRIVER_FORMULA_REF)
        self.assertEqual(held["formula_hash"], STRUCTURED_DRIVER_FORMULA_HASH)
        self.assertEqual(held["scenario_version_ref"], record["id"])
        self.assertEqual(held["scenario_version_hash"], record["content_hash"])

    def test_v03_workbook_translates_company_dag_and_refuses_eps_annual_sum(self):
        from openpyxl import load_workbook

        inputs, structure = self.authority()
        candidate = proposal(inputs)
        for line in candidate["lines"]:
            original = next(item for item in structure["lines"]
                            if item["ref"] == line["ref"])
            line["forecast_method"] = original["forecast_method"]
            line["forecast_base_ref"] = original["forecast_base_ref"]
        structure, replay = validate_financial_statement_structure(
            candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(
            company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            record = ForecastModelAuthority(store).publish(body)
            path = Path(temporary) / "structured.xlsx"
            calendar = {
                "calendar_ref": "calendar:test", "source_hash": "c" * 64,
                "as_of": "2026-09-11", "fiscal_year_end_month": 12,
            }
            calendar["content_hash"] = content_hash(calendar)
            export_fund_workbook(
                path, model=record, spec=company_spec(), inputs=inputs,
                calendar_binding=calendar)
            book = load_workbook(path, data_only=False)
        financials = book["Financials"]
        pretax_row = next(row for row in range(5, financials.max_row + 1)
                          if str(financials.cell(row, 1).value).startswith("Pretax"))
        formulas = [financials.cell(pretax_row, column).value
                    for column in range(2, financials.max_column + 1)]
        self.assertTrue(any(isinstance(value, str) and "+" in value and "-" in value
                            for value in formulas), formulas)
        eps_row = 5 + next(index for index, result in enumerate(record["results"])
                           if result["role"] == "diluted_eps")
        # No annual EPS is made by summing per-share quarters.
        self.assertIsNone(financials.cell(eps_row, 2).value)


if __name__ == "__main__":
    unittest.main()
