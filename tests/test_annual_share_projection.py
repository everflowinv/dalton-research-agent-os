from __future__ import annotations

import copy
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from dalton_core.company_financial_statement_structure import (
    forecast_structure_binding, validate_financial_statement_structure,
)
from dalton_core.company_model_annual_projection import (
    AnnualProjectionError, _forecast_outcomes, build_annual_projection,
    validate_annual_projection,
)
from dalton_core.fund_xlsx_export import export_fund_workbook
from dalton_core.model_forecast_driver import build_structured_forecast_model, ForecastModelAuthority
from dalton_core.company_model_report import render_forecast_model
from dalton_core.store import DaltonStore, content_hash
from tests.test_financial_structure_forecast_consumer import (
    ACCESSION, annual_authority_inputs, annual_forecastable_proposal, company_spec,
)


class AnnualShareProjectionTests(unittest.TestCase):
    def fixture(self):
        inputs = annual_authority_inputs()
        shares = next(line for line in inputs["filed_lines"] if line["concept"] == "shares")
        shares["duration_facts"].append({
            "period_start": "2024-01-01", "period_end": "2024-12-31",
            "period_kind": "cumulative", "value": "90", "unit": "shares",
            "source_accessions": [ACCESSION], "source_forms": ["10-K"],
        })
        candidate = annual_forecastable_proposal(inputs)
        next(line for line in candidate["lines"]
             if line["role"] == "diluted_weighted_average_shares")["annual_forecast_method"] = "annual_growth_from_filed"
        structure, replay = validate_financial_statement_structure(candidate, company_spec(), inputs)
        binding = forecast_structure_binding(structure, replay, inputs)
        body = build_structured_forecast_model(company_spec(), inputs, structure=structure, replay=replay, binding=binding)
        calendar = {"calendar_ref": "calendar:test", "source_hash": "c" * 64,
                    "as_of": "2025-12-31", "fiscal_year_end_month": 12}
        calendar["content_hash"] = content_hash(calendar)
        return inputs, structure, body, calendar

    def test_independent_annual_projection_report_and_workbook_share_exact_source(self):
        from openpyxl import load_workbook
        inputs, structure, body, calendar = self.fixture()
        with tempfile.TemporaryDirectory() as temporary:
            store = DaltonStore(str(Path(temporary) / "core.sqlite"))
            self.addCleanup(store.close)
            model = ForecastModelAuthority(store).publish(body)
            projection = build_annual_projection(model=model, inputs=inputs, calendar_binding=calendar)
            periods = {p["label"]: p for p in projection["periods"]}
            shares = periods["FY2026E"]["forecast_shares"]
            self.assertEqual(shares["method"], "annual_growth_from_filed")
            self.assertEqual(shares["as_of"], "2025-12-31")
            self.assertEqual(shares["steps"], 1)
            self.assertAlmostEqual(Decimal(shares["value"]), Decimal(1000) / 9)
            eps = periods["FY2026E"]["forecast_eps"]
            self.assertEqual(len(eps["numerator"]["source_periods"]), 4)
            self.assertEqual(Decimal(eps["value"]), Decimal(eps["numerator"]["value"]) / Decimal(shares["value"]))
            text = render_forecast_model(model, annual_projection=projection)
            self.assertIn("年度股份独立增长假设", text)
            self.assertIn("复利 1 年", text)
            target = Path(temporary) / "annual.xlsx"
            exported = export_fund_workbook(target, model=model, spec=company_spec(), inputs=inputs,
                                            calendar_binding=calendar, annual_projection=projection)
            self.assertEqual(exported["annual_projection_hash"], projection["content_hash"])
            book = load_workbook(target, data_only=False)
            rows = {result["role"]: 3 + i for i, result in enumerate(model["results"])}
            self.assertIsInstance(book["Financials"].cell(rows["diluted_weighted_average_shares"], 6).value, (int, float))
            self.assertTrue(str(book["Financials"].cell(rows["diluted_eps"], 6).value).startswith("="))
            self.assertTrue(any(row[0] == "AnnualShareGrowthAssumption" and row[3] == "2025-12-31"
                                for row in book["Sources"].values))
            forged = copy.deepcopy(projection)
            next(p for p in forged["periods"] if p["label"] == "FY2026E")["forecast_shares"]["steps"] = 2
            with self.assertRaises(AnnualProjectionError):
                validate_annual_projection(forged, model=model, inputs=inputs, calendar_binding=calendar)

    def test_missing_common_income_quarter_blocks_eps_even_when_shares_are_available(self):
        inputs, structure, model, calendar = self.fixture()
        numerator = next(row for row in model["results"] if row["role"] == "diluted_eps_numerator")
        group = ["2026-03-31", "2026-06-30", "2026-09-30", "2026-12-31"]
        cell = next(cell for cell in numerator["cells"] if cell["period"]["end"] == group[-1])
        cell.update(status="unavailable", value=None)
        shares, eps = _forecast_outcomes(model, structure, "FY2026E", group, calendar, inputs=inputs,
                                         historical_values={}, historical_refs={}, historical_starts={})
        self.assertEqual(shares["status"], "computed")
        self.assertEqual(eps["status"], "unavailable")
        self.assertIsNone(eps["value"])
        bad = copy.deepcopy(model)
        next(item for item in bad["financial_statement_structure_replay"]["forecast_methods"]
             if item.get("annual_forecast_source"))["annual_forecast_source"]["base_value"] = "1"
        with self.assertRaisesRegex(AnnualProjectionError, "source differs"):
            _forecast_outcomes(bad, structure, "FY2026E", group, calendar, inputs=inputs,
                               historical_values={}, historical_refs={}, historical_starts={})
