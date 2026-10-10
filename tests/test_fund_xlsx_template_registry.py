"""Fund Excel template registry / select (vision-gap #4, 2026-10-07).

The first three tests are the 2026-09-13 candidate's (``8e66ee01``,
``tests/test_fund_xlsx_style_registry.py``) rewritten against the Core-backed
registry: registration is inert until chosen, automation can neither register
nor choose, and a selected style flows through the renderer while the
built-in plan keeps its identity.
"""

from __future__ import annotations

import base64
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from dalton_core.company_model_inputs import build_model_inputs
from dalton_core.fund_xlsx_export import export_fund_workbook
from dalton_core.fund_xlsx_template import (
    STYLE_RESOURCE_SHA256,
    apply_fund_xlsx_template,
    build_fund_xlsx_template_plan,
    load_fund_xlsx_template_style,
)
from dalton_core.fund_xlsx_template_import import import_fund_xlsx_template_candidate
from dalton_core.fund_xlsx_template_registry import (
    BUILTIN_VERSION_REF,
    FundXlsxTemplateRegistry,
    FundXlsxTemplateRegistryConflict,
    FundXlsxTemplateRegistryError,
    builtin_version,
    export_binding,
    resolve_template,
    selection_state,
    template_versions,
    validate_mapping,
)
from dalton_core.model_forecast_driver import ForecastModelAuthority
from dalton_core.store import DaltonStore, content_hash
from tests.test_fund_xlsx_template_import import FundXlsxTemplateImportTests
from tests.test_model_forecast_driver import ledger, model, spec

OWNER = "human:owner"
ACME = "company:sec-cik:0000000001"
BOLT = "company:sec-cik:0000000002"
CORD = "company:sec-cik:0000000003"


def owner_workbook(path: Path) -> None:
    FundXlsxTemplateImportTests.workbook(None, path)  # type: ignore[arg-type]


def candidate(root: Path, slug: str = "owner-reference") -> dict:
    path = root / f"{slug}.xlsx"
    if not path.exists():
        owner_workbook(path)
    return import_fund_xlsx_template_candidate(
        path, sheets={"valuation": "SecretCo Valuation", "financials": "SecretCo Financials",
                      "driver": "SecretCo Drivers"}, template_slug=slug)


def dossier_table(store: DaltonStore, rows: dict[str, str]) -> None:
    store.connection.execute(
        "CREATE TABLE IF NOT EXISTS company_dossier_versions("
        "company_ref TEXT, version_number INTEGER, record_json TEXT)")
    for ref, word in rows.items():
        store.connection.execute(
            "INSERT INTO company_dossier_versions VALUES(?,?,?)",
            (ref, 1, json.dumps({"industry_classification": {"classification": word}})))
    store.connection.commit()


class RegistryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.store = DaltonStore(":memory:")
        self.addCleanup(self.store.close)

    # -- recovered from 8e66ee01 ---------------------------------------------

    def test_registered_template_is_inert_until_the_owner_chooses_it(self) -> None:
        registry = FundXlsxTemplateRegistry(self.store)
        first = registry.register(candidate=candidate(self.root), display_name="TMT 模板",
                                  reason="owner sample", actor_ref=OWNER)
        self.assertEqual(("registered", 1), (first["status"], first["version"]))
        self.assertEqual("fund-xlsx-template-style:owner-reference:1", first["style_ref"])
        self.assertEqual("builtin", resolve_template(self.store.connection, ACME)["basis"])
        registry.choose(scope_kind="default", scope_ref="default", reason="use it",
                        actor_ref=OWNER, template_version_ref=first["id"],
                        template_version_hash=first["content_hash"])
        chosen = resolve_template(self.store.connection, ACME)
        self.assertEqual((first["id"], "default"), (chosen["template"]["id"], chosen["basis"]))
        second = registry.register(candidate=candidate(self.root), display_name="TMT 模板",
                                   reason="revised", actor_ref=OWNER)
        self.assertEqual((2, first["id"]), (second["version"], second["prior_version_ref"]))
        self.assertTrue(second["style_ref"].endswith(":2"))
        # A new version does not move a choice that names the old one.
        self.assertEqual(first["id"], resolve_template(self.store.connection, ACME)["template"]["id"])

    def test_automation_can_neither_register_nor_choose(self) -> None:
        registry = FundXlsxTemplateRegistry(self.store)
        with self.assertRaisesRegex(FundXlsxTemplateRegistryError, "person"):
            registry.register(candidate=candidate(self.root), display_name="x",
                              reason="r", actor_ref="automation:worker")
        builtin = builtin_version()
        with self.assertRaisesRegex(FundXlsxTemplateRegistryError, "person"):
            registry.choose(scope_kind="default", scope_ref="default", reason="r",
                            actor_ref="automation:worker",
                            template_version_ref=builtin["id"],
                            template_version_hash=builtin["content_hash"])
        tampered = dict(candidate(self.root))
        tampered["activation_authorized"] = True
        with self.assertRaisesRegex(FundXlsxTemplateRegistryError, "not inert"):
            registry.register(candidate=tampered, display_name="x", reason="r", actor_ref=OWNER)

    def test_selected_style_flows_through_renderer_and_builtin_plan_remains(self) -> None:
        from openpyxl import Workbook

        record = FundXlsxTemplateRegistry(self.store).register(
            candidate=candidate(self.root), display_name="x", reason="r", actor_ref=OWNER)
        arguments = dict(sheet_names={"valuation": "Valuation", "financials": "Financials",
                                      "driver": "Driver"},
                         annual_periods=["FY2025A"], quarterly_periods=["Q1 FY2026E"],
                         unit_labels={"financials": "(USD MM)", "driver": "(USD MM)"})
        selected = build_fund_xlsx_template_plan(**arguments,
                                                 style_contract=record["style_contract"])
        builtin = build_fund_xlsx_template_plan(**arguments)
        self.assertEqual(record["style_ref"], selected["template_ref"])
        self.assertEqual(record["style_hash"], selected["style_resource_sha256"])
        self.assertEqual("fund-xlsx-template-style:amzn-source:0.1", builtin["template_ref"])
        self.assertEqual(STYLE_RESOURCE_SHA256, builtin["style_resource_sha256"])
        book = Workbook(); book.active.title = "Valuation"
        book.create_sheet("Financials"); book.create_sheet("Driver")
        apply_fund_xlsx_template(book, selected, style_contract=record["style_contract"])
        self.assertEqual(["Valuation", "Financials", "Driver"], book.sheetnames)
        # A plan bound to one style cannot be applied with another.
        with self.assertRaises(Exception):
            apply_fund_xlsx_template(book, selected)

    # -- the Core-backed authority ---------------------------------------------

    def test_builtin_is_seeded_by_code_and_a_bare_core_reads_it(self) -> None:
        bare = sqlite3.connect(":memory:")
        self.addCleanup(bare.close)
        self.assertEqual([BUILTIN_VERSION_REF], [r["id"] for r in template_versions(bare)])
        self.assertEqual("builtin", resolve_template(bare, ACME)["basis"])
        FundXlsxTemplateRegistry(self.store)
        FundXlsxTemplateRegistry(self.store)  # idempotent
        rows = self.store.connection.execute(
            "SELECT version_id, actor_ref FROM fund_xlsx_template_versions").fetchall()
        self.assertEqual([(BUILTIN_VERSION_REF, "system:builtin")], [tuple(r) for r in rows])
        self.assertEqual(builtin_version(), template_versions(self.store.connection)[0])
        with self.assertRaisesRegex(sqlite3.DatabaseError, "append-only"):
            self.store.connection.execute(
                "UPDATE fund_xlsx_template_versions SET style_hash='x'")
        with self.assertRaisesRegex(sqlite3.DatabaseError, "DaltonStore"):
            self.store.connection.execute(
                "INSERT INTO fund_xlsx_template_selections(receipt_id,sequence,scope_kind,"
                "scope_ref,action,reason,actor_ref,record_json,content_hash,created_at) "
                "VALUES('x',1,'default','default','clear','r','human:x','{}','h','t')")

    def test_select_order_is_explicit_company_industry_default_builtin(self) -> None:
        registry = FundXlsxTemplateRegistry(self.store)
        dossier_table(self.store, {ACME: "contract_compounder", BOLT: "contract_compounder",
                                   CORD: "insufficient_evidence"})
        a = registry.register(candidate=candidate(self.root, "tpl-a"), display_name="A",
                              reason="r", actor_ref=OWNER, industries=["contract_compounder"])
        b = registry.register(candidate=candidate(self.root, "tpl-b"), display_name="B",
                              reason="r", actor_ref=OWNER)
        c = registry.register(candidate=candidate(self.root, "tpl-c"), display_name="C",
                              reason="r", actor_ref=OWNER)

        def choose(kind, ref, record):
            return registry.choose(scope_kind=kind, scope_ref=ref, reason="r", actor_ref=OWNER,
                                   template_version_ref=None if record is None else record["id"],
                                   template_version_hash=None if record is None
                                   else record["content_hash"])

        conn = self.store.connection
        choose("default", "default", c)
        choose("industry", "contract_compounder", a)
        choose("company", BOLT, b)
        self.assertEqual(("tpl-a", "industry"), (resolve_template(conn, ACME)["template"]["template_slug"],
                                                 resolve_template(conn, ACME)["basis"]))
        self.assertEqual("company", resolve_template(conn, BOLT)["basis"])
        self.assertEqual("tpl-b", resolve_template(conn, BOLT)["template"]["template_slug"])
        # insufficient_evidence is no classification: the workspace default.
        self.assertEqual(("default", None), (resolve_template(conn, CORD)["basis"],
                                             resolve_template(conn, CORD)["classification"]))
        explicit = resolve_template(conn, BOLT, explicit_version_ref=BUILTIN_VERSION_REF)
        self.assertEqual(("explicit", BUILTIN_VERSION_REF),
                         (explicit["basis"], explicit["template"]["id"]))
        cleared = choose("company", BOLT, None)
        self.assertEqual("applied", cleared["outcome"])
        self.assertEqual("industry", resolve_template(conn, BOLT)["basis"])
        self.assertEqual("was_not_selected", choose("company", BOLT, None)["outcome"])
        choose("default", "default", None)
        self.assertEqual("builtin", resolve_template(conn, CORD)["basis"])
        self.assertEqual(6, selection_state(conn)["receipt_count"])

    def test_a_stale_hash_unknown_scope_and_key_reuse_are_refused(self) -> None:
        registry = FundXlsxTemplateRegistry(self.store)
        builtin = builtin_version()
        with self.assertRaisesRegex(FundXlsxTemplateRegistryError, "changed"):
            registry.choose(scope_kind="default", scope_ref="default", reason="r",
                            actor_ref=OWNER, template_version_ref=builtin["id"],
                            template_version_hash="0" * 64)
        with self.assertRaisesRegex(FundXlsxTemplateRegistryError, "industry"):
            registry.choose(scope_kind="industry", scope_ref="banks", reason="r",
                            actor_ref=OWNER, template_version_ref=builtin["id"],
                            template_version_hash=builtin["content_hash"])
        first = registry.choose(scope_kind="company", scope_ref=ACME, reason="r",
                                actor_ref=OWNER, template_version_ref=builtin["id"],
                                template_version_hash=builtin["content_hash"],
                                idempotency_key="k1")
        again = registry.choose(scope_kind="company", scope_ref=ACME, reason="r",
                                actor_ref=OWNER, template_version_ref=builtin["id"],
                                template_version_hash=builtin["content_hash"],
                                idempotency_key="k1")
        self.assertEqual(("replayed", first["id"]), (again["status"], again["id"]))
        with self.assertRaises(FundXlsxTemplateRegistryConflict):
            registry.choose(scope_kind="company", scope_ref=BOLT, reason="r",
                            actor_ref=OWNER, template_version_ref=builtin["id"],
                            template_version_hash=builtin["content_hash"],
                            idempotency_key="k1")

    def test_mapping_is_closed(self) -> None:
        self.assertEqual(validate_mapping(None), builtin_version()["mapping"])
        for bad in ({"sheet_names": {"financials": "Sources"}},
                    {"sheet_names": {"driver": "Bad:Name"}},
                    {"sheet_names": {"driver": "Valuation"}},
                    {"unit_labels": {"financials": "{ticker} MM"}},
                    {"valuation_annual_blocks": ["revenue", "revenue"]},
                    {"valuation_annual_blocks": ["free_cash_flow"]},
                    {"valuation_snapshot_cells": {"price": "A9"}},
                    {"valuation_snapshot_cells": {"price": "C2", "market_cap": "D2"}},
                    {"valuation_snapshot_cells": {"price": "A3"}},
                    {"colour": "blue"}):
            with self.subTest(bad=bad), self.assertRaises(FundXlsxTemplateRegistryError):
                validate_mapping(bad)


class ExportWithTemplateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.store = DaltonStore(":memory:")
        self.addCleanup(self.store.close)
        self.specification = spec(quarters=8)
        self.inputs = build_model_inputs(ledger(), self.specification)
        self.model = ForecastModelAuthority(self.store).publish(
            model(specification=self.specification))

    def calendar(self) -> dict:
        value = {"calendar_ref": "fiscal-calendar:acn:2026-09-10", "source_hash": "1" * 64,
                 "as_of": "2026-08-31", "fiscal_year_end_month": 8}
        return {**value, "content_hash": content_hash(value)}

    def export(self, name: str, template=None) -> tuple[dict, Path]:
        path = self.root / name
        result = export_fund_workbook(path, model=self.model, spec=self.specification,
                                      inputs=self.inputs, calendar_binding=self.calendar(),
                                      template=template)
        return result, path

    @staticmethod
    def signature(path: Path, *, names=("Valuation", "Financials", "Driver")) -> dict:
        from openpyxl import load_workbook

        book = load_workbook(path, data_only=False)
        out = {}
        for role, name in zip(("valuation", "financials", "driver"), names):
            out[role] = {cell.coordinate: (cell.value, cell.number_format, cell.font.name,
                                           cell.font.sz, cell.font.b,
                                           str(cell.fill.fgColor.rgb))
                         for row in book[name].iter_rows() for cell in row
                         if cell.value is not None}
        return out

    def test_builtin_through_the_registry_writes_what_the_exporter_always_wrote(self) -> None:
        plain, plain_path = self.export("plain.xlsx")
        bound, bound_path = self.export(
            "bound.xlsx", export_binding({"template": builtin_version(), "basis": "builtin"}))
        self.assertEqual(self.signature(plain_path), self.signature(bound_path))
        self.assertEqual(plain["formula_map_hash"], bound["formula_map_hash"])
        self.assertEqual((BUILTIN_VERSION_REF, "builtin"),
                         (plain["template_version_ref"], plain["template_basis"]))
        from openpyxl import load_workbook
        rows = {row[0].value: row[1].value for row in load_workbook(plain_path)["Formula Map"].iter_rows()
                if row[0].value}
        self.assertEqual(BUILTIN_VERSION_REF, rows["Fund template"])
        self.assertEqual(STYLE_RESOURCE_SHA256, rows["Fund template style SHA-256"])

    def test_owner_template_renames_sheets_rewrites_references_and_moves_valuation_cells(self) -> None:
        from openpyxl import load_workbook

        record = FundXlsxTemplateRegistry(self.store).register(
            candidate=candidate(self.root), display_name="TMT", reason="r", actor_ref=OWNER,
            mapping={"sheet_names": {"valuation": "估值", "financials": "Fin Model",
                                     "driver": "Drivers"},
                     "unit_labels": {"financials": "{currency} in millions"},
                     "valuation_annual_blocks": ["revenue", "diluted_eps"],
                     "valuation_snapshot_cells": {"price": "G2"}})
        resolved = resolve_template(self.store.connection, self.model["company_ref"],
                                    explicit_version_ref=record["id"])
        plain, plain_path = self.export("plain.xlsx")
        result, path = self.export("owner.xlsx", export_binding(resolved))
        book = load_workbook(path, data_only=False)
        self.assertEqual(["估值", "Fin Model", "Drivers", "Sources", "Formula Map"], book.sheetnames)
        formulas = [cell.value for name in ("估值", "Fin Model", "Drivers")
                    for row in book[name].iter_rows() for cell in row
                    if isinstance(cell.value, str) and cell.value.startswith("=")]
        self.assertTrue(formulas)
        self.assertFalse([f for f in formulas if "'Financials'!" in f or "'Driver'!" in f])
        self.assertTrue(any("'Drivers'!" in f for f in formulas))
        self.assertEqual("USD in millions", book["Fin Model"]["A1"].value)
        valuation = book["估值"]
        self.assertEqual("Total Revenue (USD MM)", valuation["A8"].value)
        self.assertEqual("GAAP EPS (USD)", valuation["A12"].value)
        self.assertIsNone(valuation["A16"].value)
        self.assertTrue(str(valuation["G2"].value).startswith("Current Price"))
        self.assertNotEqual("Current Price (USD) — N/A", valuation["A4"].value)
        # Same numbers and formulas, only the sheet names differ.
        renamed = self.signature(path, names=("估值", "Fin Model", "Drivers"))
        original = self.signature(plain_path)
        def norm(sig):
            return {coord: (str(v[0]).replace("'Fin Model'!", "'Financials'!")
                            .replace("'Drivers'!", "'Driver'!"), v[1])
                    for coord, v in sig.items() if coord != "A1"}
        self.assertEqual(norm(original["financials"]), norm(renamed["financials"]))
        self.assertEqual(norm(original["driver"]), norm(renamed["driver"]))
        cells = [row[0].value for row in book["Formula Map"].iter_rows(min_row=5) if row[0].value]
        self.assertTrue(any(str(c).startswith("Fin Model!") for c in cells))
        self.assertFalse(any(str(c).startswith("Financials!") for c in cells))
        self.assertEqual((record["id"], "explicit", record["style_ref"]),
                         (result["template_version_ref"], result["template_basis"],
                          result["template_style_ref"]))


class WriterAndCockpitTests(unittest.TestCase):
    """The two governance doors and the cockpit card, end to end in-process."""

    def test_writer_doors_bind_the_owner_and_refuse_outside_companies(self) -> None:
        from dalton_core.errors import ValidationError
        from dalton_core.writer_server import (
            HUMAN_GOVERNANCE_OPERATIONS, OPERATION_FIELDS, WriterServer,
        )

        for op in ("register_fund_xlsx_template", "select_fund_xlsx_template"):
            self.assertIn(op, HUMAN_GOVERNANCE_OPERATIONS)
            self.assertIn("actor_ref", OPERATION_FIELDS[op])
        with tempfile.TemporaryDirectory() as name:
            root = Path(name).resolve()
            store = DaltonStore(":memory:")
            self.addCleanup(store.close)
            server = object.__new__(WriterServer)
            server._store = store
            server._now = lambda: __import__("datetime").datetime.now(
                __import__("datetime").timezone.utc)
            server._active_universe_refs = lambda: {ACME}
            out = server._op_register_fund_xlsx_template({
                "candidate": candidate(root), "display_name": "TMT", "reason": "r",
                "actor_ref": OWNER, "industries": ["structural_growth"]})
            self.assertEqual("registered", out["status"])
            template = out["template"]
            picked = server._op_select_fund_xlsx_template({
                "scope_kind": "company", "scope_ref": ACME, "reason": "r", "actor_ref": OWNER,
                "template_version_ref": template["id"],
                "template_version_hash": template["content_hash"]})
            self.assertEqual("applied", picked["receipt"]["outcome"])
            with self.assertRaisesRegex(ValidationError, "coverage list"):
                server._op_select_fund_xlsx_template({
                    "scope_kind": "company", "scope_ref": BOLT, "reason": "r",
                    "actor_ref": OWNER, "template_version_ref": template["id"],
                    "template_version_hash": template["content_hash"]})
            with self.assertRaisesRegex(ValidationError, "person"):
                server._op_select_fund_xlsx_template({
                    "scope_kind": "default", "scope_ref": "default", "reason": "r",
                    "actor_ref": "automation:x", "template_version_ref": template["id"],
                    "template_version_hash": template["content_hash"]})

    def test_cockpit_card_lists_uploads_and_selects_through_the_writer(self) -> None:
        from contextlib import contextmanager

        from dalton_core import cockpit_fund_templates as card
        from dalton_core.cockpit_plane import CockpitError

        with tempfile.TemporaryDirectory() as name:
            root = Path(name).resolve()
            db = root / "core.sqlite"
            store = DaltonStore(db)
            self.addCleanup(store.close)
            registry = FundXlsxTemplateRegistry(store)
            calls: list[tuple[str, dict]] = []
            events: list[dict] = []
            members = {ACME: {"ticker": "ACME"}, BOLT: {"ticker": "BOLT"}}

            class Plane:
                journal = SimpleNamespace(record_event=lambda **kw: events.append(kw))

                @contextmanager
                def _core(self):
                    connection = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
                    connection.row_factory = sqlite3.Row
                    try:
                        yield connection
                    finally:
                        connection.close()

                def _mission(self, core):
                    return {"mission_ref": "coverage-mission:t", "universe": []}

                def _members(self, mission):
                    return members

                @staticmethod
                def _label(members, ref):
                    return members[ref]["ticker"]

                def clock(self):
                    from datetime import datetime, timezone
                    return datetime.now(timezone.utc)

                def _governance(self, login, operation, params, *, failure):
                    calls.append((operation, params))
                    if operation == "register_fund_xlsx_template":
                        record = registry.register(actor_ref=OWNER, **params)
                        return {"status": record.pop("status"), "template": record}
                    receipt = registry.choose(actor_ref=OWNER, **params)
                    return {"status": receipt.pop("status"), "receipt": receipt}

            plane = Plane()
            view = card.templates_view(plane)
            self.assertEqual([BUILTIN_VERSION_REF], [t["version_ref"] for t in view["templates"]])
            self.assertEqual({"builtin"}, {c["basis"] for c in view["companies"]})
            path = root / "upload.xlsx"
            owner_workbook(path)
            body = {"filename": "owner.xlsx",
                    "content_base64": base64.b64encode(path.read_bytes()).decode(),
                    "template_slug": "owner-tmt", "display_name": "TMT 组模板",
                    "sheets": {"valuation": "SecretCo Valuation",
                               "financials": "SecretCo Financials",
                               "driver": "SecretCo Drivers"},
                    "industries": ["structural_growth"], "reason": "组内统一", "request_id": "r1"}
            out = card.register(plane, "owner", body)
            self.assertEqual("registered", out["status"])
            sent = calls[-1][1]
            self.assertNotIn("content_base64", sent)
            self.assertNotIn("SecretCo Holdings", json.dumps(sent, ensure_ascii=False))
            with self.assertRaisesRegex(CockpitError, "找不到"):
                card.register(plane, "owner", {**body, "sheets": {**body["sheets"],
                                                                  "driver": "Nope"},
                                               "request_id": "r2"})
            view = card.templates_view(plane)
            uploaded = next(t for t in view["templates"] if t["slug"] == "owner-tmt")
            card.select(plane, "owner", {"scope_kind": "company", "scope_ref": ACME,
                                         "template_version_ref": uploaded["version_ref"],
                                         "template_version_hash": uploaded["version_hash"],
                                         "reason": "试用", "request_id": "r3"})
            view = card.templates_view(plane)
            acme = next(c for c in view["companies"] if c["company_ref"] == ACME)
            self.assertEqual(("company", uploaded["version_ref"]),
                             (acme["basis"], acme["template_version_ref"]))
            self.assertEqual(["ACME"], next(t for t in view["templates"]
                                            if t["slug"] == "owner-tmt")["used_by"])
            self.assertEqual(1, len(view["receipts"]))
            self.assertEqual(2, len(events))


if __name__ == "__main__":
    unittest.main()
