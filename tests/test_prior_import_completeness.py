from __future__ import annotations

import copy
import json
import tempfile
import unittest
import zipfile
from unittest import mock
from pathlib import Path

from dalton_core.prior_import_assets import docx_artifact_manifest, xlsx_artifact_manifest
from dalton_core.prior_model_import import (
    PriorModelAuthority,
    WorkbookReadBudget,
    read_workbook,
    workbook_digest,
)
from dalton_core.store import DaltonStore
from dalton_core.prior_research_core import MANIFEST_NAME, document_ref, read_document_artifact
from dalton_core import prior_research_core
from dalton_core.connector_governance import (
    PRIOR_RESEARCH_GET_KIND, build_governance_record,
)
from dalton_core.raw_spool import RawSpool
from dalton_core import prior_research_cli


class WorkbookCompletenessTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / "large.xlsx"
        from openpyxl import Workbook
        book = Workbook()
        for index, name in enumerate(("Support", "AMZN Driver", "AMZN Financials")):
            sheet = book.active if index == 0 else book.create_sheet()
            sheet.title = name
            sheet["A1"], sheet["B1"], sheet["C1"] = "Driver", "2025", "3Q26E"
            for row in range(2, 12):
                sheet.cell(row, 1, f"metric {row}")
                sheet.cell(row, 2, row)
                sheet.cell(row, 3, f"=B{row}*(1+10%)")
                font = copy.copy(sheet.cell(row, 3).font)
                font.color = "008000"
                sheet.cell(row, 3).font = font
        book.save(self.path)

    def test_default_read_covers_late_sheets_and_preserves_roles(self):
        artifact = read_workbook(self.path)
        self.assertTrue(artifact.metadata["complete"])
        self.assertEqual([s["name"] for s in artifact.metadata["sheets"]],
                         ["Support", "AMZN Driver", "AMZN Financials"])
        last = next(row for row in artifact if row["sheet"] == "AMZN Financials"
                    and row["cell"] == "C11")
        self.assertEqual(last["formula"], "=B11*(1+10%)")
        self.assertEqual((last["period"], last["period_status"]), ("3Q26E", "forecast"))
        self.assertEqual(last["period_basis"], "calendar_axis")
        self.assertEqual(last["source_type"], "prior_human_formula")

    def test_explicit_budget_reports_where_it_truncated(self):
        artifact = read_workbook(self.path, budget=WorkbookReadBudget(max_cells=5))
        self.assertFalse(artifact.metadata["complete"])
        self.assertEqual(artifact.metadata["truncations"][0]["reason"], "max_cells")
        self.assertEqual(artifact.metadata["truncations"][0]["at_sheet"], "Support")

    def test_generic_artifact_has_late_formulas_and_no_cached_values(self):
        manifest = xlsx_artifact_manifest(self.path)
        last = next(row for row in manifest["cells"]
                    if row["sheet"] == "AMZN Financials" and row["cell"] == "C11")
        self.assertEqual(last["formula"], "=B11*(1+10%)")
        self.assertNotIn("value", last)
        self.assertTrue(manifest["import_metadata"]["complete"])

    def test_enriched_record_replays_as_duplicate_and_never_calls_cells_actual(self):
        store = DaltonStore(str(self.root / "core.sqlite"))
        self.addCleanup(store.close)
        authority = PriorModelAuthority(store)
        kwargs = dict(company_ref="company:test", source_document_ref="prior:source",
                      as_of="2026-09-10", workbook_sha256=workbook_digest(self.path),
                      assumptions=read_workbook(self.path), actor_ref="human:owner")
        first = authority.publish(**kwargs)
        second = authority.publish(**kwargs)
        self.assertEqual((first["schema_version"], second["status"]), ("0.2", "duplicate"))
        self.assertTrue(first["import_metadata"]["complete"])
        self.assertEqual({row["kind"] for row in first["assumptions"]}, {"prior_human"})
        unknown = next(row for row in first["assumptions"] if row["cell"] == "B2")
        self.assertEqual(unknown["period_status"], "unknown")

    def test_numeric_data_is_not_mistaken_for_a_calendar(self):
        from openpyxl import load_workbook
        book = load_workbook(self.path)
        sheet = book["AMZN Financials"]
        sheet["B1"], sheet["C1"] = "2025A", "2026E"
        sheet["B2"] = 1200
        sheet["B3"], sheet["C3"] = "=B2+1", "=B2+2"
        book.save(self.path)
        artifact = read_workbook(self.path)
        by_cell = {(row["sheet"], row["cell"]): row for row in artifact}
        self.assertEqual(by_cell[("AMZN Financials", "B3")]["period"], "2025A")
        self.assertEqual(by_cell[("AMZN Financials", "B3")]["period_status"], "unknown")
        self.assertEqual(by_cell[("AMZN Financials", "C3")]["period"], "2026E")

    def test_formula_is_not_silently_shortened(self):
        from openpyxl import load_workbook
        book = load_workbook(self.path)
        formula = "=" + "+".join(["1"] * 1500)
        book["AMZN Driver"]["D2"] = formula
        book.save(self.path)
        row = next(row for row in read_workbook(self.path)
                   if row["sheet"] == "AMZN Driver" and row["cell"] == "D2")
        self.assertEqual(row["formula"], formula)

    def test_legacy_plain_rows_keep_the_old_wire_version(self):
        store = DaltonStore(str(self.root / "legacy.sqlite"))
        self.addCleanup(store.close)
        result = PriorModelAuthority(store).publish(
            company_ref="company:test", source_document_ref="prior:legacy",
            as_of="2026-09-10", workbook_sha256="a" * 64,
            assumptions=[{"sheet": "S", "cell": "B2", "label": "growth",
                          "value": "0.1", "formula": "", "unit": "percent",
                          "unit_basis": "label"}], actor_ref="human:owner")
        self.assertEqual(result["schema_version"], "0.1")
        self.assertNotIn("import_metadata", result)


class DocxArtifactTests(unittest.TestCase):
    def test_text_ceiling_returns_an_explicit_excerpt_and_preserves_source(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "long.md"
            path.write_text("x" * 100)
            with mock.patch.object(prior_research_core, "MAX_TEXT_CHARS", 80):
                text, _, _, raw = prior_research_core.read_body(path, relative_path="long.md")
            self.assertTrue(text.endswith("configured ceiling)"))
            self.assertEqual(raw, b"x" * 100)

    def test_body_order_and_assets_are_hashed_without_execution(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sample.docx"
            document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                        '<w:body><w:p><w:r><w:t>First</w:t></w:r></w:p>'
                        '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Table</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
                        '<w:p><w:r><w:drawing r:id="rId7"/></w:r></w:p></w:body></w:document>')
            rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId7" Type="http://x/image" Target="media/image1.png"/>'
                    '</Relationships>')
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/document.xml", document)
                archive.writestr("word/_rels/document.xml.rels", rels)
                archive.writestr("word/media/image1.png", b"not executed")
                archive.writestr("word/embeddings/model.xlsx", b"not opened")
            artifact = docx_artifact_manifest(path.read_bytes())
        self.assertEqual([row["kind"] for row in artifact["body"]],
                         ["paragraph", "table", "paragraph"])
        self.assertEqual([row["kind"] for row in artifact["assets"]],
                         ["embedded_workbook", "media"])
        self.assertTrue(all(len(row["sha256"]) == 64 for row in artifact["assets"]))

    def test_current_document_path_exposes_the_inert_artifact_manifest(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "TEST"
            folder.mkdir()
            path = folder / "memo.docx"
            document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        '<w:body><w:p><w:r><w:t>Memo text</w:t></w:r></w:p></w:body></w:document>')
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/document.xml", document)
                archive.writestr("word/embeddings/model.xlsx", b"opaque")
            (folder / MANIFEST_NAME).write_text(json.dumps({"documents": [{
                "path": "memo.docx", "kind": "memo", "as_of": "2026-09-10",
                "author": "human:analyst", "source_note": "sample"}]}))
            ref = document_ref("TEST", "memo.docx")
            header, text, artifact = read_document_artifact(root, ref)
        self.assertEqual(text, "Memo text")
        self.assertEqual(header["document_id"], ref)
        self.assertEqual(artifact["assets"][0]["kind"], "embedded_workbook")

    def test_governed_get_spools_original_and_replayable_structure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            corpus, state = root / "corpus", root / "state"
            folder = corpus / "TEST"
            folder.mkdir(parents=True)
            state.mkdir()
            path = folder / "memo.docx"
            document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                        '<w:body><w:p><w:hyperlink r:id="rId9"><w:r><w:t>Source</w:t></w:r></w:hyperlink></w:p>'
                        '<w:p><w:r><w:drawing r:id="rId7"/></w:r></w:p></w:body></w:document>')
            rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId7" Type="http://x/image" Target="media/image1.png"/>'
                    '<Relationship Id="rId9" Type="http://x/hyperlink" Target="https://private.invalid/x" TargetMode="External"/>'
                    '</Relationships>')
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/document.xml", document)
                archive.writestr("word/_rels/document.xml.rels", rels)
                archive.writestr("word/media/image1.png", b"image")
                archive.writestr("word/embeddings/model.xlsx", b"opaque workbook")
            original = path.read_bytes()
            (folder / MANIFEST_NAME).write_text(json.dumps({"documents": [{
                "path": "memo.docx", "kind": "memo", "as_of": "2026-09-10",
                "author": "human:analyst", "source_note": "sample"}]}))
            governance = state / "governance.json"
            governance.write_text(json.dumps(build_governance_record(
                PRIOR_RESEARCH_GET_KIND, approved_by="human:owner", status="approved")))
            args = prior_research_cli.build_parser().parse_args([
                "--state-dir", str(state), "--summary-dir", str(state),
                "--governance", str(governance), "--corpus-root", str(corpus),
                "--operation", "get_document", "--document-id", document_ref("TEST", "memo.docx")])
            summary = prior_research_cli.run(args)
            bundle = json.loads((state / "artifact-manifest.json").read_text())
            stored = RawSpool(state / "connector-spool", max_total_bytes=1_000_000_000).read_object(
                bundle["source_object"]["content_hash"])
        self.assertEqual(summary["status"], "succeeded", summary["failure_reason"])
        self.assertEqual(stored, original)
        external = next(ref for block in bundle["structure"]["body"]
                        for ref in block["relationships"] if ref["external"])
        self.assertIsNone(external["target"])
        self.assertEqual(len(external["target_sha256"]), 64)


if __name__ == "__main__":
    unittest.main()
